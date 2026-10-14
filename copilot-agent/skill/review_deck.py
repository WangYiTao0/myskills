"""生成后审查：机械检查文稿和 pptx，列出要改的地方。

运行：python scripts/review_deck.py deck.md 输出.pptx（在技能目录下）
每条问题一行：「[规则] 位置：问题 → 怎么改」。没有问题时最后一行打印 REVIEW OK。
审查只报告、不改文件，退出码总是 0（文稿本身读不了时为 1）。
规则来自 Copilot 实测里出现过的问题；需要判断的项（内容是否贴题、页序是否讲得通）
留给 SKILL.md 的审查清单，由模型逐条核对。
"""
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import layouts  # noqa: E402
import mdparse  # noqa: E402
from mdparse import DeckError  # noqa: E402

# 章节 → 不适合放的版式（目标、范围要写目标和边界，不放流程和改前改后；应用场景不放改前改后）
MISFIT = {"TARGET": {"process", "before-after"}, "SCOPE": {"process", "before-after"},
          "USE CASES": {"before-after"}}
VOICE = ["赋能", "打通", "闭环", "抓手", "拉通", "组合拳", "打造", "未来可期", "令人兴奋", "激动人心",
         "惊艳", "重磅", "大幅", "显著", "大量", "彻底解决", "完美", "全面", "零问题", "革命性",
         "散乱", "低效", "容易出错", "混乱"]
_NUMBER = re.compile(r"^\s*(?:约|~|≈|≥|≤|>|<|\+/-)?\s*[$+\-]?\d")
MAX_PAGES = 10


def where(page):
    return f"页「{page.title}」"


def english(title):
    return title.split("/")[-1].strip().upper() if "/" in title else ""


def chinese(title):
    return title.split("/")[0].strip()


def page_text(page):
    return "\n".join([page.title, page.action_title, *page.lines, page.next])


def review_md(doc):
    issues = []
    pages = doc.pages
    if len(pages) > MAX_PAGES:
        issues.append(("页数过多", "整份文稿", f"{len(pages)} 页",
                       f"合并重复的页，控制在 {MAX_PAGES} 页以内"))
    for prev, page in zip(pages, pages[1:]):
        if prev.layout == page.layout:
            issues.append(("相邻两页版式相同", where(page),
                           f"和上一页「{prev.title}」都是 {page.layout}，连着看像同一页",
                           "合并两页，或给其中一页换更贴合内容的版式"))
    seen_steps = []
    for page in pages:
        title = page.action_title.strip()
        if len(title) < 8 or title == chinese(page.title):
            issues.append(("目录式副标题", where(page), f"副标题「{title}」没有给出结论",
                           "改成一句结论，如「试点 8 周，检验单全部电子化」"))
        if not page.conclusion:
            issues.append(("缺结论行", where(page), "没有普通文字结论行，页底没有红条",
                           "在版式字段后加一行本页结论，最好带数字"))
        bad_layouts = MISFIT.get(english(page.title), set())
        if page.layout in bad_layouts:
            issues.append(("章节和版式不配", where(page),
                           f"{english(page.title)} 章节用了 {page.layout}",
                           "目标、范围写目标值和边界（stats / stats-side），流程和改前改后放到对应章节"))
        if page.layout == "process":
            names = {name for name, _ in page.fields["steps"]}
            for other_title, other in seen_steps:
                if len(names & other) * 2 >= min(len(names), len(other)):
                    issues.append(("内容重复", where(page),
                                   f"流程步骤和页「{other_title}」大半相同",
                                   "删掉重复的一页，或改写成不同的内容"))
            seen_steps.append((page.title, names))
        if page.layout in ("stats", "stats-side"):
            for item in page.fields["items"]:
                if "[数据缺失]" in item["value"]:
                    continue
                if not _NUMBER.match(item["value"]):
                    issues.append(("数字卡不是数字", where(page),
                                   f"卡片「{item['label']}」的数字栏是「{item['value']}」",
                                   "数字栏只放数字；文字内容改写进说明、结论要点或别的版式"))
        if page.layout in ("timeline", "timeline-phase"):
            keys = (page.fields["keys"] if page.layout == "timeline"
                    else [item[4] for item in page.fields["items"]])
            if len(keys) >= 3 and sum(keys) * 2 > len(keys):
                issues.append(("关键标记过多", where(page),
                               f"{sum(keys)}/{len(keys)} 个节点标了 !",
                               "只给 1–2 个真正关键的节点加 !"))
        text = page_text(page)
        if "[数据缺失]" in text:
            issues.append(("数据缺失", where(page), "有 [数据缺失] 占位",
                           "交付时告诉用户要补哪项数据；没有数据的指标不单独做数字卡"))
        words = [w for w in VOICE if w in text]
        if words:
            issues.append(("用词", where(page), "用了 " + "、".join(words),
                           "换成具体做了什么、结果是多少"))
    return issues


def review_pptx(doc, path):
    """确认 pptx 是生成脚本做的：页数对得上，标题填在模板的标题框里，banner 上没有另画的文本框。"""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(path)
    slides = list(prs.slides)
    closing = 1 if doc.meta.get("closing") or mdparse.is_monthly(doc.meta) else 0
    want = 1 + len(doc.pages) + closing
    issues = []
    if len(slides) != want:
        issues.append(("页数不符", "pptx", f"有 {len(slides)} 页，文稿应生成 {want} 页",
                       "用 generate_ppt.py 按当前文稿重新生成"))
    drawn = []
    for n, slide in enumerate(slides[1:1 + len(doc.pages)], start=2):
        titles = {ph.placeholder_format.idx: ph.text_frame.text.strip()
                  for ph in slide.placeholders if ph.has_text_frame}
        on_banner = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
                     and sh.top is not None and sh.top < 914400]
        if not titles.get(0) or on_banner:
            drawn.append(n)
    if drawn:
        issues.append(("不是生成脚本做的", "pptx",
                       f"第 {', '.join(map(str, drawn))} 页的标题没有填进模板标题框，或在 banner 上另画了文本框",
                       "不要自己写代码画幻灯片；改 md 文稿后用 generate_ppt.py 重新生成"))
    return issues


def main(argv):
    if len(argv) not in (1, 2):
        print("用法: python scripts/review_deck.py deck.md [输出.pptx]", file=sys.stderr)
        return 1
    try:
        doc = mdparse.parse(argv[0])
        for page in doc.pages:
            layouts.parse_page(page)
    except DeckError as e:
        print(f"错误: {e}", file=sys.stderr)
        return 1
    issues = review_md(doc)
    if len(argv) == 2:
        issues += review_pptx(doc, argv[1])
    for rule, place, problem, fix in issues:
        print(f"[{rule}] {place}：{problem} → {fix}")
    print(f"REVIEW: {len(issues)} 条问题" if issues else "REVIEW OK")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))

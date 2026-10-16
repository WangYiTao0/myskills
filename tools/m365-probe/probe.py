"""M365 Copilot 技能沙箱探针：只读环境、写一个输出文件，不联网、不装包。

运行：python3 probe.py   （在沙箱里由 agent 调用；本地也能跑）
打印 JSON 风格的逐项结果，最后一行是 PROBE DONE。
"""
import importlib
import os
import platform
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent


def line(key, value):
    print(f"{key}: {value}")


def try_import(name):
    try:
        mod = importlib.import_module(name)
        return f"OK {getattr(mod, '__version__', getattr(mod, 'VERSION', '?'))}"
    except Exception as e:  # noqa: BLE001 - 探针要把任何失败都报出来
        return f"FAIL {type(e).__name__}: {e}"


def main():
    line("python", sys.version.replace("\n", " "))
    line("executable", sys.executable)
    line("platform", platform.platform())
    line("cwd", os.getcwd())
    line("probe_file", __file__)
    line("skill_dir", HERE)
    line("skill_dir_files", sorted(p.name for p in HERE.iterdir()))
    # 跨 skill 引用：同级目录里能不能看到别的 skill（如 milwaukee-ppt）
    try:
        line("parent_dir", HERE.parent)
        line("parent_dir_entries", sorted(p.name for p in HERE.parent.iterdir())[:50])
    except Exception as e:  # noqa: BLE001
        line("parent_dir_entries", f"FAIL {e}")
    for env in ("HOME", "TMPDIR", "OUTPUT_DIR", "PYTHONPATH"):
        line(f"env.{env}", os.environ.get(env))

    for name in ("pptx", "lxml", "lxml.etree", "PIL", "xlsxwriter", "typing_extensions", "docx", "openpyxl"):
        line(f"import.{name}", try_import(name))

    sample = HERE / "sample.pptx"
    try:
        with zipfile.ZipFile(sample) as z:
            line("read_sample.zipfile", f"OK {len(z.namelist())} 个部件")
    except Exception as e:  # noqa: BLE001
        line("read_sample.zipfile", f"FAIL {e}")

    out_dir = Path(os.environ.get("OUTPUT_DIR") or os.getcwd())
    try:
        from pptx import Presentation
        prs = Presentation(str(sample))
        line("read_sample.pptx", f"OK {len(prs.slides)} 页，标题 {prs.slides[0].shapes.title.text!r}")
        prs.slides[0].shapes.title.text = "m365-probe output"
        out = out_dir / "m365-probe-output.pptx"
        prs.save(str(out))
        line("write_output", f"OK {out.resolve()}")
    except Exception as e:  # noqa: BLE001
        line("read_sample.pptx", f"FAIL {type(e).__name__}: {e}")
        try:
            out = out_dir / "m365-probe-output.txt"
            out.write_text("m365-probe output\n", encoding="utf-8")
            line("write_output", f"OK(txt) {out.resolve()}")
        except Exception as e2:  # noqa: BLE001
            line("write_output", f"FAIL {e2}")

    try:
        from importlib.metadata import distributions
        names = sorted({(d.metadata["Name"] or "?").lower() for d in distributions()})
        line("installed_packages", f"{len(names)} 个：{', '.join(names)}")
    except Exception as e:  # noqa: BLE001
        line("installed_packages", f"FAIL {e}")
    print("PROBE DONE")


if __name__ == "__main__":
    main()

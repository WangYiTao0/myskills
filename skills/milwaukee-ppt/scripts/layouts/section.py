"""section 版式（隔断页）：模板 layout 0「Blank Slide」，章节过渡页。

文稿写法：正好一行「编号 | 中文名 | 英文名」，三段都必填，如 `02 | 试点推进 | PILOT`。
隔断页不放结论红条：写了普通文字行或 Next: 都报错。Tag: 照常可用。

排版：正文区（1.30–5.60"）里一组横排、整体垂直居中——左边大号章节编号（96pt bold #1A1A1A），
一条 #BFBFBF 竖线，右边中文名（40pt bold #1A1A1A）在上、英文名（20pt #6B6B6B）在下。
只用黑白灰，banner 以外不用红。和 evidence 共用模板 layout 0，inspect 靠 MARK_SHAPE 区分。
"""
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

from mdparse import DeckError

from .common import (IMG_H, IMG_LEFT, IMG_TOP, IMG_W, INK, LINE_SPACING, MUTED, RULE, add_line,
                     add_page_chrome, add_run, add_textbox, chapter_row, check_chapter_row,
                     fill_titles, text_width_in)

NAME = "section"
LAYOUT_INDEX = 0
MARK_SHAPE = "Section Number"

NUMBER_PT, NAME_PT, ENGLISH_PT = 96, 40, 20
GROUP_LEFT = Inches(IMG_LEFT + 1.0)    # 整组从正文区左沿再缩进 1"
GAP = Inches(0.4)                      # 编号 → 竖线 → 名字
NAME_GAP = Inches(0.12)                # 中文名 → 英文名


def parse(page):
    found = []
    for line in page.lines:
        if not line:
            continue
        parts = chapter_row(line)
        if parts is None:
            raise DeckError(f"页「{page.title}」隔断页只写一行「编号 | 中文名 | 英文名」，"
                            f"不放结论红条；这一行不对: {line!r}")
        found.append(parts)
    if page.next:
        raise DeckError(f"页「{page.title}」隔断页不放结论红条，去掉 Next:")
    if len(found) != 1:
        raise DeckError(f"页「{page.title}」隔断页须正好一行「编号 | 中文名 | 英文名」，"
                        f"现在 {len(found)} 行")
    check_chapter_row(page, "隔断页", found[0])
    number, name, english = found[0]
    page.fields = {"number": number, "name": name, "english": english}


def render(prs, page, ctx):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    fill_titles(slide, page)
    f = page.fields
    mid = Inches(IMG_TOP + IMG_H / 2)
    right = Inches(IMG_LEFT + IMG_W)

    num_w = Inches(text_width_in(f["number"], NUMBER_PT) + 0.1)
    num_h = Pt(NUMBER_PT * LINE_SPACING)
    tf = add_textbox(slide, "Section Number", GROUP_LEFT, mid - num_h / 2, num_w, num_h,
                     anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    add_run(tf.paragraphs[0], f["number"], NUMBER_PT, True, INK)

    name_h, eng_h = Pt(NAME_PT * LINE_SPACING), Pt(ENGLISH_PT * LINE_SPACING)
    block_h = name_h + NAME_GAP + eng_h
    rule_x = GROUP_LEFT + num_w + GAP
    add_line(slide, "Section Rule", rule_x, mid - num_h / 2, rule_x, mid + num_h / 2, RULE, 1.5)

    text_x = rule_x + GAP
    text_w = right - text_x
    top = mid - block_h / 2
    tf = add_textbox(slide, "Section Name", text_x, top, text_w, name_h,
                     anchor=MSO_ANCHOR.BOTTOM)
    add_run(tf.paragraphs[0], f["name"], NAME_PT, True, INK)
    tf = add_textbox(slide, "Section English", text_x, top + name_h + NAME_GAP, text_w, eng_h)
    add_run(tf.paragraphs[0], f["english"], ENGLISH_PT, False, MUTED)

    add_page_chrome(slide, page, ctx)
    return slide

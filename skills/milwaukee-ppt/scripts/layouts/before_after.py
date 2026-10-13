"""before-after 版式：模板 layout 3「Before / After」，槽位约定见 references/before-after.md。

文稿写法：
    第一个 ### 之前的普通文字      → lead（正文顶部一行粗体）
    ### Before[: 自定义标题]        → 改前小节；标题默认 BEFORE（idx 11）
        一张图 ![](...) 或一段短文  → 内容区 idx 13
        Caption: 一句说明           → idx 15
    ### After[: 自定义标题]         → 改后小节，同上（idx 12 / 14 / 16）
    最后一个小节之后的普通文字      → 底部红条结论
小节在 Caption: 行或空行处结束；两个小节之间不能有普通文字。

字段：fields = {"lead": str, "before": Side, "after": Side}，
Side = {"title", "image", "text": [段落], "caption"}。
"""
import re

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches

from mdparse import DeckError

from .common import (IMG_LEFT, IMG_W, add_page_chrome, add_run, caption_lines, fill_titles,
                     fit_picture, load_image, text_lines, text_width_in, warn)

NAME = "before-after"
LAYOUT_INDEX = 3

SIDES = ("before", "after")
SLOTS = {"before": {"title": 11, "content": 13, "caption": 15},
         "after": {"title": 12, "content": 14, "caption": 16}}

# 正文 lead：模板样例页「Lead」文本框的位置与样式（英寸）。写了 Tag 时右端让到月份标识左侧。
LEAD_LEFT, LEAD_TOP, LEAD_W, LEAD_H = 0.56, 1.30, 11.90, 0.42
TAG_LEFT = IMG_LEFT + IMG_W - 1.7          # common.add_month_tag 的文本框左缘
LEAD_PT = 20
INK = "1A1A1A"
CAPTION_PT = 16                            # 模板 caption 槽：16pt bold、不换行
CAPTION_INSET = 0.06 * 2                   # 槽的左右内边距（英寸）
CONTENT_PT = 18                            # 模板内容槽：18pt、居中
LINE_SPACING = 1.2

# 有底部红条时（红条最低从 5.88" 起、最高可长到 5.60"），模板 6.12" 的 caption 槽会被盖住：
# caption 收进对比框底部（与内容区同宽），内容区相应变矮。框本身是 layout 形状，位置不变。
BAR_SAFE_BOTTOM = 5.55
CAPTION_IN_FRAME_H = 0.35
CONTENT_CAPTION_GAP = 0.10

_IMG = re.compile(r"^!\[[^\]]*\]\(([^)]+)\)$")
_SECTION = re.compile(r"^###\s+(\S+?)\s*(?:[:：]\s*(.*))?$")
_CAPTION = re.compile(r"(?i)^caption\s*[:：]\s*(.*)$")


def parse(page):
    def err(msg):
        raise DeckError(f"页「{page.title}」{msg}")

    lead, pending, sides = [], [], {}
    cur, open_ = None, False          # 当前小节；它是否还在收内容
    for line in page.lines:
        sec = _SECTION.match(line) if line.startswith("###") else None
        if sec:
            key = sec.group(1).lower()
            if key not in SIDES:
                err(f"的小节只能是 ### Before / ### After, 实际: {line!r}")
            if key in sides:
                err(f"有两个 ### {sec.group(1)} 小节")
            if pending:
                err(f"的两个小节之间有普通文字: {pending[0]!r}")
            cur = sides[key] = {"title": (sec.group(2) or "").strip() or key.upper(),
                                "image": None, "text": [], "caption": ""}
            open_ = True
            continue
        if cur is None:               # 第一个小节之前：lead
            if line:
                lead.append(line)
            continue
        if not open_:                 # 小节已结束：结论候选
            if line:
                pending.append(line)
            continue
        if not line:
            open_ = bool(not cur["image"] and not cur["text"])   # 内容前的空行不结束小节
            continue
        cap = _CAPTION.match(line)
        if cap:
            cur["caption"] = cap.group(1).strip()
            open_ = False
            continue
        img = _IMG.match(line)
        if img:
            if cur["image"] or cur["text"]:
                err("的每个小节只能放一张图或一段短文")
            cur["image"] = img.group(1).strip()
        else:
            if cur["image"]:
                err("的每个小节只能放一张图或一段短文")
            cur["text"].append(line)

    for key in SIDES:
        if key not in sides:
            err(f"缺少 ### {key.capitalize()} 小节")
        if not sides[key]["image"] and not sides[key]["text"]:
            err(f"的 ### {key.capitalize()} 小节没有内容（一张图或一段短文）")
    page.conclusion.extend(pending)
    page.fields = {"lead": " ".join(lead), **sides}


def _check_caption(page, side, slot):
    avail = Emu(slot.width).inches - CAPTION_INSET
    if side["caption"] and text_width_in(side["caption"], CAPTION_PT) > avail:
        warn(f"页「{page.title}」{side['title']} 的 Caption 估计超过一行, "
             f"建议精简到约 {int(avail * 72 / CAPTION_PT)} 个汉字以内: {side['caption']}")


def _check_text(page, side, slot):
    avail_w = Emu(slot.width).inches - CAPTION_INSET
    lines = sum(text_lines(t, CONTENT_PT, avail_w) for t in side["text"])
    if lines * CONTENT_PT * LINE_SPACING / 72 > Emu(slot.height).inches:
        warn(f"页「{page.title}」{side['title']} 的短文估计溢出内容区 (约 {int(lines)} 行), "
             f"建议精简")


def _add_lead(slide, page):
    text = page.fields["lead"]
    if not text:
        return None
    right = min(LEAD_LEFT + LEAD_W, TAG_LEFT - 0.1) if page.tag else LEAD_LEFT + LEAD_W
    tb = slide.shapes.add_textbox(Inches(LEAD_LEFT), Inches(LEAD_TOP),
                                  Inches(right - LEAD_LEFT), Inches(LEAD_H))
    tb.name = "Lead"
    if text_width_in(text, LEAD_PT) > right - LEAD_LEFT - CAPTION_INSET:
        warn(f"页「{page.title}」的 lead 估计超过一行, 建议精简: {text}")
    tf = tb.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    add_run(p, text, LEAD_PT, True, INK)
    return tb


def _fit_above_bar(content, caption):
    """把 caption 槽挪进对比框底部，内容区缩到 caption 上方（都在红条可能的最高位置之上）。"""
    cap_top = Inches(BAR_SAFE_BOTTOM - CAPTION_IN_FRAME_H)
    left, top, width = content.left, content.top, content.width
    caption.left, caption.top = left, cap_top
    caption.width, caption.height = width, Inches(CAPTION_IN_FRAME_H)
    content.left, content.top, content.width = left, top, width
    content.height = cap_top - Inches(CONTENT_CAPTION_GAP) - top


def render(prs, page, ctx):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    fill_titles(slide, page)
    has_bar = bool(caption_lines(page))
    for key in SIDES:
        side, slots = page.fields[key], SLOTS[key]
        slide.placeholders[slots["title"]].text = side["title"]
        content = slide.placeholders[slots["content"]]
        caption = slide.placeholders[slots["caption"]]
        if has_bar:
            _fit_above_bar(content, caption)
        content.text = "\n".join(side["text"])     # 放图时为空：先清空槽内文字
        content.text_frame.word_wrap = True            # 模板槽是不换行的，短文要在槽宽内换行
        _check_text(page, side, content)
        if side["image"]:
            fit_picture(slide, load_image(side["image"], ctx, page.title),
                        content.left, content.top, content.width, content.height)
        caption.text = side["caption"]
        _check_caption(page, side, caption)
    _add_lead(slide, page)
    add_page_chrome(slide, page, ctx)
    return slide

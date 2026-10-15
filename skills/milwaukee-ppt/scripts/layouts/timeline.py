"""timeline 版式：模板 layout 6「Timeline」，槽位 idx 19「Timeline Area」。

字段：3–6 条无序列表，每条「[!]日期 | 标题 | 说明 [| 状态]」→
fields["items"] = [(日期, 标题, 说明)]、fields["keys"]（行首 ! = 关键里程碑）、
fields["statuses"]（common.parse_statuses：整页都没写为 None）；其余普通文字行进红条结论。

画法：槽位按条目数等分成列（列间距 0.3"），每列中心一个节点；日期在线上方，标题和说明
在线下方，都与节点居中对齐。全部原生形状，画完删掉空的槽位占位框。

`!` 只认一个（common.key_mark），写成 `!!` 报错。

- 不写状态（旧写法）：黑白灰，一条灰横线连起首末节点；不写状态也不写 ! 时输出与 #24 之前
  逐字节一致（tests/test_timeline.py 的 LEGACY_SHA 守着）。带 ! 时会画菱形：放大的深灰菱形。
  形状名 Timeline Line / Node N / Date N / Text N。
- 写了状态（Spec #20，见 references/layouts.md「timeline」）：已完成绿色实心圆 + 白对勾、
  首节点到进度终点是绿色粗线（Progress Line），之后灰色虚线（Remaining Line）；进行中放大的琥珀圆 + 外圈
  （Current Ring），上方近黑胶囊“当前”（Current Tag），日期加粗放大；未完成白底红描边。
  关键里程碑是放大的菱形，颜色按状态。右上角图例 4 项（common.add_legend：Legend Icon N /
  Legend Label N，N = 1 已完成、2 进行中、3 未完成、4 关键里程碑；已完成图标上的对勾 Legend Check 1）。节点相关形状名 Milestone Node N /
  Check N / Date N / Text N（N 为条目序号，对勾只在已完成节点上）。
"""
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from mdparse import DeckError

from .common import (DOING, DONE, DONE_GREEN, KEY_LABEL, LIST_ITEM, STATUS_COLORS, STATUS_LABELS,
                     TODO, add_legend, add_line, add_page_chrome, add_run, add_shape, add_textbox,
                     fill_titles, key_mark, parse_statuses, strip_theme_style, fit_body, take_slot,
                     text_lines, warn)

NAME = "timeline"
LAYOUT_INDEX = 6
SLOT_IDX = 19
MIN_ITEMS, MAX_ITEMS = 3, 6

GAP = 0.3             # 列间距（统一块间距）
LINE_Y = 1.6          # 横线距槽位上沿（内容大致落在槽位中部，不全挤在上半）
NODE_D = 0.24         # 节点直径
KEY_SCALE = 1.5       # 关键里程碑：菱形放大

DATE_H = 0.35         # 日期框高（12pt 一行）
TEXT_GAP = 0.2        # 日期 / 正文与节点之间的空隙

LINE_COLOR = "BFBFBF"  # 旧写法的横线、状态版的“之后”虚线
NODE_COLOR = "333333"  # 旧写法的节点
WHITE = "FFFFFF"
INK = "1A1A1A"        # 标题
MUTED = "6B6B6B"      # 日期、说明
LINE_SPACING = 1.2    # 行高 ≈ 字号 × 1.2（估算溢出用）
DESC_SPACE_BEFORE = 6  # 说明段前（pt）

# (字号 pt, 粗体, 颜色)：#13 字号阶梯
DATE_STYLE = (12, True, MUTED)
TITLE_STYLE = (18, True, INK)
DESC_STYLE = (14, False, MUTED)



def parse(page):
    items, keys, raws = [], [], []
    for line in page.lines:
        if not line:
            continue
        m = LIST_ITEM.match(line)
        if not m:
            page.conclusion.append(line)
            continue
        key, body = key_mark(m.group(1), page)
        parts = [x.strip() for x in body.split("|")]
        if len(parts) not in (3, 4) or not parts[0] or not parts[1]:
            raise DeckError(f"页「{page.title}」的时间线条目要写成「日期 | 标题 | 说明」"
                            f"（可加第 4 段状态）: {line}")
        items.append(tuple(parts[:3]))
        keys.append(key)
        raws.append(parts[3] if len(parts) == 4 else "")
    if not MIN_ITEMS <= len(items) <= MAX_ITEMS:
        raise DeckError(f"页「{page.title}」的时间线条目有 {len(items)} 条, "
                        f"应为 {MIN_ITEMS}–{MAX_ITEMS} 条")
    page.fields = {"items": items, "keys": keys, "statuses": parse_statuses(page, raws)}


def _overflow_warnings(page, items, col_w, date_h, body_h):
    """日期超过一行、标题 + 说明估计超出线下区域时警告（不中止）。"""
    width_in = Emu(col_w).inches

    def height_pt(text, style):
        return text_lines(text, style[0], width_in) * style[0] * LINE_SPACING

    for i, (date, title, desc) in enumerate(items, start=1):
        date_pt = height_pt(date, DATE_STYLE)
        body_pt = height_pt(title, TITLE_STYLE)
        if desc:
            body_pt += DESC_SPACE_BEFORE + height_pt(desc, DESC_STYLE)
        if date_pt > Emu(date_h).pt or body_pt > Emu(body_h).pt:
            warn(f"页「{page.title}」时间线第 {i} 条文字估计溢出, 建议精简日期 / 标题 / 说明")


def _plain(shape):
    """去掉主题样式引用、关掉继承阴影：填充和描边只用显式写的颜色。"""
    strip_theme_style(shape)
    shape.shadow.inherit = False
    return shape


def _line(slide, name, x1, x2, y, color, pt, dash=None):
    """水平线；时间线的线也关掉继承阴影（与 #24 之前的输出一致）。"""
    ln = add_line(slide, name, x1, y, x2, y, color, pt, dash)
    ln.shadow.inherit = False
    return ln


def render(prs, page, ctx):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    fill_titles(slide, page)
    # 空槽位不留在页面上；方向 C 下上下沿换成正文区
    rect = fit_body(take_slot(slide, SLOT_IDX), page, ctx)
    if page.fields["statuses"] is None:
        _render_plain(slide, page, rect)
    else:
        _render_status(slide, page, rect)
    add_page_chrome(slide, page, ctx)
    return slide


def _render_plain(slide, page, rect):
    """旧写法（不写状态）：黑白灰，输出与 #24 之前逐字节一致。"""
    left, top, width, slot_h = rect
    items = page.fields["items"]
    n = len(items)
    gap = Inches(GAP)
    col_w = (width - gap * (n - 1)) // n
    centers = [left + i * (col_w + gap) + col_w // 2 for i in range(n)]
    line_y = top + Inches(LINE_Y)

    _line(slide, "Timeline Line", centers[0], centers[-1], line_y, LINE_COLOR, 1.5)

    for i, (cx, key) in enumerate(zip(centers, page.fields["keys"]), start=1):
        d = Inches(NODE_D * KEY_SCALE) if key else Inches(NODE_D)   # 关键里程碑：放大的菱形
        # 2pt 白描边，让节点和横线之间留一圈空隙
        _shape(slide, MSO_SHAPE.DIAMOND if key else MSO_SHAPE.OVAL, f"Timeline Node {i}",
               cx, line_y, d, NODE_COLOR, WHITE, 2)

    d = Inches(NODE_D)
    date_h = Inches(DATE_H)
    date_top = line_y - d // 2 - Inches(TEXT_GAP) - date_h
    body_top = line_y + d // 2 + Inches(TEXT_GAP)
    body_h = top + slot_h - body_top
    _overflow_warnings(page, items, col_w, date_h, body_h)
    for i, (cx, (date, title, desc)) in enumerate(zip(centers, items), start=1):
        x = cx - col_w // 2
        tf = add_textbox(slide, f"Timeline Date {i}", x, date_top, col_w, date_h,
                      MSO_ANCHOR.BOTTOM)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        add_run(tf.paragraphs[0], date, *DATE_STYLE)

        tf = add_textbox(slide, f"Timeline Text {i}", x, body_top, col_w, body_h,
                      MSO_ANCHOR.TOP)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_run(p, title, *TITLE_STYLE)
        if desc:
            p = tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            p.space_before = Pt(DESC_SPACE_BEFORE)
            add_run(p, desc, *DESC_STYLE)


# ---- 状态着色（Spec #20）-----------------------------------------------------------

S_LINE_Y = 2.25       # 横线距槽位上沿（上方留给图例、“当前”标签和日期）
S_NODE_D = 0.32       # 普通节点直径
CUR_SCALE = 1.25      # 进行中：节点放大
RING_SCALE = 1.55     # 进行中外圈直径 / 放大前直径
S_DATE_GAP = 0.65     # 日期框底边距横线
TAG_GAP = 0.23        # “当前”标签中心比日期框顶再高出的距离
TAG_W, TAG_H = 0.7, 0.26
S_DATE_STYLE = DATE_STYLE              # 普通日期 12pt 粗体灰
CUR_DATE_STYLE = (14, True, INK)       # 进行中日期 14pt 粗体近黑
TAG_STYLE = (12, True, WHITE)          # “当前”：近黑胶囊白字
PROGRESS_PT = 4       # 已完成段：绿色实线
REMAIN_PT = 3         # 之后：灰色虚线
# 对勾折线三点（相对节点中心，单位 = 节点边长）；关键里程碑菱形里缩到 0.75
CHECK_PTS = ((-0.26, 0.0), (-0.07, 0.2), (0.28, -0.2))
CHECK_PT = 2.25
S_BODY_GAP = 0.6      # 正文框顶距横线（让开最大的节点 / 外圈）
S_TITLE_STYLE = (16, True, INK)        # 普通标题；关键里程碑 / 进行中用 TITLE_STYLE 18pt
TODO_TITLE_COLOR = MUTED               # 未完成的标题灰色

# 图例（common.add_legend）：槽位右上角一行，右对齐；图标边长 0.2"（关键里程碑菱形 0.26"）
LEGEND_ICON, LEGEND_KEY_ICON = 0.2, 0.26
LEGEND_SWATCH = 0.22  # 图标占位宽
LEGEND_TAG_DROP = 0.25  # 页面写了 Tag 时右上角有月份标识（1.12–1.42"），图例下移让开


def _shape(slide, kind, name, cx, cy, d, fill, line, line_pt):
    """以 (cx, cy) 为中心、边长 d 的自选图形；颜色全部显式给，无主题样式、无阴影。"""
    return add_shape(slide, kind, name, cx - d // 2, cy - d // 2, d, d, fill, line, line_pt)


def _node(slide, name, cx, cy, status, key, d):
    """一个状态节点：已完成 / 进行中实心 + 白描边；未完成白底红描边（不实心填红）。"""
    kind = MSO_SHAPE.DIAMOND if key else MSO_SHAPE.OVAL
    color = STATUS_COLORS[status]
    if status == TODO:
        return _shape(slide, kind, name, cx, cy, d, WHITE, color, 2.5)
    return _shape(slide, kind, name, cx, cy, d, color, WHITE, 1.5)


def _render_status(slide, page, rect):
    left, top, width, slot_h = rect
    items, keys, sts = (page.fields[k] for k in ("items", "keys", "statuses"))
    n = len(items)
    gap = Inches(GAP)
    col_w = (width - gap * (n - 1)) // n
    centers = [left + i * (col_w + gap) + col_w // 2 for i in range(n)]
    line_y = top + Inches(S_LINE_Y)

    _legend(slide, left + width, top + (Inches(LEGEND_TAG_DROP) if page.tag else 0))

    # 进度：到“进行中”为止；没有进行中时到最后一个已完成
    reached = [i for i, st in enumerate(sts) if st in (DONE, DOING)]
    end = sts.index(DOING) if DOING in sts else (max(reached) if reached else 0)
    if end > 0:
        _line(slide, "Progress Line", centers[0], centers[end], line_y, DONE_GREEN,
              PROGRESS_PT)
    if end < n - 1:
        _line(slide, "Remaining Line", centers[end], centers[-1], line_y, LINE_COLOR,
              REMAIN_PT, MSO_LINE_DASH_STYLE.DASH)

    date_h = Inches(DATE_H)
    date_top = line_y - Inches(S_DATE_GAP) - date_h
    for i, (cx, st, key) in enumerate(zip(centers, sts, keys), start=1):
        d = S_NODE_D * (KEY_SCALE if key else 1)
        if st == DOING:
            _shape(slide, MSO_SHAPE.OVAL, "Current Ring", cx, line_y, Inches(d * RING_SCALE),
                   WHITE, STATUS_COLORS[DOING], 2)
            d *= CUR_SCALE
        _node(slide, f"Milestone Node {i}", cx, line_y, st, key, Inches(d))
        if st == DONE:
            _check(slide, f"Milestone Check {i}", cx, line_y, Inches(d * (0.75 if key else 1)))

    for i, (cx, st, (date, _, _)) in enumerate(zip(centers, sts, items), start=1):
        x = cx - col_w // 2
        tf = add_textbox(slide, f"Milestone Date {i}", x, date_top, col_w, date_h,
                      MSO_ANCHOR.BOTTOM)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        add_run(tf.paragraphs[0], date, *(CUR_DATE_STYLE if st == DOING else S_DATE_STYLE))
        if st == DOING:
            _current_tag(slide, cx, date_top - Inches(TAG_GAP))

    body_top = line_y + Inches(S_BODY_GAP)
    body_h = top + slot_h - body_top
    _overflow_warnings(page, items, col_w, date_h, body_h)
    for i, (cx, st, key, (_, title, desc)) in enumerate(zip(centers, sts, keys, items),
                                                       start=1):
        tf = add_textbox(slide, f"Milestone Text {i}", cx - col_w // 2, body_top, col_w, body_h,
                      MSO_ANCHOR.TOP)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        pt, bold, color = TITLE_STYLE if (key or st == DOING) else S_TITLE_STYLE
        add_run(p, title, pt, bold, TODO_TITLE_COLOR if st == TODO else color)
        if desc:
            p = tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            p.space_before = Pt(DESC_SPACE_BEFORE)
            add_run(p, desc, *DESC_STYLE)


def _legend(slide, right, top):
    """槽位右上角的状态图例：已完成（带对勾）/ 进行中 / 未完成 / 关键里程碑。"""
    def status_icon(st):
        def draw(name, cx, cy):
            _node(slide, name, cx, cy, st, False, Inches(LEGEND_ICON))
            if st == DONE:
                _check(slide, name.replace("Icon", "Check"), cx, cy, Inches(LEGEND_ICON))
        return draw

    def key_icon(name, cx, cy):
        _shape(slide, MSO_SHAPE.DIAMOND, name, cx, cy, Inches(LEGEND_KEY_ICON), INK, None, 0)

    entries = [(LEGEND_SWATCH, status_icon(st), STATUS_LABELS[st]) for st in (DONE, DOING, TODO)]
    add_legend(slide, right, top, entries + [(LEGEND_SWATCH, key_icon, KEY_LABEL)])


def _check(slide, name, cx, cy, size):
    """节点中心的白色对勾（两段折线，无填充）。"""
    pts = [(cx + int(dx * size), cy + int(dy * size)) for dx, dy in CHECK_PTS]
    fb = slide.shapes.build_freeform(*pts[0], scale=1.0)
    fb.add_line_segments(pts[1:], close=False)
    sp = fb.convert_to_shape()
    sp.name = name
    sp.fill.background()
    sp.line.color.rgb = RGBColor.from_string(WHITE)
    sp.line.width = Pt(CHECK_PT)
    _plain(sp)
    return sp


def _current_tag(slide, cx, cy):
    """进行中节点上方的“当前”胶囊标签（近黑底白字）。"""
    w, h = Inches(TAG_W), Inches(TAG_H)
    sp = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, "Current Tag", cx - w // 2, cy - h // 2,
                   w, h, INK)
    sp.adjustments[0] = 0.5
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    add_run(tf.paragraphs[0], "当前", *TAG_STYLE)

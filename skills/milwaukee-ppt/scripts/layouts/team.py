"""team 版式（团队页）：模板 layout 0「Blank Slide」，列出成员、角色和分工。本期不带头像。

文稿写法：2–8 行无序列表「- 姓名 | 角色 | 职责」，一人一行，三段都必填，如
`- 周航 | 现场推进 | 试点车间贴码、培训和巡检`；其余普通文字行进底部红条（分工结论），
Next: 照常可用。没写列表符号的「姓名 | 角色 | 职责」行报错（免得悄悄进红条）；
团队页不标重点（`!`）、不写子列表。

排版：每人一张卡片（#F5F5F5 底 + #E5E5E5 细描边），卡片宽固定为一行 4 张时的宽度：
2–4 人排一行，5–8 人排两行（第一行 = 总数一半向上取整，第二行从同一左沿排起）；
整组在正文区（common.body_area；月报模式 1.30–5.60"）里水平、垂直居中。卡片里自上而下：姓名（20pt bold #1A1A1A）、
角色（14pt #6B6B6B）、一条 #BFBFBF 细线、职责（12pt #1A1A1A，最多 3 行）。
只用黑白灰，banner 和底部红条以外不用红。和 evidence、section、agenda 共用模板 layout 0，
inspect 靠 MARK_SHAPE（`Member Name 1`，两种模式都有）区分。以上是月报模式（逐字节不变）。

通用 / 项目汇报模式（方向 C，#50）：不画卡片。每人一列：黑底白字的姓名首字圆形头像
（Member Avatar N），下面依次居中放姓名（Member Name N，粗体 #1A1A1A）、角色（Member Role N，
#6B6B6B）、一条灰色短线（Member Rule N）、职责（Member Duty N，#333333）。2–4 人排一行，
5–8 人排两行（第一行 = 一半向上取整），每行各自水平居中；列宽 = 正文区宽 / 每行人数
（不超过 3 人一行时的列宽）。头像和字号按 AVATAR_TIERS 从大到小试，取第一档
「姓名、角色一行放得下，整组高度放得进正文区」的；最小一档也放不下才警告。
整组在正文区（common.body_area）里垂直居中。
"""
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from mdparse import DeckError

from .common import (DESC_INK, IMG_W, INK, LINE_SPACING, LIST_ITEM, MUTED, RULE, add_line,
                     add_page_chrome, add_run, add_shape, add_textbox, body_area, fill_titles,
                     text_lines, text_width_in, warn)

NAME = "team"
LAYOUT_INDEX = 0
MARK_SHAPE = "Member Name 1"   # 两种模式都有：inspect_ppt 凭这个形状认出本版式

MIN_MEMBERS, MAX_MEMBERS = 2, 8
ROW_MAX = 4                     # 一行最多 4 张卡片

NAME_PT, ROLE_PT, DUTY_PT = 20, 14, 12
MAX_DUTY_LINES = 3
CARD_FILL, CARD_LINE = "F5F5F5", "E5E5E5"

# 几何（英寸）
GAP = 0.3                       # 卡片之间（横、竖）
CARD_W = (IMG_W - GAP * (ROW_MAX - 1)) / ROW_MAX
PAD = 0.2                       # 卡片内边距
TEXT_W = CARD_W - 2 * PAD
NAME_H = NAME_PT * LINE_SPACING / 72
ROLE_H = ROLE_PT * LINE_SPACING / 72
DUTY_H = MAX_DUTY_LINES * DUTY_PT * LINE_SPACING / 72
ROLE_GAP = 0.04                 # 姓名 → 角色
RULE_GAP = 0.12                 # 角色 → 细线 → 职责
CARD_H = PAD + NAME_H + ROLE_GAP + ROLE_H + 2 * RULE_GAP + DUTY_H + PAD   # ≈ 1.90"

# 方向 C（英寸 / pt）：头像和字号分档，从大到小试
# (头像直径, 首字, 姓名, 角色, 职责, 头像 → 姓名的间距；角色 → 短线 → 职责各半个间距，两行之间 1.5 个)
AVATAR_TIERS = ((1.4, 44, 24, 16, 18, 0.25),
                (1.1, 36, 22, 16, 16, 0.2),
                (0.85, 30, 20, 14, 14, 0.15),
                (0.5, 20, 16, 12, 12, 0.1))
C_ROLE_GAP = 0.04               # 姓名 → 角色
C_TEXT_PAD = 0.2                # 文字框比列宽左右各窄
C_RULE_W = 0.7                  # 短线长
C_MIN_COLS = 3                  # 2 人时列宽按 3 列算，免得两人隔得太远


def _fail(page, msg):
    raise DeckError(f"页「{page.title}」{msg}")


def parse(page):
    members = []
    for i, line in enumerate(page.lines):
        if not line:
            continue
        m = LIST_ITEM.match(line)
        if m and page.indents[i] > 0:
            _fail(page, f"团队页不写子列表，职责写在成员行的第三段: {line!r}")
        if not m:
            if "|" in line:
                _fail(page, f"团队页成员行要以列表符号开头，写成「- 姓名 | 角色 | 职责」: {line!r}")
            page.conclusion.append(line)
            continue
        body = m.group(1)
        if body.startswith("!"):
            _fail(page, f"团队页不标重点，去掉 !: {line!r}")
        parts = [p.strip() for p in body.split("|")]
        if len(parts) != 3:
            _fail(page, f"团队页成员行写成「- 姓名 | 角色 | 职责」，这一行不对: {line!r}")
        members.append(parts)
    if not MIN_MEMBERS <= len(members) <= MAX_MEMBERS:
        _fail(page, f"团队页须 {MIN_MEMBERS}–{MAX_MEMBERS} 行「- 姓名 | 角色 | 职责」，"
                    f"现在 {len(members)} 行")
    for name, role, duty in members:
        for label, value in (("姓名", name), ("角色", role), ("职责", duty)):
            if not value:
                _fail(page, f"团队页「{name} | {role} | {duty}」的{label}为空")
        for label, value, pt, hint in (("姓名", name, NAME_PT, "约 8 个汉字"),
                                       ("角色", role, ROLE_PT, "约 12 个汉字")):
            if text_width_in(value, pt) > TEXT_W:
                _fail(page, f"团队页的{label}「{value}」一行放不下，最多{hint}")
        if text_lines(duty, DUTY_PT, TEXT_W) > MAX_DUTY_LINES:
            _fail(page, f"团队页「{name}」的职责「{duty}」超过 {MAX_DUTY_LINES} 行，"
                        f"最多约 40 个汉字；写一句分工就够")
    page.fields = {"members": [{"name": n, "role": r, "duty": d} for n, r, d in members]}


def render(prs, page, ctx):
    if ctx.lead:
        return _render_avatars(prs, page, ctx)
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    fill_titles(slide, page)
    members = page.fields["members"]
    per_row = len(members) if len(members) <= ROW_MAX else (len(members) + 1) // 2
    rows = -(-len(members) // per_row)
    group_w = per_row * CARD_W + (per_row - 1) * GAP
    group_h = rows * CARD_H + (rows - 1) * GAP
    area = body_area(page, ctx)
    left0 = area.left + (area.width - group_w) / 2
    top0 = area.top + (area.height - group_h) / 2
    for i, mb in enumerate(members):
        n = i + 1
        row, col = divmod(i, per_row)
        x = left0 + col * (CARD_W + GAP)
        y = top0 + row * (CARD_H + GAP)
        add_shape(slide, MSO_SHAPE.RECTANGLE, f"Member Card {n}", Inches(x), Inches(y),
                  Inches(CARD_W), Inches(CARD_H), CARD_FILL, CARD_LINE, 0.75)
        tx, tw = Inches(x + PAD), Inches(TEXT_W)
        top = y + PAD
        tf = add_textbox(slide, f"Member Name {n}", tx, Inches(top), tw, Inches(NAME_H),
                         anchor=MSO_ANCHOR.BOTTOM, wrap=False)
        add_run(tf.paragraphs[0], mb["name"], NAME_PT, True, INK)
        top += NAME_H + ROLE_GAP
        tf = add_textbox(slide, f"Member Role {n}", tx, Inches(top), tw, Inches(ROLE_H),
                         wrap=False)
        add_run(tf.paragraphs[0], mb["role"], ROLE_PT, False, MUTED)
        top += ROLE_H + RULE_GAP
        add_line(slide, f"Member Rule {n}", tx, Inches(top), tx + tw, Inches(top), RULE, 0.75)
        top += RULE_GAP
        tf = add_textbox(slide, f"Member Duty {n}", tx, Inches(top), tw, Inches(DUTY_H))
        add_run(tf.paragraphs[0], mb["duty"], DUTY_PT, False, INK)
    add_page_chrome(slide, page, ctx)
    return slide


# ---- 方向 C：姓名首字圆形头像 ----------------------------------------------------

def _rows(n):
    """n 人分几行、每行几人：2–4 人一行，5–8 人两行（第一行一半向上取整）。"""
    if n <= ROW_MAX:
        return [n]
    first = (n + 1) // 2
    return [first, n - first]


def _h(pt, lines=1):
    return lines * pt * LINE_SPACING / 72


def _tier_layout(members, rows, col, tier):
    """按这一档排时 (各段高度, 整组高)；姓名或角色一行放不下返回 None。"""
    d, _, name_pt, role_pt, duty_pt, gap = tier
    text_w = col - 2 * C_TEXT_PAD
    if any(text_width_in(m["name"], name_pt) > text_w or text_width_in(m["role"], role_pt) > text_w
           for m in members):
        return None
    duty_lines = max(int(text_lines(m["duty"], duty_pt, text_w)) for m in members)
    row_h = (d + gap + _h(name_pt) + C_ROLE_GAP + _h(role_pt) + gap
             + _h(duty_pt, duty_lines))
    return duty_lines, row_h, len(rows) * row_h + (len(rows) - 1) * 1.5 * gap


def _pick_tier(members, rows, col, area_h):
    """从大到小取第一档放得下的；都放不下用最小一档（返回 fits=False）。"""
    for tier in AVATAR_TIERS:
        got = _tier_layout(members, rows, col, tier)
        if got and got[2] <= area_h:
            return tier, got, True
    tier = AVATAR_TIERS[-1]
    return tier, _tier_layout(members, rows, col, tier) or (MAX_DUTY_LINES, 0, 0), False


def _render_avatars(prs, page, ctx):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    fill_titles(slide, page)
    members = page.fields["members"]
    rows = _rows(len(members))
    area = body_area(page, ctx)
    col = area.width / max(rows[0], C_MIN_COLS)
    tier, (duty_lines, row_h, group_h), fits = _pick_tier(members, rows, col, area.height)
    if not fits:
        warn(f"页「{page.title}」团队成员的姓名、角色和职责估计溢出正文区, 建议精简职责")
    d, initial_pt, name_pt, role_pt, duty_pt, gap = tier
    text_w = col - 2 * C_TEXT_PAD
    top0 = area.top + max(0, (area.height - group_h) / 2)
    n = 0
    for r, count in enumerate(rows):
        top = top0 + r * (row_h + 1.5 * gap)
        left = area.left + (area.width - count * col) / 2
        for c in range(count):
            mb = members[n]
            n += 1
            cx = left + (c + 0.5) * col
            av = add_shape(slide, MSO_SHAPE.OVAL, f"Member Avatar {n}", Inches(cx - d / 2),
                           Inches(top), Inches(d), Inches(d), INK)
            tf = av.text_frame
            tf.word_wrap = False
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            add_run(tf.paragraphs[0], mb["name"][0], initial_pt, True, "FFFFFF")
            y = top + d + gap
            _centered(slide, f"Member Name {n}", mb["name"], name_pt, True, INK,
                      cx, y, text_w, _h(name_pt), wrap=False)
            y += _h(name_pt) + C_ROLE_GAP
            _centered(slide, f"Member Role {n}", mb["role"], role_pt, False, MUTED,
                      cx, y, text_w, _h(role_pt), wrap=False)
            y += _h(role_pt) + gap / 2
            rule_w = min(C_RULE_W, text_w)
            add_line(slide, f"Member Rule {n}", Inches(cx - rule_w / 2), Inches(y),
                     Inches(cx + rule_w / 2), Inches(y), RULE, 0.75)
            y += gap / 2
            _centered(slide, f"Member Duty {n}", mb["duty"], duty_pt, False, DESC_INK,
                      cx, y, text_w, _h(duty_pt, duty_lines))
    add_page_chrome(slide, page, ctx)
    return slide


def _centered(slide, name, text, pt, bold, color, cx, top, width, height, wrap=True):
    tf = add_textbox(slide, name, Inches(cx - width / 2), Inches(top), Inches(width),
                     Inches(height), wrap=wrap)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    add_run(tf.paragraphs[0], text, pt, bold, color)

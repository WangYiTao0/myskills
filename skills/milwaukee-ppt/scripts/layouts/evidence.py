"""evidence 版式（默认）：1–3 张图 + 结论（月报底部红条 / 通用、项目模式 lead），基于模板 layout 0。

字段：fields["images"] = 文稿里 ![](...) 的图片引用列表；其余普通文字行是结论（月报进底部红条，通用 / 项目模式画成标题下的 lead）。
"""
import re

from pptx.util import Inches

from .common import IMG_GAP, add_page_chrome, body_area, fill_titles, load_image, warn

NAME = "evidence"
LAYOUT_INDEX = 0          # 模板 layout 0「Blank Slide」：idx 0 主标题、idx 10 action title
MAX_IMAGES = 3

_IMG = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")


def parse(page):
    images = []
    for line in page.lines:
        if not line:
            continue
        m = _IMG.match(line)
        if m:
            images.append(m.group(1).strip())
        else:
            page.conclusion.append(line)
    page.fields = {"images": images}


def render(prs, page, ctx):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    fill_titles(slide, page)
    place_images(slide, page.fields.get("images", []), ctx, page.title, body_area(page, ctx))
    add_page_chrome(slide, page, ctx)
    return slide


def _cell_widths(files, usable):
    """多图宽度分配：宽高比差异大时按比例分（宽图多占），否则等分。

    需要 Pillow；没装就等分。下限防最窄图被挤没。"""
    n = len(files)
    widths = [usable / n] * n
    if n < 2:
        return widths
    try:
        from PIL import Image
    except ImportError:
        return widths
    aspects = []
    for p in files:
        with Image.open(p) as im:
            aspects.append(im.width / im.height)
    if max(aspects) / min(aspects) <= 1.5:
        return widths
    floor = 0.28 if n == 2 else 0.22
    shares = [max(a / sum(aspects), floor) for a in aspects]
    total = sum(shares)
    return [usable * s / total for s in shares]


def place_images(slide, refs, ctx, page_title, area):
    """1–3 张图排在正文区 area（common.Area，英寸）里：各自保持比例、在 cell 内居中；
    整行按最高图垂直居中。"""
    if not refs:
        return
    if len(refs) > MAX_IMAGES:
        warn(f"页「{page_title}」有 {len(refs)} 张图, 仅使用前 {MAX_IMAGES} 张")
        refs = refs[:MAX_IMAGES]
    files = [load_image(r, ctx, page_title) for r in refs]
    widths = _cell_widths(files, area.width - IMG_GAP * (len(files) - 1))

    placed = []
    cell_left = area.left
    for p, cell_w in zip(files, widths):
        pic = slide.shapes.add_picture(p, Inches(cell_left), Inches(area.top))
        scale = min(Inches(cell_w) / pic.width, Inches(area.height) / pic.height)
        pic.width = int(pic.width * scale)
        pic.height = int(pic.height * scale)
        pic.left = Inches(cell_left) + int((Inches(cell_w) - pic.width) / 2)
        placed.append(pic)
        cell_left += cell_w + IMG_GAP
    row_h = max(pic.height for pic in placed)
    row_top = Inches(area.top) + int((Inches(area.height) - row_h) / 2)
    for pic in placed:
        pic.top = row_top



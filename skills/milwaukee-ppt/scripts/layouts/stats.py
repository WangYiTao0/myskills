"""stats 版式：数据对比 / Key Results，基于模板 layout 5「Stats」，分栏式（Spec #20，见 references/layouts/stats.md）。

文稿写法：2–4 条无序列表「标签 | 数字 | 一句说明」；最多一条在列表符号后加 `!`，
标为重点（它那栏的顶线和数字用品牌红）。卡片行下面缩进的 `- ` 子列表是这张卡的解读，
每张最多 3 条；`口径：` / `Source:` 开头的一行是口径；其余普通文字行是结论（月报进底部红条，通用 / 项目模式画成标题下的 lead）。

    - Lead time | 5 d | Mock average per request
      - Mock reading of this number
    - !Rework | -40% | Sample reduction after the change
    口径：Mock source

字段：fields["items"] = [{"label", "value", "note", "key", "points"}...]，fields["source"]
渲染：槽位 idx 18 等分成 n 栏，每栏自上而下：顶部粗线、标签、大数字（单位缩小）、说明；
有解读时再画细分隔线和带项目符号的解读；口径通栏贴在内容区底部。整页都没有解读时
（旧写法）上半截在槽位里垂直居中。画完删掉空的槽位占位框。

形状名（细调用）：Stat Rule N、Stat Label N、Stat Number N、Stat Note N、Stat Divider N、
Stat Points N（N 从 1 起，左到右），Stat Source。
"""
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from mdparse import DeckError

from types import SimpleNamespace

from .common import (BRAND_RED, DESC_INK, LATIN_FONT, UNIT_PT, add_line, add_number, add_page_chrome,
                     SOURCE_LINE, add_run, add_shape, add_textbox, check_stat_cards,
                     fill_titles, parse_stat_card, split_number, sub_item, fit_body, take_slot,
                     text_lines, warn)

NAME = "stats"
LAYOUT_INDEX = 5
SLOT_IDX = 18
MIN_ITEMS, MAX_ITEMS = 2, 4
MAX_POINTS = 3            # 每栏解读最多几条

# 几何（英寸，Spec #20 定稿，见 references/layouts/stats.md）
GAP = 0.45                # 栏间距
RULE_H = 0.06             # 顶部粗线
RULE_GAP = 0.19           # 顶线 → 标签
LABEL_MIN_H = 0.3         # 标签框最小高（一行 16pt）
NUMBER_GAP = 0.07         # 大数字 → 说明
NOTE_MIN_H = 0.3          # 说明框最小高（一行 14pt）
DIVIDER_GAP = 0.2         # 说明 → 分隔线
POINTS_GAP = 0.18         # 分隔线 → 解读
POINT_SPACE_PT = 10       # 解读条与条之间（段前）
LINE = 1.2                # 行高 = 字号 × 1.2

LABEL_PT, NUMBER_PT, NOTE_PT = 16, 60, 14
NUMBER_MIN_PT = 28        # 数字长时统一缩小，最低到这里（说明字号的 2 倍）
POINT_PT = 14
SOURCE_PT = 12
SOURCE_H = 0.28           # 口径一行高，底边贴内容区底
SOURCE_GAP = 0.1          # 解读区 → 口径
INK, GRAY, RULE = "1A1A1A", "6B6B6B", "E5E5E5"

# 两种模式的字号与颜色：月报沿用上面的定稿值（输出逐字节不变）；方向 C（通用 / 项目模式，
# #51）数字 88pt、标签 20pt、说明和解读 16pt 近黑、顶线 0.08"。
MONTHLY = SimpleNamespace(number_pt=NUMBER_PT, label_pt=LABEL_PT, label_min_h=LABEL_MIN_H,
                          note_pt=NOTE_PT, note_min_h=NOTE_MIN_H, note_color=GRAY,
                          point_pt=POINT_PT, point_color=INK, rule_h=RULE_H)
GENERAL = SimpleNamespace(number_pt=88, label_pt=20, label_min_h=0.36,
                          note_pt=16, note_min_h=0.34, note_color=DESC_INK,
                          point_pt=16, point_color=DESC_INK, rule_h=0.08)



def parse(page):
    items, sources = [], []
    for i, line in enumerate(page.lines):
        if not line:
            continue
        sub = sub_item(page, i)
        if sub is not None:                 # 缩进子列表 = 上一张卡的解读
            if not items:
                raise DeckError(f"页「{page.title}」的解读要写在某条数据下面（缩进的 - ）, "
                                f"前面还没有数据条目: {line}")
            if sub:
                items[-1]["points"].append(sub)
            continue
        if SOURCE_LINE.match(line):
            sources.append(line)
            continue
        card = parse_stat_card(page, line)
        if card is None:
            page.conclusion.append(line)
            continue
        items.append({**card, "points": []})
    check_stat_cards(page, items, MIN_ITEMS, MAX_ITEMS)
    for n, it in enumerate(items, start=1):
        if len(it["points"]) > MAX_POINTS:
            raise DeckError(f"页「{page.title}」第 {n} 条数据「{it['label']}」的解读最多 "
                            f"{MAX_POINTS} 条, 现在 {len(it['points'])} 条")
    if len(sources) > 1:
        raise DeckError(f"页「{page.title}」口径只能写一行, 现在 {len(sources)} 行")
    page.fields = {"items": items, "source": sources[0] if sources else ""}


def _glyph_em(ch):
    """逐字宽（Arial，单位 em），用来定卡片几何：大数字字号、标签 / 说明行数 → 卡片高度。

    这些估算直接决定形状尺寸，按字形细分比通用口径（common.text_width_in，
    拉丁一律 0.55em）准，所以本版式保留这张表；别处的溢出警告一律用 common 的口径。"""
    if ord(ch) > 0x2E7F:
        return 1.0
    if ch in " .,:;'!|iljtfrI()[]":
        return 0.3
    if ch in "mwMW%@":
        return 0.9
    if ch.isupper():
        return 0.72
    return 0.58


def _glyph_width_in(text, pt):
    return sum(_glyph_em(ch) for ch in text) * pt / 72


def _glyph_lines(text, pt, width_in):
    return max(1, -(-_glyph_width_in(text, pt) // width_in))


def _number_size(items, col_w_in, page_title, target=NUMBER_PT):
    """所有数字共用一个字号：target（月报 60pt、方向 C 88pt）放不下最长的那个时统一缩小
    （整数 pt）；缩到 60pt 以下才警告（方向 C 从 88 缩到 60 之间不警告）。
    宽度按 common.split_number 拆开量：大数字随字号变，前缀和单位固定 UNIT_PT。"""
    def fit_pt(value):
        prefix, number, unit = split_number(value)
        affix_w = _glyph_width_in(prefix + unit, UNIT_PT)
        return int((col_w_in - affix_w) / _glyph_width_in(number, 1))

    tightest = min(items, key=lambda it: fit_pt(it["value"]))
    fit, longest = fit_pt(tightest["value"]), tightest["value"]
    if fit >= target:
        return target
    size = max(fit, NUMBER_MIN_PT)
    if size >= NUMBER_PT:
        return size
    if fit < NUMBER_MIN_PT:
        warn(f"页「{page_title}」数字「{longest}」太长, 大数字已统一缩小到 {size}pt 仍可能放不下, "
             f"建议缩短数字或减少条数")
    else:
        warn(f"页「{page_title}」数字「{longest}」在 {NUMBER_PT}pt 放不下, "
             f"大数字已统一缩小到 {size}pt")
    return size


def _text_box(slide, name, x, y, w, h, wrap=True, anchor=MSO_ANCHOR.TOP):
    tf = add_textbox(slide, name, x, y, w, h, anchor, wrap)
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    return tf


def _text(slide, name, x, y, w, h, text, pt, bold, color):
    tf = _text_box(slide, name, x, y, w, h)
    add_run(tf.paragraphs[0], text, pt, bold, color)
    return tf


def _bullets(tf, texts, pt, color):
    """一条一段，带 • 项目符号和悬挂缩进（换行后和正文对齐；缩进 = 字号 × 1.4）。"""
    indent = int(Pt(pt * 1.4))
    for j, text in enumerate(texts):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        if j:
            p.space_before = Pt(POINT_SPACE_PT)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(indent))
        pPr.set("indent", str(-indent))
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": LATIN_FONT}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        add_run(p, text, pt, False, color)


def _points_height_in(points, width_in, pt):
    """解读区估计高度：各条按 common.text_lines 的行数 × 行高，加条间段前距。"""
    text_w = width_in - pt * 1.4 / 72
    lines = sum(text_lines(p, pt, text_w) for p in points)
    return _height_in(lines, pt) + (len(points) - 1) * POINT_SPACE_PT / 72


def _height_in(lines, pt):
    return lines * pt * LINE / 72


def render(prs, page, ctx):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    fill_titles(slide, page)
    # 方向 C 下上下沿换成正文区（common.fit_body），字号换成 GENERAL 一档
    left, top, width, height = fit_body(take_slot(slide, SLOT_IDX), page, ctx)
    st = GENERAL if ctx.lead else MONTHLY

    items = page.fields["items"]
    n = len(items)
    gap = Inches(GAP)
    col_w = (width - gap * (n - 1)) // n
    col_w_in = Emu(col_w).inches

    number_pt = _number_size(items, col_w_in, page.title, st.number_pt)
    label_h = max(st.label_min_h, _height_in(
        max(_glyph_lines(it["label"], st.label_pt, col_w_in) for it in items), st.label_pt))
    note_h = max(st.note_min_h, _height_in(
        max(_glyph_lines(it["note"], st.note_pt, col_w_in) for it in items), st.note_pt))
    has_points = any(it["points"] for it in items)
    source = page.fields["source"]
    body_h = Emu(height).inches - (SOURCE_H + SOURCE_GAP if source else 0)   # 口径以上

    def need_points(pt):
        """数字 pt 时解读区还差多少高（英寸，<=0 放得下）。"""
        points_y = (st.rule_h + RULE_GAP + label_h + _height_in(1, pt) + NUMBER_GAP + note_h
                    + DIVIDER_GAP + POINTS_GAP)
        return max((_points_height_in(it["points"], col_w_in, st.point_pt)
                    for it in items if it["points"]), default=0) - (body_h - points_y)

    # 方向 C：88pt 时解读放不下，大数字统一往下缩（不低于月报的 60pt），先保解读
    while ctx.lead and has_points and number_pt > NUMBER_PT and need_points(number_pt) > 1e-9:
        number_pt -= 1
    number_h = _height_in(1, number_pt)

    # 栏内自上而下的位置（相对内容区顶，英寸）
    label_y = st.rule_h + RULE_GAP
    number_y = label_y + label_h
    note_y = number_y + number_h + NUMBER_GAP
    divider_y = note_y + note_h + DIVIDER_GAP
    points_y = divider_y + POINTS_GAP
    y0 = top
    if not has_points:      # 旧写法只有上半截：整块在口径以上的区域里垂直居中
        y0 = top + Inches(max(0, (body_h - note_y - note_h) / 2))
    points_h = body_h - Emu(y0 - top).inches - points_y
    for it in items:
        if it["points"] and _points_height_in(it["points"], col_w_in, st.point_pt) > points_h:
            warn(f"页「{page.title}」数据「{it['label']}」的解读估计放不下, "
                 f"建议缩短或减少条数")
    if source and text_lines(source, SOURCE_PT, Emu(width).inches) > 1:
        warn(f"页「{page.title}」口径超过一行, 建议缩短")

    for i, it in enumerate(items, start=1):
        x = left + (col_w + gap) * (i - 1)
        w = (left + width - x) if i == n else col_w     # 最后一栏吃掉取整余数
        add_shape(slide, MSO_SHAPE.RECTANGLE, f"Stat Rule {i}", x, y0, w, Inches(st.rule_h),
                  BRAND_RED if it["key"] else INK)
        _text(slide, f"Stat Label {i}", x, y0 + Inches(label_y), w, Inches(label_h),
              it["label"], st.label_pt, True, INK)
        tf = _text_box(slide, f"Stat Number {i}", x, y0 + Inches(number_y), w,
                       Inches(number_h), wrap=False, anchor=MSO_ANCHOR.BOTTOM)
        add_number(tf.paragraphs[0], it["value"], number_pt,
                   BRAND_RED if it["key"] else INK)
        _text(slide, f"Stat Note {i}", x, y0 + Inches(note_y), w, Inches(note_h),
              it["note"], st.note_pt, False, st.note_color)
        if has_points:
            dy = y0 + Inches(divider_y)
            add_line(slide, f"Stat Divider {i}", x, dy, x + w, dy, RULE, 1)
        if it["points"]:
            tf = _text_box(slide, f"Stat Points {i}", x, y0 + Inches(points_y), w,
                           Inches(max(points_h, _height_in(1, st.point_pt))))
            _bullets(tf, it["points"], st.point_pt, st.point_color)
    if source:
        _text(slide, "Stat Source", left, top + height - Inches(SOURCE_H), width,
              Inches(SOURCE_H), source, SOURCE_PT, False, GRAY)

    add_page_chrome(slide, page, ctx)
    return slide

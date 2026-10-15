"""process 版式：模板 layout 4「Process」，3–5 步等宽卡片 + 序号圆，不画箭头。

文稿写法：3–5 条有序列表，每条「步骤名 | 一句说明」（说明可省略）；
其余普通文字行进底部红条结论。

排版（#13 定稿）：在槽位 idx 17 的矩形内按步数等分，卡片间距 0.3"、内边距 0.2"；
序号圆直径 0.5"、#333333 底、白字 16pt bold，圆心压在卡片上沿；
卡片 #F5F5F5 底 + #E5E5E5 细描边；步骤名 18pt bold #1A1A1A，说明 14pt #6B6B6B。
只用黑白灰。画完删掉空的槽位占位框，免得编辑视图里显示提示文字。
"""
import re

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from mdparse import DeckError

from .common import (add_page_chrome, add_run, fill_titles, strip_theme_style, fit_body, take_slot,
                     text_lines, warn)

NAME = "process"
LAYOUT_INDEX = 4
SLOT_IDX = 17
MIN_STEPS, MAX_STEPS = 3, 5

GAP = Inches(0.3)
PAD = Inches(0.2)
CIRCLE = Inches(0.5)
MIN_CARD_H = Inches(2.0)    # 内容少时卡片也不低于这个高度，免得显得单薄

NAME_PT, DESC_PT, NUM_PT = 18, 14, 16
INK, MUTED, WHITE = "1A1A1A", "6B6B6B", "FFFFFF"
CARD_FILL, CARD_LINE, CIRCLE_FILL = (
    RGBColor(0xF5, 0xF5, 0xF5), RGBColor(0xE5, 0xE5, 0xE5), RGBColor(0x33, 0x33, 0x33))
LINE_SPACING = 1.2          # 行高 ≈ 字号 × 1.2（估算溢出用）
NAME_SPACE_AFTER = Pt(6)

_STEP = re.compile(r"^\d+[.)、]\s*(.+)$")


def parse(page):
    steps = []
    for line in page.lines:
        if not line:
            continue
        m = _STEP.match(line)
        if m:
            name, _, desc = m.group(1).partition("|")
            steps.append((name.strip(), desc.strip()))
        else:
            page.conclusion.append(line)
    if not MIN_STEPS <= len(steps) <= MAX_STEPS:
        raise DeckError(f"页「{page.title}」流程步骤须 {MIN_STEPS}–{MAX_STEPS} 条"
                        f"（有序列表「1. 步骤名 | 说明」），现在 {len(steps)} 条")
    page.fields = {"steps": steps}


def render(prs, page, ctx):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    fill_titles(slide, page)
    # 方向 C 下上下沿换成正文区（common.fit_body）
    left, top, width, height = fit_body(take_slot(slide, SLOT_IDX), page, ctx)

    steps = page.fields["steps"]
    n = len(steps)
    card_w = (width - GAP * (n - 1)) // n
    max_h = height - CIRCLE // 2
    needs = [_card_height(name, desc, card_w) for name, desc in steps]
    for i, ((name, _), need) in enumerate(zip(steps, needs), start=1):
        if need > max_h:
            warn(f"页「{page.title}」第 {i} 步「{name}」的说明估计会溢出卡片, 建议精简")
    # 卡片等高：取最高的那张（不低于 MIN_CARD_H，不超槽位）；序号圆 + 卡片整组在槽位里垂直居中
    card_h = min(max_h, max(MIN_CARD_H, *needs))
    group_top = top + (height - (card_h + CIRCLE // 2)) // 2
    card_top = group_top + CIRCLE // 2
    for i, (name, desc) in enumerate(steps, start=1):
        x = left + (card_w + GAP) * (i - 1)
        _add_card(slide, i, name, desc, x, card_top, card_w, card_h)
        _add_number(slide, i, x + PAD, group_top)
    add_page_chrome(slide, page, ctx)
    return slide


def _add_card(slide, i, name, desc, x, y, w, h):
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    card.name = f"Process Card {i}"
    strip_theme_style(card)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_FILL
    card.line.color.rgb = CARD_LINE
    card.line.width = Pt(0.75)
    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_bottom = PAD
    tf.margin_top = CIRCLE // 2 + PAD
    head = tf.paragraphs[0]
    head.alignment = PP_ALIGN.LEFT        # 自选图形默认居中，正文一律左对齐
    head.space_after = NAME_SPACE_AFTER
    add_run(head, name, NAME_PT, True, INK)
    if desc:
        body = tf.add_paragraph()
        body.alignment = PP_ALIGN.LEFT
        add_run(body, desc, DESC_PT, False, MUTED)


def _add_number(slide, i, x, y):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, CIRCLE, CIRCLE)
    c.name = f"Process Number {i}"
    strip_theme_style(c)
    c.fill.solid()
    c.fill.fore_color.rgb = CIRCLE_FILL
    c.line.fill.background()
    tf = c.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    add_run(tf.paragraphs[0], str(i), NUM_PT, True, WHITE)


def _card_height(name, desc, card_w):
    """估算卡片需要的高度（EMU）：上内边距（含半个序号圆）+ 文字 + 下内边距。"""
    text_w = Emu(card_w - 2 * PAD).inches
    text_pt = (text_lines(name, NAME_PT, text_w) * NAME_PT * LINE_SPACING
               + NAME_SPACE_AFTER.pt
               + text_lines(desc, DESC_PT, text_w) * DESC_PT * LINE_SPACING)
    return CIRCLE // 2 + 2 * PAD + Pt(text_pt)

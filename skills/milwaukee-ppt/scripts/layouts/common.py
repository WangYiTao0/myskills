"""各版式共用的部件：Context、警告、图片加载、标题填充、底部红条、月份标识，
以及画原生形状的小工具（取槽位、删形状、等比放图、写带中文字体的 run、去主题样式、
估算文字宽度和行数、写“大数字 + 缩小单位”）和品牌红、字体名常量；
缩进子列表、时间线状态解析和状态色、章节行「编号 | 中文名 | 英文名」解析，
以及 section / agenda / team / table 共用的黑白灰和行高倍数。

版式模块只从这里 import，互相之间不 import。"""
import hashlib
import os
import re
import sys
import urllib.request
from dataclasses import dataclass

from xml.sax.saxutils import escape

from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from mdparse import DeckError

BRAND_RED = "DB011C"          # Milwaukee Red（与 banner 同色）
LATIN_FONT = "Arial"
EA_FONT = "PingFang SC"       # 中文字体（a:ea）

# 黑白灰（section / agenda / team / table）：正文和标题、次要文字（英文名、角色）、细线
INK, MUTED, RULE = "1A1A1A", "6B6B6B", "BFBFBF"
# 方向 C 正文说明文字（process 说明、team 职责……）：近黑，比 MUTED 深，投影上看得清
DESC_INK = "333333"
# 这几个版式排文本框 / 表格行高用的行高倍数：字号 × 1.25。审核脚本估算溢出按字号 × 1.2，
# 比这里紧 4%，所以生成器按 1.25 排出的框审核不会报溢出（见 scripts/README.md 审核脚本一节）。
LINE_SPACING = 1.25

# 章节行「编号 | 中文名 | 英文名」（section 隔断页、agenda 目录页同一写法）：各段名称和最多字数
CHAPTER_FIELDS = (("编号", 3), ("中文名", 12), ("英文名", 40))

# 正文区几何（英寸）：banner 占顶部 1.0"，底部留给红条（各版式槽位见模板 layout 4–7）
IMG_LEFT, IMG_TOP, IMG_W, IMG_H = 0.56, 1.3, 12.2, 4.3
IMG_GAP = 0.3

# 底部红条：月报模板固定样式，逐项复刻，不要改。
# 品牌红 DB011C（与 banner 同色）+ 0.75pt 黑描边，白字 ❖（Wingdings）14pt，行距 150%，垂直居中。
# 底边固定；文字多时向上加高，但不高过图片区底缘。
CAPTION_BOTTOM = 5379048 + 1097952  # EMU，底边位置固定
CAPTION_MIN_CY = 1097952            # 原版高度 = 最小高度
CAPTION_MAX_TOP = 5121840           # 不高过正文图片区底缘
CAPTION_LINE_EMU = 266700           # 每行 14pt × 150% 行距 = 21pt
CAPTION_INSETS_EMU = 45720 * 2      # tIns + bIns
CAPTION_CHARS_PER_LINE = 105        # 14pt Arial 在 12.3in 可用宽度下的保守估计

CAPTION_SP_XML = """<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <p:nvSpPr>
    <p:cNvPr id="{id}" name="Caption Bar"/>
    <p:cNvSpPr txBox="1"><a:spLocks noChangeArrowheads="1"/></p:cNvSpPr>
    <p:nvPr/>
  </p:nvSpPr>
  <p:spPr bwMode="auto">
    <a:xfrm><a:off x="155575" y="{y}"/><a:ext cx="11874874" cy="{cy}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
    <a:solidFill><a:srgbClr val="{red}"/></a:solidFill>
    <a:ln w="9525" cap="flat" cmpd="sng" algn="ctr">
      <a:solidFill><a:schemeClr val="tx1"/></a:solidFill>
      <a:prstDash val="solid"/><a:round/>
      <a:headEnd type="none" w="med" len="med"/><a:tailEnd type="none" w="med" len="med"/>
    </a:ln>
    <a:effectLst/>
  </p:spPr>
  <p:txBody>
    <a:bodyPr vert="horz" wrap="square" lIns="91440" tIns="45720" rIns="91440" bIns="45720" numCol="1" rtlCol="0" anchor="ctr" anchorCtr="0" compatLnSpc="1"><a:prstTxWarp prst="textNoShape"><a:avLst/></a:prstTxWarp></a:bodyPr>
    <a:lstStyle>
      <a:defPPr><a:defRPr lang="en-US"/></a:defPPr>
      <a:lvl1pPr marL="342900" indent="-342900" eaLnBrk="0" hangingPunct="0">
        <a:lnSpc><a:spcPct val="150000"/></a:lnSpc>
        <a:buFont typeface="Wingdings" panose="05000000000000000000" pitchFamily="2" charset="2"/>
        <a:buChar char="v"/>
        <a:defRPr sz="1400" b="0">
          <a:solidFill><a:schemeClr val="bg1"/></a:solidFill>
          <a:ea typeface="微软雅黑" panose="020B0503020204020204" pitchFamily="34" charset="-122"/>
          <a:cs typeface="Arial" panose="020B0604020202020204" pitchFamily="34" charset="0"/>
        </a:defRPr>
      </a:lvl1pPr>
    </a:lstStyle>
  </p:txBody>
</p:sp>"""

CAPTION_P_XML = ('<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                 '<a:r><a:rPr lang="en-US" altLang="zh-CN" dirty="0"/>'
                 '<a:t>{text}</a:t></a:r></a:p>')


@dataclass
class Context:
    source_dir: str      # 文稿所在目录，相对图片路径以它为基准
    year: str            # Tag 不含年份时补上的年份
    lead: bool = False   # 方向 C（通用 / 项目汇报模式）：结论画成 lead、Next 画在页底一行；
                         # False = 月报模式，结论和 Next 进底部红条


def fill_titles(slide, page):
    """主标题进 idx 0、action title 进 idx 10（所有正文版式都有这两个占位框）。"""
    slide.placeholders[0].text = page.title
    if page.action_title:
        slide.placeholders[10].text = page.action_title


def warn(msg):
    print(f"警告: {msg}", file=sys.stderr)


def url_cache_dir():
    """URL 图片的下载缓存目录（PPT 必须内嵌图片字节；缓存后可离线重复生成）。"""
    return os.environ.get("MILWAUKEE_PPT_CACHE") or os.path.join(
        os.path.expanduser("~"), ".cache", "milwaukee-ppt", "url-images")


def load_image(ref, ctx, page_title):
    """图片引用 → 本地路径。支持相对路径（相对文稿）、绝对路径、http(s) URL。"""
    if ref.startswith(("http://", "https://")):
        cache_dir = url_cache_dir()
        os.makedirs(cache_dir, exist_ok=True)
        ext = os.path.splitext(ref.split("?")[0])[1] or ".png"
        cache = os.path.join(cache_dir, hashlib.md5(ref.encode()).hexdigest() + ext)
        if not os.path.exists(cache):
            try:
                req = urllib.request.Request(ref, headers={"User-Agent": "Mozilla/5.0"})
                with urllib.request.urlopen(req, timeout=30) as resp:
                    data = resp.read()
                with open(cache, "wb") as fh:
                    fh.write(data)
            except Exception as e:                      # noqa: BLE001
                raise DeckError(f"页「{page_title}」下载图片失败: {ref}\n  {e}")
        return cache
    p = ref if os.path.isabs(ref) else os.path.join(ctx.source_dir, ref)
    if not os.path.exists(p):
        raise DeckError(f"页「{page_title}」引用的图片不存在: {ref}")
    return p


def _next_shape_id(slide):
    ids = [int(e.get("id")) for e in slide.shapes._spTree.iter(qn("p:cNvPr"))
           if e.get("id")]
    return (max(ids) + 1) if ids else 100


def _caption_geometry(lines, page_title):
    """底边固定，按估算行数向上加高（CJK 宽字符按 2 计，额外半行余量防贴框）。

    这是月报红条固定的算法（几何直接决定输出），不改用下面的 text_lines，
    以保证月报模式的输出不变。"""
    rows = 0
    for text in lines:
        w = sum(2 if ord(ch) > 0x2E7F else 1 for ch in text)
        rows += max(1, -(-w // CAPTION_CHARS_PER_LINE))
    cy = max(CAPTION_MIN_CY,
             rows * CAPTION_LINE_EMU + CAPTION_INSETS_EMU + CAPTION_LINE_EMU // 2)
    y = CAPTION_BOTTOM - cy
    if y < CAPTION_MAX_TOP:
        warn(f"页「{page_title}」红条文字过多 ({rows} 行), 已顶到图片区底缘, "
             f"建议精简结论/Next")
        y = CAPTION_MAX_TOP
        cy = CAPTION_BOTTOM - y
    return y, cy


def conclusion_text(page):
    """本页一句话结论（普通文字行拼成一行）；没写为 ""。月报进红条，方向 C 画成 lead。"""
    return " ".join(page.conclusion).strip()


def next_text(page):
    """「Next: …」/「下一步：…」一行（保留作者写的前缀）；没写 Next 为 ""。"""
    if not page.next:
        return ""
    sep = "：" if page.next_label == "下一步" else ": "
    return f"{page.next_label}{sep}{page.next}"


def caption_lines(page):
    """页面的一句话结论 + 可选 Next，组成红条的各行。"""
    return [line for line in (conclusion_text(page), next_text(page)) if line]


def add_caption_bar(slide, page):
    """有结论或 Next 时画底部红条；样式全部由 CAPTION_SP_XML 的 lstStyle 继承。"""
    lines = caption_lines(page)
    if not lines:
        return None
    y, cy = _caption_geometry(lines, page.title)
    sp = parse_xml(CAPTION_SP_XML.format(id=_next_shape_id(slide), y=y, cy=cy,
                                         red=BRAND_RED))
    tx_body = sp.find(qn("p:txBody"))
    for text in lines:
        tx_body.append(parse_xml(CAPTION_P_XML.format(text=escape(text))))
    slide.shapes._spTree.append(sp)
    return sp


def add_month_tag(slide, tag, year, right=None, top=None):
    """正文区右上角极简月份标识（无框、灰字）；只在页面写了 Tag 时显示。

    Tag 不含 4 位年份时补上文稿年份，如 "Mar" → "Mar 2026"。
    right / top（EMU）默认是正文区右缘和 1.12"；有 PDCA 图标时由调用方让到图标左侧。"""
    if not tag:
        return None
    text = tag if re.search(r"\d{4}", tag) else f"{tag} {year}"
    w, h = Inches(1.7), Inches(0.3)
    right = Inches(IMG_LEFT + IMG_W) if right is None else right
    top = Inches(1.12) if top is None else top
    tb = slide.shapes.add_textbox(right - w, top, w, h)
    tb.name = "Month Tag"
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Pt(0)
    tf.margin_top = tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = text
    r.font.size = Pt(11)
    r.font.bold = False
    r.font.name = LATIN_FONT
    r.font.color.rgb = RGBColor(0x6E, 0x6E, 0x74)
    return tb


def add_page_chrome(slide, page, ctx):
    """所有版式通用的页面装饰：右上 PDCA 图标和月份标识 + 结论 / Next。

    月报模式：结论和 Next 进底部红条（叠在最上层）。方向 C（ctx.lead）：结论画成标题下的
    lead（add_lead），Next 画在页底一行（add_next_line）。版式渲染函数画完自己的内容后最后调用它；
    内容要排在 body_area(page, ctx) 里，才不会和 lead / Next 行重叠。"""
    if page.pdca:
        left = add_pdca_icons(slide, page.pdca)
        add_month_tag(slide, page.tag, ctx.year, right=left - Inches(PDCA_TAG_GAP),
                      top=Inches(PDCA_TAG_TOP))
    else:
        add_month_tag(slide, page.tag, ctx.year)
    if ctx.lead:
        add_lead(slide, page)
        add_next_line(slide, page)
    else:
        add_caption_bar(slide, page)


# ---- 方向 C：lead、页底 Next 行、正文区 -------------------------------------------
# lead：标题下、正文区顶部，黑色粗体 22pt；左侧一条品牌红短竖线（Lead Accent）。
# 写了 Tag 时 lead 右端让到月份标识框左侧（PDCA 图标在 1.30" 以上，不用让）。
# 一行时框高 0.56"（1.30–1.86"），正文区从 2.05" 起；多一行框加高一行、正文区跟着下移。
LEAD_TOP = 1.30
LEAD_PT = 22
LEAD_ONE_LINE_H = 0.56
LEAD_LINE_H = LEAD_PT * LINE_SPACING / 72
LEAD_MAX_LINES = 2            # 超过就警告
ACCENT_W = 0.08               # 红竖线宽
ACCENT_INSET = 0.06           # 红竖线比 lead 框上下各短
LEAD_LEFT = IMG_LEFT + 0.24   # lead 文字左缘（竖线 → 文字留 0.16"）
LEAD_TAG_GAP = 0.1            # lead 右端 → 月份标识框左缘
TAG_W = 1.7                   # add_month_tag 的框宽
LEAD_BODY_GAP = 0.19          # lead 下沿 → 正文区上沿
# Next 行：页底一行黑色粗体 16pt，下沿固定 6.94"；上方一条 RULE 色细线（Next Rule）。
# 一行时框 6.58–6.94"、细线在 6.50"，正文区下沿 6.40"；文字多一行框向上加高一行。
NEXT_PT = 16
NEXT_BOTTOM = 6.94
NEXT_ONE_LINE_H = 0.36
NEXT_LINE_H = NEXT_PT * LINE_SPACING / 72
NEXT_RULE_GAP = 0.08          # 细线 → Next 行
NEXT_RULE_PT = 0.75
NEXT_BODY_GAP = 0.10          # 正文区下沿 → 细线
BODY_BOTTOM = 6.95            # 没有 Next 时正文区下沿


@dataclass(frozen=True)
class Area:
    """矩形区域（英寸）。emu() → (left, top, width, height) 的 EMU 整数，可直接给 python-pptx。"""
    left: float
    top: float
    width: float
    height: float

    @property
    def right(self):
        return self.left + self.width

    @property
    def bottom(self):
        return self.top + self.height

    def emu(self):
        return tuple(int(Inches(v)) for v in (self.left, self.top, self.width, self.height))


def _lead_right(page):
    """lead 文字框右缘（英寸）：写了 Tag 时让到月份标识框左侧。"""
    right = IMG_LEFT + IMG_W
    if not page.tag:
        return right
    tag_right = right
    if page.pdca:
        n = len(page.pdca)
        tag_right = right - n * PDCA_SIZE - (n - 1) * PDCA_GAP - PDCA_TAG_GAP
    return min(right, tag_right - TAG_W - LEAD_TAG_GAP)


def lead_area(page):
    """方向 C 下 lead 文字框的矩形；本页没有结论返回 None。"""
    text = conclusion_text(page)
    if not text:
        return None
    width = _lead_right(page) - LEAD_LEFT
    lines = int(text_lines(text, LEAD_PT, width))
    return Area(LEAD_LEFT, LEAD_TOP, width, LEAD_ONE_LINE_H + (lines - 1) * LEAD_LINE_H)


def next_area(page):
    """方向 C 下 Next 行文字框的矩形；本页没写 Next 返回 None。"""
    text = next_text(page)
    if not text:
        return None
    lines = int(text_lines(text, NEXT_PT, IMG_W))
    height = NEXT_ONE_LINE_H + (lines - 1) * NEXT_LINE_H
    return Area(IMG_LEFT, NEXT_BOTTOM - height, IMG_W, height)


def body_area(page, ctx):
    """本页正文区（英寸 Area）：版式内容都排在这里面。

    月报模式：原来的正文区 0.56 / 1.30"、12.20 × 4.30"（下面留给底部红条）。
    方向 C：左右不变；上沿在 lead 下方（没有结论时 1.30"），下沿在 Next 细线上方
    （没有 Next 时 6.95"）。导航版式（section / agenda）不用它，正文区不变。"""
    if not ctx.lead:
        return Area(IMG_LEFT, IMG_TOP, IMG_W, IMG_H)
    lead, nxt = lead_area(page), next_area(page)
    top = lead.bottom + LEAD_BODY_GAP if lead else IMG_TOP
    bottom = nxt.top - NEXT_RULE_GAP - NEXT_BODY_GAP if nxt else BODY_BOTTOM
    return Area(IMG_LEFT, top, IMG_W, bottom - top)


def fit_body(rect, page, ctx):
    """模板槽位矩形 (left, top, width, height)（EMU）→ 版式实际排内容的矩形：
    月报模式原样返回；方向 C 下左右沿用槽位，上下沿换成 body_area。"""
    if not ctx.lead:
        return rect
    area = body_area(page, ctx)
    left, _, width, _ = rect
    return left, int(Inches(area.top)), width, int(Inches(area.height))


def add_lead(slide, page):
    """方向 C 的 lead：结论文字（Lead）+ 左侧品牌红短竖线（Lead Accent）。没有结论不画。"""
    area = lead_area(page)
    if area is None:
        return None
    text = conclusion_text(page)
    lines = text_lines(text, LEAD_PT, area.width)
    if lines > LEAD_MAX_LINES:
        warn(f"页「{page.title}」的结论（lead）估计 {int(lines)} 行, 正文区被压低, "
             f"建议精简到 {LEAD_MAX_LINES} 行以内: {text}")
    add_shape(slide, MSO_SHAPE.RECTANGLE, "Lead Accent", Inches(IMG_LEFT),
              Inches(area.top + ACCENT_INSET), Inches(ACCENT_W),
              Inches(area.height - 2 * ACCENT_INSET), BRAND_RED)
    tf = add_textbox(slide, "Lead", *area.emu(), anchor=MSO_ANCHOR.MIDDLE)
    add_run(tf.paragraphs[0], text, LEAD_PT, True, INK)
    return tf


def add_next_line(slide, page):
    """方向 C 的页底 Next：细线（Next Rule）+ 一行黑色粗体（Next Line）。没写 Next 不画。"""
    area = next_area(page)
    if area is None:
        return None
    y = Inches(area.top - NEXT_RULE_GAP)
    add_line(slide, "Next Rule", Inches(area.left), y, Inches(area.right), y, RULE, NEXT_RULE_PT)
    tf = add_textbox(slide, "Next Line", *area.emu(), anchor=MSO_ANCHOR.MIDDLE)
    add_run(tf.paragraphs[0], next_text(page), NEXT_PT, True, INK)
    return tf


# PDCA 图标（项目汇报模式，`PDCA:`）：assets/pdca/<字母>_icon.png（由同名 SVG 预转，
# 转换命令见 scripts/README.md），运行时只插 PNG。
# 放在 banner 下沿和正文槽位上沿之间的窄条里（1.02–1.30"），右缘贴正文区右缘；
# 各版式的内容都从 IMG_TOP（1.30"）起排，图标不压内容，版式不用为它让位。
# 组合（D/C）按文稿顺序从左到右排。写了 Tag 时月份标识让到图标左侧、纵向中线和图标对齐。
PDCA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "assets", "pdca")
PDCA_SIZE = 0.28        # 图标边长（英寸）
PDCA_TOP = 1.02         # 图标上沿：banner 下沿 1.0" 下留 0.02"，下沿正好到 IMG_TOP
PDCA_GAP = 0.06         # 组合图标之间的间距
PDCA_TAG_GAP = 0.12     # 月份标识文字右缘到最左图标的间距
PDCA_TAG_TOP = 1.07     # 有图标时月份标识框上沿：11pt 一行的中线落在图标中线 1.16"


def add_pdca_icons(slide, letters):
    """右上角按顺序插 PDCA 图标（PDCA Icon 1..n），返回最左图标的左缘（EMU）。"""
    size, gap = Inches(PDCA_SIZE), Inches(PDCA_GAP)
    right = Inches(IMG_LEFT + IMG_W)
    left = right - len(letters) * size - (len(letters) - 1) * gap
    for n, letter in enumerate(letters, start=1):
        pic = slide.shapes.add_picture(os.path.join(PDCA_DIR, f"{letter}_icon.png"),
                                       left + (n - 1) * (size + gap), Inches(PDCA_TOP),
                                       size, size)
        pic.name = f"PDCA Icon {n}"
    return left


# ---- 画原生形状的小工具 ---------------------------------------------------------

def slot(slide, idx):
    """按 idx 取占位框（槽位）；没有返回 None。按 idx 找，不按名字找。"""
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    return None


def remove_shape(shape):
    """从页面上删掉形状（空槽位画完就删，免得编辑视图里露出提示文字）。"""
    shape._element.getparent().remove(shape._element)


def take_slot(slide, idx):
    """取出槽位矩形 (left, top, width, height)（EMU）并删掉槽位占位框。"""
    ph = slot(slide, idx)
    rect = (ph.left, ph.top, ph.width, ph.height)
    remove_shape(ph)
    return rect


def strip_theme_style(shape):
    """去掉自选图形自带的主题样式引用（p:style 会带出主题阴影和 accent 色）；
    填充、描边、文字颜色全部由调用方显式给。"""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    return shape


def add_run(paragraph, text, pt, bold, color):
    """追加一个 run：Arial + 中文 PingFang SC（latin 之后写 a:ea），字号、粗细、颜色显式给。

    color 是 6 位十六进制字符串，如 "1A1A1A"。"""
    r = paragraph.add_run()
    r.text = text
    f = r.font
    f.size = Pt(pt)
    f.bold = bold
    f.name = LATIN_FONT
    f.color.rgb = RGBColor.from_string(color)
    latin = r._r.get_or_add_rPr().find(qn("a:latin"))
    latin.addnext(latin.makeelement(qn("a:ea"), {"typeface": EA_FONT}))
    return r


def add_textbox(slide, name, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    """起好名字、四边内边距 0 的文本框，返回 text_frame（坐标 EMU，可为浮点）。"""
    tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def add_shape(slide, kind, name, x, y, w, h, fill, line=None, line_pt=None):
    """扁平自选图形：去主题样式、关阴影，填充 / 描边只用显式颜色（fill 为 None = 无填充，
    line 为 None = 无描边）。坐标 EMU，可为浮点。"""
    sp = slide.shapes.add_shape(kind, int(x), int(y), int(w), int(h))
    sp.name = name
    strip_theme_style(sp)
    sp.shadow.inherit = False
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(line_pt)
    else:
        sp.line.fill.background()
    return sp


def add_line(slide, name, x1, y1, x2, y2, color, pt, dash=None):
    """直线连接线：去主题样式，颜色、线宽（pt）、虚线样式显式给。"""
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x1), int(y1), int(x2), int(y2))
    c.name = name
    strip_theme_style(c)
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(pt)
    if dash is not None:
        c.line.dash_style = dash
    return c


UNIT_PT = 28                  # 数字前后的小字（约、≥ 这类前缀和件、%、秒这类单位）：28pt 灰色、不加粗
UNIT_COLOR = "6B6B6B"

# 大数字 = 可选的 $ + - 前缀、数字、千分位逗号、小数点，再接数量级后缀（K/M/B/k/m/b/x/X，
# 后面不能紧跟字母）、+、/数字（24/7、3/5）；数字前可有一个前缀符号（约、~、≈、≥、≤、>、<、
# +/-、±），连同其后的空白一起作前缀小字。
_PREFIX = r"(?:约|~|～|≈|≥|≤|>|<|＞|＜|\+/-|±)\s*"
_NUMBER = re.compile(r"^(" + _PREFIX + r")?"
                     r"([$+\-]*\d[\d,]*(?:\.\d+)?(?:[KMBkmbxX](?![A-Za-z])|\+|/\d[\d,]*(?:\.\d+)?)*)"
                     r"(.*)$", re.S)
# 算单位的尾巴：% 或 CJK 开头（前面可有空白），或空白后面的英文单词；空白原样保留
_UNIT = re.compile(r"^(\s*[%\u2E80-\uFFFF]|\s+[A-Za-z])")


def chapter_row(line):
    """「编号 | 中文名 | 英文名」→ [编号, 中文名, 英文名]（各段去掉首尾空白）；不是三段返回 None。"""
    parts = [p.strip() for p in line.split("|")]
    return parts if len(parts) == 3 else None


def check_chapter_row(page, kind, row, show_row=False):
    """查章节行三段都不空、不超字数，否则抛 DeckError（含页标题和 kind，如「隔断页」）。
    show_row=True 时空段报错里带上整行原文（目录页一页多行，要指明是哪行）。"""
    for (label, limit), value in zip(CHAPTER_FIELDS, row):
        if not value:
            which = f"「{' | '.join(row)}」" if show_row else ""
            raise DeckError(f"页「{page.title}」{kind}{which}的{label}为空")
        if len(value) > limit:
            raise DeckError(f"页「{page.title}」{kind}的{label}「{value}」有 {len(value)} 字，"
                            f"最多 {limit} 字")


def is_number(value):
    """数字栏是不是「（前缀 +）数字」开头：和 split_number 同一条规则，审核也用它判定。"""
    return _NUMBER.match(value.strip()) is not None


def split_number(value):
    """拆成 (前缀, 大数字, 单位)，三段拼起来等于去掉首尾空白的原文（不补、不删空白）：

    「120 件」→ ("", "120", " 件")；「120件」→ ("", "120", "件")；「≥95%」→ ("≥", "95", "%")；
    「约 5 天」→ ("约 ", "5", " 天")；「24/7」「12.3M」「1,000+」「5x」整个是大数字；
    「Q3」「N/A」这类不以（前缀 +）数字开头的 → ("", 原文, "")。
    数字后面的尾巴只有 % / CJK 开头、或空白 + 英文单词时才算单位，否则并进大数字。"""
    v = value.strip()
    m = _NUMBER.match(v)
    if not m:
        return "", v, ""
    prefix, number, rest = m.group(1) or "", m.group(2), m.group(3)
    if rest and not _UNIT.match(rest):
        number, rest = number + rest, ""
    return prefix, number, rest


def number_width_in(value, pt, unit_pt=UNIT_PT, width=None):
    """「前缀 + 大数字 + 单位」一行的宽度估算（英寸）：大数字按 pt、前缀和单位按 unit_pt 量。
    width(text, pt) 是量宽函数，默认 text_width_in。"""
    width = width or text_width_in
    prefix, number, unit = split_number(value)
    return width(number, pt) + (width(prefix + unit, unit_pt) if prefix or unit else 0)


def add_number(paragraph, value, pt, color, unit_pt=UNIT_PT, unit_color=UNIT_COLOR):
    """在段落里写「前缀 + 大数字 + 单位」（拆法见 split_number）：大数字 pt 号加粗、color 色；
    前缀和单位 unit_pt 号灰色不加粗；同一段落所以同一基线。空白照原文，不补空格。
    返回 (数字 run, 单位 run 或 None)；有前缀时前缀 run 在数字 run 之前。"""
    prefix, number, unit = split_number(value)
    if prefix:
        add_run(paragraph, prefix, unit_pt, False, unit_color)
    num_run = add_run(paragraph, number, pt, True, color)
    unit_run = add_run(paragraph, unit, unit_pt, False, unit_color) if unit else None
    return num_run, unit_run


def fit_picture(slide, path, left, top, width, height):
    """图在矩形内等比缩放、居中，完整显示不裁切（原生图片，不用 insert_picture）。"""
    pic = slide.shapes.add_picture(path, left, top)
    scale = min(width / pic.width, height / pic.height)
    pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
    pic.left = left + (width - pic.width) // 2
    pic.top = top + (height - pic.height) // 2
    return pic


# ---- 文字估算（只用于判断放不放得下 / 是否警告，统一一个口径）-------------------

CJK_EM, LATIN_EM = 1.0, 0.55     # CJK / 全角按 1em，拉丁字符按约 0.55em


def _is_wide(ch):
    return ord(ch) > 0x2E7F


def text_width_in(text, pt):
    """单行文字宽度估算（英寸）。"""
    return sum(CJK_EM if _is_wide(ch) else LATIN_EM for ch in text) * pt / 72


def text_lines(text, pt, width_in):
    """text 在 width_in 宽、pt 号字下约占几行；空文字 0 行。"""
    if not text:
        return 0
    return max(1, -(-text_width_in(text, pt) // width_in))


# ---- 正文行：缩进子列表 ---------------------------------------------------------

_SUB_ITEM = re.compile(r"^[-*+]\s+(.*)$")


def sub_item(page, i):
    """page.lines[i] 是缩进的列表行（`  - 文字`，缩进 ≥1 列）时返回去掉列表符号的文字，
    否则 None。不缩进的 `- ` 行是本版式的条目行，不算子列表。"""
    if page.indents[i] <= 0:
        return None
    m = _SUB_ITEM.match(page.lines[i])
    return m.group(1).strip() if m else None


# ---- 条目行：! 重点标记、数据卡片（stats 类版式）------------------------------------

def key_mark(body, page):
    """条目正文开头的 `!`（重点 / 关键里程碑）：`!文字` → (True, 文字)，没写 → (False, 原文)。
    `!` 后的空白一并去掉；只认一个 `!`，写成 `!!` 抛 DeckError，不把多出来的 ! 当成文字。"""
    if not body.startswith("!"):
        return False, body
    rest = body[1:].lstrip()
    if rest.startswith("!"):
        raise DeckError(f"页「{page.title}」的条目开头只写一个 ! 标重点, 现在写了多个: {body}")
    return True, rest


LIST_ITEM = re.compile(r"^[-*+]\s+(.*)$")         # 不缩进的列表行：条目
SOURCE_LINE = re.compile(r"^(口径|source)\s*[:：]", re.I)   # 口径行


def parse_stat_card(page, line):
    """数据卡片行「- [!]标签 | 数字 | 一句说明」→ {"label", "value", "note", "key"}；
    不是列表行返回 None（由版式决定进红条还是别的）。DeckError：三段不全；
    重点条漏写列表符号（`!标签 | …`，别悄悄进红条）；`!!`。"""
    m = LIST_ITEM.match(line)
    if not m:
        if line.startswith("!") and "|" in line:
            raise DeckError(f"页「{page.title}」的重点数据条要以列表符号开头, "
                            f"写成「- !标签 | 数字 | 一句说明」: {line}")
        return None
    key, body = key_mark(m.group(1), page)
    parts = [x.strip() for x in body.split("|")]
    if len(parts) != 3 or not all(parts):
        raise DeckError(f"页「{page.title}」的数据条目要写成「标签 | 数字 | 一句说明」: {line}")
    return {"label": parts[0], "value": parts[1], "note": parts[2], "key": key}


def check_stat_cards(page, items, lo, hi):
    """卡片条数在 lo–hi 之间、最多一条标 ! 重点，否则抛 DeckError。"""
    if not lo <= len(items) <= hi:
        raise DeckError(f"页「{page.title}」的数据条目要 {lo}–{hi} 条, 现在 {len(items)} 条")
    keys = sum(it["key"] for it in items)
    if keys > 1:
        raise DeckError(f"页「{page.title}」最多一条数据标 ! 重点, 现在 {keys} 条")


# ---- 时间线状态（只用于时间线类版式）---------------------------------------------

DONE, DOING, TODO = "done", "doing", "todo"
DONE_GREEN = "2E7D32"         # 已完成：深绿，上面的字 / 对勾用白色
DOING_AMBER = "F2A900"        # 进行中：琥珀，上面的字用深色 DOING_INK
DOING_INK = "1A1A1A"
STATUS_COLORS = {DONE: DONE_GREEN, DOING: DOING_AMBER, TODO: BRAND_RED}  # 未完成只描边
STATUS_LABELS = {DONE: "已完成", DOING: "进行中", TODO: "未完成"}          # 图例文字
KEY_LABEL = "关键里程碑"                                                   # 图例第 4 项

# 状态图例：一行右对齐，每项 = 图形 + 灰色小字
LEGEND_H = 0.26               # 一行高
LEGEND_GAP_IN = 0.08          # 图形 → 文字
LEGEND_GAP_OUT = 0.3          # 项与项之间
LEGEND_TEXT_PAD = 0.1         # 文字框比估算宽度多留的余量（默认），防止估算偏小时被裁
LEGEND_STYLE = (12, False, "6B6B6B")

_STATUS_WORDS = {"完成": DONE, "已完成": DONE, "done": DONE,
                 "进行中": DOING, "doing": DOING,
                 "未完成": TODO, "todo": TODO}


def parse_status(text, page, n):
    """一条的状态值 → DONE / DOING / TODO；空串 → None（没写）。不区分大小写。
    认不出的值抛 DeckError（含页标题和第 n 条）。"""
    text = text.strip()
    if not text:
        return None
    st = _STATUS_WORDS.get(text.lower())
    if st is None:
        raise DeckError(f"页「{page.title}」第 {n} 条的状态写错了: {text!r}, "
                        f"可选: 完成 / 进行中 / 未完成 (或 done / doing / todo)")
    return st


def parse_statuses(page, raws, required=False):
    """一页各条的状态原文（没写为 ""）→ 状态列表；整页都没写且 required=False 时返回 None
    （旧写法，按黑白灰画）。DeckError（均含页标题）：值写错、只有部分条写了、
    required=True 却有条没写、“进行中”超过一个。"""
    sts = [parse_status(raw, page, n) for n, raw in enumerate(raws, start=1)]
    missing = [n for n, st in enumerate(sts, start=1) if st is None]
    if len(missing) == len(sts) and not required:
        return None
    if missing:
        need = "每条都必须写状态" if required else "状态要么全写、要么全不写 (部分写了)"
        raise DeckError(f"页「{page.title}」{need}, 第 {' / '.join(map(str, missing))} 条没写")
    doing = sts.count(DOING)
    if doing > 1:
        raise DeckError(f"页「{page.title}」“进行中”最多一个, 现在 {doing} 个")
    return sts


def add_legend(slide, right, top, entries, text_pad=LEGEND_TEXT_PAD):
    """时间线类版式的状态图例：一行、右端对齐到 right，上沿 top（EMU）。

    entries = [(图形占位宽（英寸）, draw, 文字)...]，draw(name, cx, cy) 以 (cx, cy) 为中心画图形；
    第 n 项的图形叫 `Legend Icon n`（draw 收到这个名字；图形上另加的形状也带编号 n，
    如 `Legend Check 1`），文字框叫 `Legend Label n`。text_pad：文字框比估算宽度多留的余量（英寸）。"""
    gap_in, gap_out = Inches(LEGEND_GAP_IN), Inches(LEGEND_GAP_OUT)
    slots = [Inches(w) for w, _, _ in entries]
    text_w = [Inches(text_width_in(label, LEGEND_STYLE[0]) + text_pad)
              for _, _, label in entries]
    x = right - sum(sw + gap_in + tw for sw, tw in zip(slots, text_w)) - gap_out * (len(entries) - 1)
    h = Inches(LEGEND_H)
    cy = top + h // 2
    for n, ((_, draw, label), sw, tw) in enumerate(zip(entries, slots, text_w), start=1):
        draw(f"Legend Icon {n}", x + sw // 2, cy)
        tf = add_textbox(slide, f"Legend Label {n}", x + sw + gap_in, top, tw, h,
                         MSO_ANCHOR.MIDDLE, wrap=False)
        add_run(tf.paragraphs[0], label, *LEGEND_STYLE)
        x += sw + gap_in + tw + gap_out

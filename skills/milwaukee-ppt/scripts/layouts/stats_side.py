"""stats-side 版式：左边竖排数据横卡、右边带编号的结论区，基于模板 layout 5「Stats」。

文稿写法：

    - Lead time | 5 d | Mock average per request       ← 2–3 张横卡「标签 | 数字 | 一句说明」
    - !Rework | -40% | Sample reduction after the change   ← `!` 标唯一重点（红竖条 + 红数字）
    ### Mock findings                                    ← 可选，结论区小标题，默认「结论」
    > Mock finding one.                                  ← 结论要点，自动编号，1–4 条
    口径：Mock source note.                             ← 可选，一行（或 Source:），灰色小字贴结论区底
    Mock one-line conclusion.                           ← 其余普通文字行进底部红条

卡片不带解读：卡片行下面缩进的 `- ` 子列表报错（解读写成 `> ` 要点，或改用 stats）。
字段：fields = {"items": [{"label", "value", "note", "key"}...], "points": [...],
"source": str 或 None, "heading": str}
渲染：槽位 idx 18 左侧 6.0" 竖排横卡（浅灰底），卡内大数字（单位缩小）在左、标签和说明在右；
右侧是结论区：小标题、细分隔线、编号要点，口径贴底。几何以原型第 3 页为参照。
形状名：Side Card N、Side Accent N（只有重点卡）、Side Number N、Side Text N、
Side Conclusion Title、Side Conclusion Rule、Side Points、Stat Source。
"""
import re

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from mdparse import DeckError

from .common import (BRAND_RED, UNIT_PT, add_number, add_page_chrome, add_run, fill_titles,
                     split_number, strip_theme_style, sub_item, take_slot,
                     text_lines, text_width_in, warn)

NAME = "stats-side"
LAYOUT_INDEX = 5
SLOT_IDX = 18
MIN_ITEMS, MAX_ITEMS = 2, 3
MAX_POINTS = 4

# 几何（英寸，原型第 3 页）
LEFT_W = 6.0              # 左侧卡片列宽
CARD_GAP = 0.3            # 卡片竖向间距
ACCENT_W = 0.06           # 重点卡左侧红竖条
NUM_X, NUM_W = 0.25, 2.5  # 卡内数字框：距卡左边、宽
TEXT_X = 2.85             # 卡内标签 / 说明框距卡左边
TEXT_PAD_R = 0.2          # 标签 / 说明框距卡右边

# 字号与颜色
NUMBER_PT, LABEL_PT, NOTE_PT = 60, 16, 14
NUMBER_MIN_PT = 36        # 数字长时统一缩小，最低到这里
LABEL_NOTE_GAP_PT = 4     # 标签与说明之间
LINE = 1.2                # 行高 = 字号 × 1.2（估算用）
INK, GRAY = "1A1A1A", "6B6B6B"
CARD_FILL = "F5F5F5"
RULE_COLOR = "E5E5E5"
HEADING_PT, POINT_PT, SOURCE_PT = 18, 16, 12
DEFAULT_HEADING = "结论"

# 结论区（英寸）
COL_GAP = 0.5             # 左卡片列与结论区之间
HEADING_H = 0.35
RULE_Y = 0.5              # 分隔线距槽位顶
POINTS_Y = 0.7            # 要点距槽位顶
POINT_SPACE_PT = 14       # 要点段前距
POINT_LINE = 1.15         # 要点行距倍数
SOURCE_H = 0.5            # 口径框高（文字贴底）

_ITEM = re.compile(r"^[-*+]\s+(!?)\s*(.*)$")
_POINT = re.compile(r"^>\s*(.*)$")
_HEADING = re.compile(r"^###\s+(.*)$")
_SOURCE = re.compile(r"^(口径|source)\s*[:：]", re.I)


def parse(page):
    def err(msg):
        raise DeckError(f"页「{page.title}」{msg}")

    items, points, sources, headings = [], [], [], []
    for i, line in enumerate(page.lines):
        if not line:
            continue
        if sub_item(page, i) is not None:
            err(f"的 stats-side 卡片不带解读子列表, 解读改写成 `> ` 开头的结论要点, "
                f"或改用 Layout: stats: {line}")
        point, heading = _POINT.match(line), _HEADING.match(line)
        if point:
            if point.group(1).strip():
                points.append(point.group(1).strip())
        elif heading:
            headings.append(heading.group(1).strip())
        elif _SOURCE.match(line):
            sources.append(line)
        elif _ITEM.match(line):
            m = _ITEM.match(line)
            parts = [x.strip() for x in m.group(2).split("|")]
            if len(parts) != 3 or not all(parts):
                err(f"的数据条目要写成「标签 | 数字 | 一句说明」: {line}")
            items.append({"label": parts[0], "value": parts[1], "note": parts[2],
                          "key": bool(m.group(1))})
        elif line.startswith("!") and "|" in line:     # 重点条漏写列表符号，别悄悄进红条
            err(f"的重点数据条要以列表符号开头, 写成「- !标签 | 数字 | 一句说明」: {line}")
        else:
            page.conclusion.append(line)

    if not MIN_ITEMS <= len(items) <= MAX_ITEMS:
        err(f"的数据条目要 {MIN_ITEMS}–{MAX_ITEMS} 条, 现在 {len(items)} 条")
    keys = sum(it["key"] for it in items)
    if keys > 1:
        err(f"最多一条数据标 ! 重点, 现在 {keys} 条")
    if not points:
        err("缺少结论要点, 用 `> ` 开头的行写 1–4 条")
    if len(points) > MAX_POINTS:
        err(f"的结论要点最多 {MAX_POINTS} 条 (`> ` 开头的行), 现在 {len(points)} 条")
    if len(sources) > 1:
        err(f"只能有一行口径 (口径： / Source:), 现在 {len(sources)} 行")
    if len(headings) > 1:
        err(f"只能有一个结论区小标题 (### ), 现在 {len(headings)} 个")
    page.fields = {"items": items, "points": points,
                   "source": sources[0] if sources else None,
                   "heading": headings[0] if headings else DEFAULT_HEADING}


def _shape(slide, name, x, y, w, h, fill):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    sp.name = name
    strip_theme_style(sp)          # 去掉主题样式引用：LibreOffice 会按它补阴影
    sp.shadow.inherit = False
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    sp.line.fill.background()
    sp.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT   # add_shape 默认居中
    return sp


def _box(slide, name, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def _number_width_in(value, pt):
    number, unit = split_number(value)
    return text_width_in(number, pt) + (text_width_in(" " + unit, UNIT_PT) if unit else 0)


def _number_size(items, title):
    """所有大数字共用一个字号：60pt 放不下最长的那个时统一缩小（整数 pt），并警告。"""
    def fits(pt):
        return all(_number_width_in(it["value"], pt) <= NUM_W for it in items)

    pt = NUMBER_PT
    while pt > NUMBER_MIN_PT and not fits(pt):
        pt -= 1
    if pt < NUMBER_PT:
        longest = max(items, key=lambda it: _number_width_in(it["value"], NUMBER_PT))["value"]
        tail = "" if fits(pt) else " 仍可能放不下, 建议缩短数字"
        warn(f"页「{title}」数字「{longest}」在 {NUMBER_PT}pt 放不下, "
             f"大数字已统一缩小到 {pt}pt{tail}")
    return pt


def _cards(slide, items, left, top, height, title):
    n = len(items)
    card_w = Inches(LEFT_W)
    card_h = (height - Inches(CARD_GAP) * (n - 1)) // n
    number_pt = _number_size(items, title)
    text_w_in = LEFT_W - TEXT_X - TEXT_PAD_R
    for i, it in enumerate(items, start=1):
        need = (text_lines(it["label"], LABEL_PT, text_w_in) * LABEL_PT * LINE
                + LABEL_NOTE_GAP_PT
                + text_lines(it["note"], NOTE_PT, text_w_in) * NOTE_PT * LINE) / 72
        if need > Emu(card_h).inches:
            warn(f"页「{title}」第 {i} 张卡的标签和说明估计放不下, 建议缩短说明")
    for i, it in enumerate(items, start=1):
        y = top + (card_h + Inches(CARD_GAP)) * (i - 1)
        _shape(slide, f"Side Card {i}", left, y, card_w, card_h, CARD_FILL)
        if it["key"]:
            _shape(slide, f"Side Accent {i}", left, y, Inches(ACCENT_W), card_h, BRAND_RED)
        tf = _box(slide, f"Side Number {i}", left + Inches(NUM_X), y, Inches(NUM_W), card_h,
                  MSO_ANCHOR.MIDDLE, wrap=False)
        add_number(tf.paragraphs[0], it["value"], number_pt, BRAND_RED if it["key"] else INK)
        tf = _box(slide, f"Side Text {i}", left + Inches(TEXT_X), y,
                  card_w - Inches(TEXT_X + TEXT_PAD_R), card_h, MSO_ANCHOR.MIDDLE)
        add_run(tf.paragraphs[0], it["label"], LABEL_PT, True, INK)
        p = tf.add_paragraph()
        p.space_before = Pt(LABEL_NOTE_GAP_PT)
        add_run(p, it["note"], NOTE_PT, False, GRAY)


def _numbered(paragraph, pt):
    """悬挂缩进 + 自动编号（1 2 3，灰色 Arial），换行后文字对齐。"""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Pt(pt * 1.4))))
    pPr.set("indent", str(-int(Pt(pt * 1.4))))
    clr = pPr.makeelement(qn("a:buClr"), {})
    clr.append(clr.makeelement(qn("a:srgbClr"), {"val": GRAY}))
    pPr.append(clr)
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buAutoNum"), {"type": "arabicPlain"}))


def _conclusion(slide, fields, x, top, w, bottom, title):
    tf = _box(slide, "Side Conclusion Title", x, top, w, Inches(HEADING_H))
    add_run(tf.paragraphs[0], fields["heading"], HEADING_PT, True, INK)
    _shape(slide, "Side Conclusion Rule", x, top + Inches(RULE_Y), w, Pt(1), RULE_COLOR)

    points_top = top + Inches(POINTS_Y)
    points_bottom = bottom - (Inches(SOURCE_H) if fields["source"] else 0)
    tf = _box(slide, "Side Points", x, points_top, w, points_bottom - points_top)
    text_w_in = Emu(w).inches - POINT_PT * 1.4 / 72        # 去掉悬挂缩进
    need = (sum(text_lines(t, POINT_PT, text_w_in) for t in fields["points"])
            * POINT_PT * LINE * POINT_LINE
            + POINT_SPACE_PT * (len(fields["points"]) - 1)) / 72
    if need > Emu(points_bottom - points_top).inches:
        warn(f"页「{title}」的结论要点估计放不下, 建议缩短要点或减少条数")
    for j, text in enumerate(fields["points"]):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        if j:
            p.space_before = Pt(POINT_SPACE_PT)
        p.line_spacing = POINT_LINE
        _numbered(p, POINT_PT)
        add_run(p, text, POINT_PT, False, INK)

    if fields["source"]:
        if text_lines(fields["source"], SOURCE_PT, Emu(w).inches) * SOURCE_PT * LINE / 72 > SOURCE_H:
            warn(f"页「{title}」的口径估计放不下, 建议缩短到两行以内")
        tf = _box(slide, "Stat Source", x, bottom - Inches(SOURCE_H), w, Inches(SOURCE_H),
                  MSO_ANCHOR.BOTTOM)
        add_run(tf.paragraphs[0], fields["source"], SOURCE_PT, False, GRAY)


def render(prs, page, ctx):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    fill_titles(slide, page)
    left, top, width, height = take_slot(slide, SLOT_IDX)
    _cards(slide, page.fields["items"], left, top, height, page.title)
    x = left + Inches(LEFT_W + COL_GAP)
    _conclusion(slide, page.fields, x, top, left + width - x, top + height, page.title)
    add_page_chrome(slide, page, ctx)
    return slide

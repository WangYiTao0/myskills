"""agenda 版式（Agenda 目录页）：模板 layout 0「Blank Slide」，列出本次汇报的各章节。

文稿写法：3–8 行「编号 | 中文名 | 英文名」，一行一章，写法和隔断页（section）完全相同，
如 `03 | 试点推进 | PILOT`；编号不能重复。目录页不放结论红条：写了普通文字行或 Next: 都报错。
Tag: 照常可用。

排版：正文区（1.30–5.60"）里，3–4 条排一列，5–8 条排两列（先填满左列再填右列，左列条数
= 总数一半向上取整）；整组垂直居中。每条一行：左边章节编号（32pt bold #1A1A1A），一条
#BFBFBF 竖线，右边中文名（20pt bold #1A1A1A）在上、英文名（12pt #6B6B6B）在下。
只用黑白灰，banner 以外不用红。和 evidence、section 共用模板 layout 0，inspect 靠 MARK_SHAPE 区分。
"""
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

from mdparse import DeckError

from .common import (IMG_H, IMG_LEFT, IMG_TOP, IMG_W, INK, LINE_SPACING, MUTED, RULE, add_line,
                     add_page_chrome, add_run, add_textbox, chapter_row, check_chapter_row,
                     fill_titles, text_width_in)

NAME = "agenda"
LAYOUT_INDEX = 0
MARK_SHAPE = "Agenda Number 1"

MIN_ITEMS, MAX_ITEMS = 3, 8
ONE_COLUMN_MAX = 4                                  # 超过 4 条分两列

NUMBER_PT, NAME_PT, ENGLISH_PT = 32, 20, 12
INDENT = 1.0          # 单列时从正文区左沿缩进（英寸），与 section 对齐
COLUMN_GAP = 0.4      # 两列之间（英寸）
GAP = 0.2             # 编号 → 竖线 → 名字（英寸）
NAME_GAP = Inches(0.04)
ROW_PITCH = 0.95      # 行距（英寸）：4 行 3.8"，放得进 4.30" 正文区


def parse(page):
    items = []
    for line in page.lines:
        if not line:
            continue
        parts = chapter_row(line)
        if parts is None:
            raise DeckError(f"页「{page.title}」目录页每行写「编号 | 中文名 | 英文名」，"
                            f"不放结论红条；这一行不对: {line!r}")
        items.append(parts)
    if page.next:
        raise DeckError(f"页「{page.title}」目录页不放结论红条，去掉 Next:")
    if not MIN_ITEMS <= len(items) <= MAX_ITEMS:
        raise DeckError(f"页「{page.title}」目录页须 {MIN_ITEMS}–{MAX_ITEMS} 行"
                        f"「编号 | 中文名 | 英文名」，现在 {len(items)} 行")
    for row in items:
        check_chapter_row(page, "目录页", row, show_row=True)
    numbers = [it[0] for it in items]
    dup = sorted({n for n in numbers if numbers.count(n) > 1})
    if dup:
        raise DeckError(f"页「{page.title}」目录页编号重复: {' / '.join(dup)}")
    page.fields = {"items": [{"number": n, "name": c, "english": e} for n, c, e in items]}


def render(prs, page, ctx):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    fill_titles(slide, page)
    items = page.fields["items"]
    two = len(items) > ONE_COLUMN_MAX
    rows = (len(items) + 1) // 2 if two else len(items)
    if two:
        col_w = (IMG_W - COLUMN_GAP) / 2
        lefts = [IMG_LEFT, IMG_LEFT + col_w + COLUMN_GAP]
    else:
        col_w = IMG_W - INDENT
        lefts = [IMG_LEFT + INDENT]
    num_w = max(text_width_in(it["number"], NUMBER_PT) for it in items) + 0.1   # 编号框统一宽
    top0 = IMG_TOP + (IMG_H - rows * ROW_PITCH) / 2

    num_h = Pt(NUMBER_PT * LINE_SPACING)
    name_h, eng_h = Pt(NAME_PT * LINE_SPACING), Pt(ENGLISH_PT * LINE_SPACING)
    block_h = name_h + NAME_GAP + eng_h
    for i, it in enumerate(items):
        n = i + 1
        col, row = divmod(i, rows)
        x = Inches(lefts[col])
        mid = Inches(top0 + (row + 0.5) * ROW_PITCH)
        tf = add_textbox(slide, f"Agenda Number {n}", x, mid - num_h / 2, Inches(num_w), num_h,
                         anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        add_run(tf.paragraphs[0], it["number"], NUMBER_PT, True, INK)
        rule_x = x + Inches(num_w + GAP)
        add_line(slide, f"Agenda Rule {n}", rule_x, mid - block_h / 2, rule_x, mid + block_h / 2,
                 RULE, 1.0)
        text_x = rule_x + Inches(GAP)
        text_w = Inches(lefts[col] + col_w) - text_x
        top = mid - block_h / 2
        tf = add_textbox(slide, f"Agenda Name {n}", text_x, top, text_w, name_h,
                         anchor=MSO_ANCHOR.BOTTOM, wrap=False)
        add_run(tf.paragraphs[0], it["name"], NAME_PT, True, INK)
        tf = add_textbox(slide, f"Agenda English {n}", text_x, top + name_h + NAME_GAP, text_w,
                         eng_h, wrap=False)
        add_run(tf.paragraphs[0], it["english"], ENGLISH_PT, False, MUTED)

    add_page_chrome(slide, page, ctx)
    return slide

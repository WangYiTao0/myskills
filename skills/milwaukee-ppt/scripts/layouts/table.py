"""table 版式（表格页）：模板 layout 0「Blank Slide」，md 管道表格 → 原生 pptx 表格。

文稿写法：md 管道表格，第一行是表头，表头下的分隔行（`| --- | --- |`）可写可不写；
2–6 列、表头以外 2–8 行，每行列数相同，表头每格必填。最多一行在第一格开头写 `!`
（`| !B 仓 | … |`）标重点。其余普通文字行进底部红条，Next: / Tag: 照常可用。
漏写开头 `|` 的表格行、列表行报错（免得悄悄进红条）。

排版：表格占正文区全宽（左边留出红竖条的位置），从正文区上沿排起（写了 Tag 时下移到
Month Tag 下面）；列等宽。14pt，表头加粗、#E5E5E5 底、下边 1.5pt #1A1A1A；正文 #1A1A1A、
无底色、每行下边 0.75pt #BFBFBF 细线。每格最多 2 行，行高按该行最多的行数定，
月报模式表格总高放不进正文区报错（渲染时按 body_area 判断）。方向 C（通用 / 项目模式）
字号按 GENERAL_TIERS 从 18pt 往下取放得下的一档，整表在正文区里垂直居中，最小一档也放不下
只警告。`!` 行文字加粗，左侧一条品牌红竖条（Table Key Bar）——
红色只作锚点（SKILL.md §6.3），表格本身只用黑白灰。表格不引用主题表格样式（主题样式会带出
accent 色），边框、底色全部显式写。和 evidence、section、agenda、team 共用模板 layout 0，
inspect 靠 MARK_SHAPE 区分。
"""
import re

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches

from mdparse import DeckError

from .common import (BRAND_RED, IMG_LEFT, IMG_TOP, IMG_W, INK, LINE_SPACING, LIST_ITEM, RULE,
                     add_page_chrome, add_run, add_shape, body_area, fill_titles, key_mark,
                     text_lines, warn)

NAME = "table"
LAYOUT_INDEX = 0
MARK_SHAPE = "Data Table"

MIN_COLS, MAX_COLS = 2, 6
MIN_ROWS, MAX_ROWS = 2, 8       # 表头以外的行
MAX_CELL_LINES = 2

FONT_PT = 14
HEAD_FILL = "E5E5E5"
HEAD_RULE_PT, RULE_PT = 1.5, 0.75
NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # 内置「无样式，无网格」

# 几何（英寸）
KEY_BAR_W = 0.06                # 红竖条宽
KEY_GAP = 0.06                  # 红竖条 → 表格
TABLE_LEFT = IMG_LEFT + KEY_BAR_W + KEY_GAP
TABLE_W = IMG_W - KEY_BAR_W - KEY_GAP
CELL_PAD_X, CELL_PAD_Y = 0.1, 0.06
TAG_BOTTOM = 1.42               # common.add_month_tag 的框下沿
TAG_GAP = 0.08
# 方向 C（通用 / 项目模式，#51）字号随行数放大：(字号 pt, 单元格上下内边距 ")，从大到小取
# 第一档「每格不超过 2 行、整表放得进正文区」的；最小一档就是月报的 14pt / 0.06"。
GENERAL_TIERS = ((18, 0.09), (16, 0.08), (FONT_PT, CELL_PAD_Y))

_ROW = re.compile(r"^\|(.*?)\|?$")
_SEP_CELL = re.compile(r"^:?-+:?$")


def _fail(page, msg):
    raise DeckError(f"页「{page.title}」{msg}")


def _cells(line):
    return [c.strip() for c in _ROW.match(line).group(1).split("|")]


def col_text_w(ncols):
    return TABLE_W / ncols - 2 * CELL_PAD_X


def table_top(page, ctx=None):
    """表格可排区域的上沿（英寸）：正文区上沿；写了 Tag 时不高于 Month Tag 下面。
    ctx 为 None 按月报口径算。方向 C 下整表在这里到正文区下沿之间垂直居中。"""
    if ctx is not None and ctx.lead:
        return max(body_area(page, ctx).top, TAG_BOTTOM + TAG_GAP if page.tag else 0)
    return TAG_BOTTOM + TAG_GAP if page.tag else IMG_TOP


def row_height(lines, pt=FONT_PT, pad=CELL_PAD_Y):
    return lines * pt * LINE_SPACING / 72 + 2 * pad


def _heights(rows, pt, pad):
    """各行行高（英寸）；有单元格超过 MAX_CELL_LINES 行返回 None。"""
    width = col_text_w(len(rows[0]))
    out = []
    for cells in rows:
        lines = [text_lines(c, pt, width) for c in cells]
        if max(lines) > MAX_CELL_LINES:
            return None
        out.append(row_height(max(1, *lines), pt, pad))
    return out


def parse(page):
    rows = []
    for line in page.lines:
        if not line:
            continue
        if line.startswith("|"):
            cells = _cells(line)
            if len(rows) == 1 and all(_SEP_CELL.match(c) for c in cells):
                continue                       # 表头下的分隔行
            rows.append(cells)
            continue
        if LIST_ITEM.match(line):
            _fail(page, f"表格页不写列表，内容写进表格: {line!r}")
        if "|" in line:
            _fail(page, f"表格行要以 | 开头，写成「| 格 | 格 |」: {line!r}")
        page.conclusion.append(line)
    if not rows:
        _fail(page, "表格页没有表格，写一个 md 管道表格（第一行是表头）")
    header, body = rows[0], rows[1:]
    ncols = len(header)
    if not MIN_COLS <= ncols <= MAX_COLS:
        _fail(page, f"表格须 {MIN_COLS}–{MAX_COLS} 列，现在 {ncols} 列；列多就拆成两页或删掉次要列")
    if not MIN_ROWS <= len(body) <= MAX_ROWS:
        _fail(page, f"表格表头以外须 {MIN_ROWS}–{MAX_ROWS} 行，现在 {len(body)} 行；"
                    f"行多就按类别拆页，或只留最关键的几行")
    for cells in body:
        if len(cells) != ncols:
            _fail(page, f"表格每行须和表头一样 {ncols} 列，这一行 {len(cells)} 列: "
                        f"{'| ' + ' | '.join(cells) + ' |'!r}")
    if any(c.startswith("!") for c in header):
        _fail(page, "表头不标重点，! 写在要强调那一行的第一格开头")
    if not all(header):
        _fail(page, "表头每一格都要写列名")
    key_row = None
    for i, cells in enumerate(body):
        if any(c.startswith("!") for c in cells[1:]):
            _fail(page, f"! 只写在一行的第一格开头: {'| ' + ' | '.join(cells) + ' |'!r}")
        key, cells[0] = key_mark(cells[0], page)
        if key:
            if key_row is not None:
                _fail(page, "表格最多一行标 !（只强调最该被记住或风险最高的一行）")
            key_row = i
    width = col_text_w(ncols)
    heights = []
    for cells in [header] + body:
        lines = [text_lines(c, FONT_PT, width) for c in cells]
        for c, n in zip(cells, lines):
            if n > MAX_CELL_LINES:
                _fail(page, f"表格单元格「{c}」超过 {MAX_CELL_LINES} 行（{ncols} 列时每格约 "
                            f"{int(width * MAX_CELL_LINES / (FONT_PT / 72))} 个汉字）；"
                            f"精简文字，或减少列数")
        heights.append(row_height(max(1, *lines)))
    page.fields = {"header": header, "rows": body, "key": key_row, "heights": heights}


def _cell_border(tcPr, tag, color, pt):
    ln = tcPr.makeelement(qn(f"a:{tag}"), {"w": str(int(pt * 12700)), "cmpd": "sng"})
    if color is None:
        ln.append(ln.makeelement(qn("a:noFill"), {}))
    else:
        fill = ln.makeelement(qn("a:solidFill"), {})
        fill.append(fill.makeelement(qn("a:srgbClr"), {"val": color}))
        ln.append(fill)
    return ln


def _style_cell(cell, text, bold, fill, rule, rule_pt, pt=FONT_PT, pad=CELL_PAD_Y):
    cell.margin_left = cell.margin_right = Inches(CELL_PAD_X)
    cell.margin_top = cell.margin_bottom = Inches(pad)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    if text:
        add_run(tf.paragraphs[0], text, pt, bold, INK)
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        tcPr.remove(child)
    for tag in ("lnL", "lnR", "lnT"):
        tcPr.append(_cell_border(tcPr, tag, None, 0))
    tcPr.append(_cell_border(tcPr, "lnB", rule, rule_pt))
    if fill is None:
        tcPr.append(tcPr.makeelement(qn("a:noFill"), {}))
    else:
        sf = tcPr.makeelement(qn("a:solidFill"), {})
        sf.append(sf.makeelement(qn("a:srgbClr"), {"val": fill}))
        tcPr.append(sf)


def monthly_height_problem(page, ctx):
    """月报模式下表格总高放不进正文区时返回原因（中文短句，带拆页建议），否则 None；
    方向 C 不报错（只警告），也返回 None。生成器报错和审核「表格太高」都用它。"""
    if ctx.lead:
        return None
    total = sum(page.fields["heights"])
    room = body_area(page, ctx).bottom - table_top(page, ctx)
    if total <= room + 1e-9:
        return None
    return (f"表格太高（约 {total:.2f}\"，正文区只有 {room:.2f}\"）；"
            f"把多行的格子写短，或拆成两页")


def render(prs, page, ctx):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    fill_titles(slide, page)
    f = page.fields
    header, body, heights = f["header"], f["rows"], f["heights"]
    ncols = len(header)
    y0 = table_top(page, ctx)
    room = body_area(page, ctx).bottom - y0
    pt, pad = FONT_PT, CELL_PAD_Y
    if not ctx.lead:
        problem = monthly_height_problem(page, ctx)
        if problem:
            _fail(page, problem)
    else:
        # 方向 C：字号从 18pt 往下试，取放得下的最大一档；整表在正文区里垂直居中
        for pt, pad in GENERAL_TIERS:
            hs = _heights([header] + body, pt, pad)
            if hs is not None and sum(hs) <= room + 1e-9:
                heights = hs
                break
        if sum(heights) > room + 1e-9:
            warn(f"页「{page.title}」表格太高（约 {sum(heights):.2f}\"，正文区只有 {room:.2f}\"），"
                 f"会压到页底的 Next 行；把多行的格子写短，或精简结论 / Next")
        y0 += max(0, room - sum(heights)) / 2
    top = Inches(y0)
    row_emu = [Inches(h) for h in heights]
    col_emu = Inches(TABLE_W / ncols)
    frame = slide.shapes.add_table(len(heights), ncols, Inches(TABLE_LEFT), top,
                                   col_emu * ncols, sum(row_emu))
    frame.name = MARK_SHAPE
    tbl = frame.table
    tbl.first_row, tbl.horz_banding = True, False
    style = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
    style.text = NO_STYLE
    for c in range(ncols):
        tbl.columns[c].width = col_emu
    for r, cells in enumerate([header] + body):
        tbl.rows[r].height = row_emu[r]
        head = r == 0
        bold = head or (r - 1 == f["key"])
        for c, text in enumerate(cells):
            _style_cell(tbl.cell(r, c), text, bold, HEAD_FILL if head else None,
                        INK if head else RULE, HEAD_RULE_PT if head else RULE_PT, pt, pad)
    if f["key"] is not None:
        r = f["key"] + 1
        add_shape(slide, MSO_SHAPE.RECTANGLE, "Table Key Bar", Inches(IMG_LEFT),
                  top + sum(row_emu[:r]), Inches(KEY_BAR_W), row_emu[r], BRAND_RED)
    add_page_chrome(slide, page, ctx)
    return slide

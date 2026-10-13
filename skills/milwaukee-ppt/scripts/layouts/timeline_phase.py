"""timeline-phase 版式：模板 layout 6「Timeline」，槽位 idx 19「Timeline Area」。

字段：3–6 条「- [!]日期 | 标题 | 说明 | 状态」，每条都必须写状态（说明可空）；
`!` 标关键里程碑 → fields["items"] = [(日期, 标题, 说明, 状态, 是否关键)]；
其余普通文字行进红条结论。

画法（原型第 5 页）：槽位横向等分成 n 段箭头（首段 homePlate、其余 chevron），段上写标题：
已完成绿底白字，进行中琥珀底深色字且加高，未完成浅粉底 + 红描边 + 红字。
日期在段上方；段下方细线连到标记（关键里程碑黑色大菱形 + “关键里程碑”，普通节点灰色小圆点），
说明在最下面。槽位右上一行图例，图形是缩小的箭头段 / 菱形。标题、文字估计放不下时警告。

形状名（N 从 1 起，对应第 N 条）：Phase Segment N（箭头段，标题在段内）、Phase Date N、
Phase Stem N（细线）、Phase Marker N、Phase Key Label N（只有关键里程碑有）、Phase Text N（说明）；
图例 Legend Swatch 1–4 / Legend Label 1–4（已完成 / 进行中 / 未完成 / 关键里程碑）。
"""
import re

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from mdparse import DeckError

from .common import (BRAND_RED, DOING, DOING_AMBER, DOING_INK, DONE, DONE_GREEN,
                     STATUS_LABELS, TODO, add_page_chrome, add_run, fill_titles, parse_statuses,
                     strip_theme_style, take_slot, text_lines, text_width_in, warn)

NAME = "timeline-phase"
LAYOUT_INDEX = 6
SLOT_IDX = 19
MIN_ITEMS, MAX_ITEMS = 3, 6

# 几何（英寸，相对槽位上沿 / 左沿；以原型第 5 页为准）
SEG_GAP = 0.06        # 相邻箭头段之间的空隙
SEG_CY = 1.8          # 箭头段中线
SEG_H = 0.72          # 箭头段高
SEG_H_DOING = 0.95    # 进行中段加高

TODO_FILL = "FCE6E8"  # 未完成段浅粉底（红色只用于描边和字）
TODO_LINE_PT = 1.5

# 段的样式：(底色, 描边色, 字色)
SEG_STYLE = {DONE: (DONE_GREEN, None, "FFFFFF"),
             DOING: (DOING_AMBER, None, DOING_INK),
             TODO: (TODO_FILL, BRAND_RED, BRAND_RED)}
SEG_TITLE_PT, SEG_TITLE_PT_DOING = 16, 18
# 箭头段的视觉中心相对外框中心的偏移：chevron 左侧有凹口往右偏，homePlate 右侧尖头往左偏。
# 段上标题、日期、细线、标记、说明都对齐这条中线。
MID_SHIFT_CHEVRON, MID_SHIFT_FIRST = 0.1, -0.05

DATE_TOP = -0.95      # 日期框上沿（相对段中线）
DATE_H = 0.3
MARK_CY = 0.85        # 标记中心（相对段中线）
KEY_D, PLAIN_D = 0.36, 0.14
LABEL_TOP = 0.25      # “关键里程碑”上沿（相对标记中心）
LABEL_H = 0.3
TEXT_TOP = 0.6        # 说明上沿（相对标记中心）
TEXT_INSET = 0.12     # 日期 / 说明框比一格窄的左右余量

INK, GRAY, LIGHT = "1A1A1A", "6B6B6B", "BFBFBF"
# (字号 pt, 粗体, 颜色)
DATE_STYLE = (12, True, GRAY)
DATE_STYLE_DOING = (14, True, INK)
LABEL_STYLE = (12, True, INK)
TEXT_STYLE = (14, False, GRAY)
KEY_LABEL = "关键里程碑"

# 图例（槽位右上；下移 LEGEND_TOP 让开右上角的 Month Tag）：图形是缩小的箭头段 / 菱形，
# 与阶段条一一对应
LEGEND_TOP = 0.2
LEGEND_H = 0.26
LEGEND_SEG_W, LEGEND_SEG_H = 0.34, 0.22
LEGEND_KEY_D = 0.24
LEGEND_GAP_IN, LEGEND_GAP_OUT = 0.08, 0.3   # 图形与文字、各项之间
LEGEND_STYLE = (12, False, GRAY)
LINE_SPACING = 1.2    # 行高 ≈ 字号 × 1.2（估算溢出用）

_ITEM = re.compile(r"^[-*+]\s+(.*)$")


def parse(page):
    rows = []
    for line in page.lines:
        if not line:
            continue
        m = _ITEM.match(line)
        if not m:
            page.conclusion.append(line)
            continue
        body = m.group(1).strip()
        key = body.startswith("!")
        parts = [x.strip() for x in body.lstrip("!").split("|")]
        if len(parts) == 3:          # 少了状态：按“缺状态”报错，而不是格式错
            parts.append("")
        if len(parts) != 4 or not parts[0] or not parts[1]:
            raise DeckError(f"页「{page.title}」的阶段条目要写成「[!]日期 | 标题 | 说明 | 状态」: "
                            f"{line}")
        rows.append((parts, key))
    if not MIN_ITEMS <= len(rows) <= MAX_ITEMS:
        raise DeckError(f"页「{page.title}」的阶段条目有 {len(rows)} 条, "
                        f"应为 {MIN_ITEMS}–{MAX_ITEMS} 条")
    statuses = parse_statuses(page, [p[3] for p, _ in rows], required=True)
    page.fields = {"items": [(p[0], p[1], p[2], st, key)
                             for (p, key), st in zip(rows, statuses)]}


def _shape(slide, kind, name, x, y, w, h, fill, line=None, line_pt=None):
    """原生自选图形：去主题样式和阴影，填充 / 描边只用显式颜色（坐标为 EMU）。"""
    sp = slide.shapes.add_shape(kind, int(x), int(y), int(w), int(h))
    strip_theme_style(sp)
    sp.shadow.inherit = False
    sp.name = name
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(line_pt)
    else:
        sp.line.fill.background()
    return sp


def _center_text(tf, text, pt, bold, color, anchor=MSO_ANCHOR.MIDDLE, wrap=False):
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, text, pt, bold, color)


def _textbox(slide, name, x, y, w, h, text, style, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tb.name = name
    _center_text(tb.text_frame, text, *style, anchor=anchor, wrap=wrap)
    return tb


def _stem(slide, name, x, y1, y2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x), int(y1), int(x), int(y2))
    strip_theme_style(c)
    c.name = name
    c.line.color.rgb = RGBColor.from_string(LIGHT)
    c.line.width = Pt(1)
    return c


def _legend(slide, right, y):
    """右对齐的一行图例：Legend Swatch N + Legend Label N（N = 1…4）。"""
    entries = [(MSO_SHAPE.CHEVRON, SEG_STYLE[st][:2], STATUS_LABELS[st])
               for st in (DONE, DOING, TODO)]
    entries.append((MSO_SHAPE.DIAMOND, (INK, None), KEY_LABEL))
    pt = LEGEND_STYLE[0]
    widths = []
    for kind, _, label in entries:
        sw = LEGEND_KEY_D if kind == MSO_SHAPE.DIAMOND else LEGEND_SEG_W
        widths.append((sw, text_width_in(label, pt) + 0.05))
    total = sum(sw + LEGEND_GAP_IN + tw for sw, tw in widths) + LEGEND_GAP_OUT * (len(entries) - 1)
    x = right - Inches(total)
    cy = y + Inches(LEGEND_H) / 2
    for n, ((kind, (fill, line), label), (sw, tw)) in enumerate(zip(entries, widths), start=1):
        sh = LEGEND_KEY_D if kind == MSO_SHAPE.DIAMOND else LEGEND_SEG_H
        _shape(slide, kind, f"Legend Swatch {n}", x, cy - Inches(sh) / 2, Inches(sw),
               Inches(sh), fill, line, TODO_LINE_PT)
        x += Inches(sw + LEGEND_GAP_IN)
        tb = _textbox(slide, f"Legend Label {n}", x, y, Inches(tw), Inches(LEGEND_H), label,
                      LEGEND_STYLE, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        tb.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        x += Inches(tw + LEGEND_GAP_OUT)


def _overflows(date, title, desc, st, seg_w, seg_h, text_w, text_h):
    """标题放不进箭头段、日期超过一行、说明超出下方区域时为 True。
    箭头段在中线处的可用宽度 = 段宽 − 左侧凹口深（chevron 尖角默认为段高的一半）。"""
    seg_in, h_in = Emu(seg_w).inches, Emu(seg_h).inches
    title_pt = SEG_TITLE_PT_DOING if st == DOING else SEG_TITLE_PT
    date_pt = (DATE_STYLE_DOING if st == DOING else DATE_STYLE)[0]
    tw = Emu(text_w).inches
    desc_pt = text_lines(desc, TEXT_STYLE[0], tw) * TEXT_STYLE[0] * LINE_SPACING
    return (text_width_in(title, title_pt) > seg_in - h_in / 2
            or text_width_in(date, date_pt) > tw
            or desc_pt > Emu(text_h).pt)


def render(prs, page, ctx):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    fill_titles(slide, page)
    left, top, width, slot_h = take_slot(slide, SLOT_IDX)   # 空槽位不留在页面上

    items = page.fields["items"]
    n = len(items)
    pitch = width / n
    seg_w = pitch - Inches(SEG_GAP)
    cy = top + Inches(SEG_CY)
    _legend(slide, left + width, top + Inches(LEGEND_TOP))
    for i, (date, title, desc, st, key) in enumerate(items, start=1):
        x = left + (i - 1) * pitch
        h = Inches(SEG_H_DOING if st == DOING else SEG_H)
        fill, line, ink = SEG_STYLE[st]
        kind = MSO_SHAPE.PENTAGON if i == 1 else MSO_SHAPE.CHEVRON
        seg = _shape(slide, kind, f"Phase Segment {i}", x, cy - h / 2, seg_w, h,
                     fill, line, TODO_LINE_PT)
        shift = Inches(MID_SHIFT_FIRST if i == 1 else MID_SHIFT_CHEVRON)
        mid = x + seg_w / 2 + shift
        _center_text(seg.text_frame, title,
                     SEG_TITLE_PT_DOING if st == DOING else SEG_TITLE_PT, True, ink)
        tf = seg.text_frame             # 用不对称内边距把标题挪到视觉中线
        tf.margin_left, tf.margin_right = (2 * shift, 0) if shift > 0 else (0, -2 * shift)

        tw = pitch - 2 * Inches(TEXT_INSET)
        _textbox(slide, f"Phase Date {i}", mid - tw / 2, cy + Inches(DATE_TOP), tw,
                 Inches(DATE_H), date, DATE_STYLE_DOING if st == DOING else DATE_STYLE,
                 anchor=MSO_ANCHOR.BOTTOM)

        my = cy + Inches(MARK_CY)
        d = Inches(KEY_D if key else PLAIN_D)
        _stem(slide, f"Phase Stem {i}", mid, cy + h / 2, my - d / 2)
        if key:
            _shape(slide, MSO_SHAPE.DIAMOND, f"Phase Marker {i}", mid - d / 2, my - d / 2,
                   d, d, INK)
            _textbox(slide, f"Phase Key Label {i}", mid - tw / 2, my + Inches(LABEL_TOP), tw,
                     Inches(LABEL_H), KEY_LABEL, LABEL_STYLE)
        else:
            _shape(slide, MSO_SHAPE.OVAL, f"Phase Marker {i}", mid - d / 2, my - d / 2,
                   d, d, LIGHT)
        text_top = my + Inches(TEXT_TOP)
        text_h = top + slot_h - text_top
        _textbox(slide, f"Phase Text {i}", mid - tw / 2, text_top, tw, text_h, desc,
                 TEXT_STYLE)
        if _overflows(date, title, desc, st, seg_w, h, tw, text_h):
            warn(f"页「{page.title}」阶段条第 {i} 条文字估计溢出, 建议精简日期 / 标题 / 说明")

    add_page_chrome(slide, page, ctx)
    return slide

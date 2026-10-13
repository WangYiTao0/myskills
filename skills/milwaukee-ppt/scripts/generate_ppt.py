#!/usr/bin/env python3
"""md 文稿 → Milwaukee 格式 pptx。

用法:
    python3 generate_ppt.py <文稿.md>              # 输出名按 frontmatter 推导
    python3 generate_ppt.py <文稿.md> -o out.pptx  # 指定输出

frontmatter 两种模式：有 title 为通用模式（closing 决定收尾页）；
有 month + author、没有 title 为月报模式（行为同旧月报脚本）。

三层：mdparse（md → 文档模型）、layouts/（每种版式一个模块，注册表分派）、
本文件（装配：封面 → 各页 → 收尾，并清理模板样例页）。约定见 scripts/README.md。
"""
import argparse
import datetime
import os
import re
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

import layouts
import mdparse
from mdparse import DeckError, is_monthly

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "..", "assets", "MS_Template.pptx")

MONTHS = ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

LAYOUT_TITLE_SLIDE = 1


def delete_template_slides(prs):
    """删除模板自带样例页，同时断开关系，避免保存时部件名冲突。"""
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        rid = sld_id.get(qn("r:id"))
        if rid:
            prs.part.drop_rel(rid)
        sld_id_lst.remove(sld_id)


def add_cover(prs, meta):
    """月报模式：「<Mon>. YYYY Working Summary」+「date   author」。
    通用模式：title +（subtitle 一行、「date   author」一行，空的省略）。"""
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE_SLIDE])
    date = meta.get("date", "").strip()
    byline = f"{date}   {meta.get('author', '').strip()}".strip()
    if is_monthly(meta):
        year, mon = meta["month"].split("-")
        slide.placeholders[0].text = f"{MONTHS[int(mon)]}. {year} Working Summary"
        slide.placeholders[1].text = byline
        return
    slide.placeholders[0].text = meta["title"]
    lines = [ln for ln in (meta.get("subtitle", "").strip(), byline) if ln]
    slide.placeholders[1].text = "\n".join(lines)


def add_closing(prs, text):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE_SLIDE])
    slide.placeholders[0].text = text
    slide.placeholders[1].text = ""


def closing_text(meta):
    """月报模式固定 Thank You；通用模式由 closing 决定，不写就不加收尾页。"""
    if is_monthly(meta):
        return "Thank You"
    return meta.get("closing", "").strip()


def deck_year(meta):
    """Tag 不含年份时补的年份：月报取 month；通用模式取 date 里的 4 位年份，否则今年。"""
    if is_monthly(meta):
        return meta["month"].split("-")[0]
    m = re.search(r"\b(\d{4})\b", meta.get("date", ""))
    return m.group(1) if m else str(datetime.date.today().year)


def build(doc):
    prs = Presentation(TEMPLATE)
    delete_template_slides(prs)
    add_cover(prs, doc.meta)
    ctx = layouts.Context(source_dir=doc.source_dir, year=deck_year(doc.meta))
    for page in doc.pages:
        layouts.render(prs, page, ctx)
    closing = closing_text(doc.meta)
    if closing:
        add_closing(prs, closing)
    return prs


def default_output(doc):
    meta = doc.meta
    if is_monthly(meta):
        year, mon = meta["month"].split("-")
        name = (f"Monthly Working Summary_{meta['author']}_"
                f"{MONTHS[int(mon)]}_{year}.pptx")
    else:
        name = re.sub(r'[\\/:*?"<>|]', "_", meta["title"]).strip() + ".pptx"
    return os.path.join(doc.source_dir, name)


def main(argv=None):
    ap = argparse.ArgumentParser(description="md 文稿 → Milwaukee 格式 pptx")
    ap.add_argument("md", help="输入 markdown 文稿")
    ap.add_argument("-o", "--output", help="输出 pptx 路径")
    args = ap.parse_args(argv)
    try:
        doc = mdparse.parse(args.md)
        prs = build(doc)
        out = args.output or default_output(doc)
        if os.path.isdir(out):
            raise DeckError(f"输出路径是目录, 请指定文件名: {out}")
        try:
            prs.save(out)
        except OSError as e:
            raise DeckError(f"保存失败 ({out}): {e}")
    except DeckError as e:
        print(f"错误: {e}", file=sys.stderr)
        return 1
    print(f"已生成: {out}  ({len(doc.pages)} 页正文)")
    return 0


if __name__ == "__main__":
    sys.exit(main())

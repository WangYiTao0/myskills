#!/usr/bin/env python3
"""生成后审核：机械检查 md 文稿和 pptx，列出要改的地方。

用法:
    python3 review_deck.py deck.md [输出.pptx]

每条问题一行：「[规则] 位置：问题 → 怎么改」。没有问题时最后一行打印 REVIEW OK，
有问题时最后一行是「REVIEW: N 条问题」。只报告、不改文件；审完退出码总是 0，
文稿本身读不了（找不到、写错）时为 1。和生成器共用解析器与版式注册表。
给了 pptx 时另查真伪（是不是生成器按当前文稿做的）和几何（出界、压 banner / 页脚、
文字溢出、重叠、错位；只用 python-pptx 和计算，不依赖 LibreOffice）。
需要判断的项（内容是否贴题、页序是否讲得通）留给审核清单，由模型逐条核对。
"""
import math
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import mdparse  # noqa: E402
from layouts.common import (BRAND_RED, conclusion_text, is_number, next_text,  # noqa: E402
                            text_width_in)  # 和生成器同一套结论 / Next 文字、数字判定、字宽估算
from layouts.process import desc_wrap_problem  # noqa: E402  和生成器的折行警告同一判断
from mdparse import DeckError  # noqa: E402

# 项目汇报模式下，章节 → 不适合放的版式（目标、范围要写目标和边界，不放流程和改前改后；应用场景不放改前改后）
MISFIT = {"TARGET": {"process", "before-after"}, "SCOPE": {"process", "before-after"},
          "USE CASES": {"before-after"}}
VOICE = ["赋能", "打通", "闭环", "抓手", "拉通", "组合拳", "打造", "未来可期", "令人兴奋", "激动人心",
         "惊艳", "重磅", "大幅", "显著", "大量", "彻底解决", "完美", "全面", "零问题", "革命性",
         "散乱", "低效", "容易出错", "混乱"]
MAX_PAGES = 10
NAV_LAYOUTS = {"section", "agenda"}     # 导航页（隔断页、目录页）：不放结论红条、不计页数、不要求 PDCA
EMU_PER_IN = 914400
BANNER_BOTTOM = EMU_PER_IN           # 模板顶部红 banner 占 0–1.0"（EMU）
FOOTER_TOP = int(7.25 * EMU_PER_IN)  # 母版页脚文本框上沿
REDRAW = "不要自己写代码画幻灯片；改 md 文稿后用 generate_ppt.py 重新生成"


def where(page):
    return f"页「{page.title}」"


def english(title):
    return title.split("/")[-1].strip().upper() if "/" in title else ""


def chinese(title):
    return title.split("/")[0].strip()


def page_text(page):
    return "\n".join([page.title, page.action_title, *page.lines, page.next])


def image_refs(page):
    """本页引用的图片（evidence 的图、image-text 的图、before-after 两侧的图）。"""
    f = page.fields
    if page.layout == "evidence":
        return list(f["images"])
    if page.layout == "image-text":
        return [f["image"]]
    if page.layout == "before-after":
        return [f[side]["image"] for side in ("before", "after") if f[side]["image"]]
    return []


def check_images(doc, page):
    """配图检查：证据页必须有图（image-text 没图解析时已报错）；本地图片文件要在。"""
    issues = []
    if page.layout == "evidence" and not page.fields["images"]:
        issues.append(("缺配图", where(page), "证据页（evidence，不写 Layout: 时的默认版式）没有放图",
                       "加一行 ![](图片路径) 放截图或现场照片；只有文字就改用 process / stats 等版式"))
    for ref in image_refs(page):
        if ref.startswith(("http://", "https://")):
            continue
        path = ref if os.path.isabs(ref) else os.path.join(doc.source_dir, ref)
        if not os.path.exists(path):
            issues.append(("图片不存在", where(page), f"引用的图片 {ref} 找不到",
                           "检查路径（相对文稿所在目录），把图片放到位后重新生成"))
    return issues


# 目录式副标题：太短（不到 6 字，如「项目目标」「试点总结」），或只是主标题中文部分
# 加两三个字（如「项目目标介绍」）。6 字以上的短结论（「两项指标均达标」）不报。
MIN_SUBTITLE = 6


def is_directory_subtitle(title, topic):
    return (len(title) < MIN_SUBTITLE
            or (bool(topic) and topic in title and len(title) - len(topic) <= 2))


def check_agenda_sections(doc):
    """有目录页时，每张隔断页的章节名都要在目录里出现。"""
    names = {item["name"] for p in doc.pages if p.layout == "agenda" for item in p.fields["items"]}
    if not names:
        return []
    return [("目录和隔断页不一致", where(p), f"隔断页章节名「{p.fields['name']}」不在目录页里",
             "隔断页和目录页的章节名写成一样，或在目录页补上这一节")
            for p in doc.pages if p.layout == "section" and p.fields["name"] not in names]


def check_lead(page):
    """方向 C（通用 / 项目汇报模式）下的 lead：不能和副标题一样；对比页只能有一句 lead。"""
    issues = []
    text = conclusion_text(page)
    if text and text == page.action_title.strip():
        issues.append(("lead 和副标题相同", where(page), f"结论行（lead）和副标题都是「{text}」",
                       "副标题写本页结论，lead 重申结论并补上数据支撑（SKILL.md §4），两句不要一样"))
    if page.layout == "before-after" and page.fields["lead"]:
        issues.append(("对比页两个 lead", where(page),
                       f"第一个 ### 前写了「{page.fields['lead']}」，结论行又是一句 lead；"
                       f"前一句不会画出来",
                       "合并成一句结论写在最后一个小节之后，删掉第一个 ### 前的文字"))
    return issues


def check_process_wrap(page):
    """方向 C 的流程页：说明折行难看（超过两行或末行只剩 ≤2 字），和生成器的警告同一判断。"""
    if page.layout != "process":
        return []
    steps = page.fields["steps"]
    return [("流程说明折行", where(page), f"步骤「{name}」的说明「{desc}」{problem}",
             "改短或改写说明，让它一行写完或两行大致等长")
            for name, desc in steps
            for problem in [desc and desc_wrap_problem(desc, len(steps))] if problem]


def review_md(doc):
    """内容检查：只看文稿（版式字段由版式模块解析好）。返回 [(规则, 位置, 问题, 怎么改)]。"""
    issues = []
    pages = doc.pages
    content = [p for p in pages if p.layout not in NAV_LAYOUTS]
    if len(content) > MAX_PAGES:
        issues.append(("页数过多", "整份文稿", f"{len(content)} 页（不含目录页、隔断页）",
                       f"合并重复的页，控制在 {MAX_PAGES} 页以内"))
    for prev, page in zip(pages, pages[1:]):
        if prev.layout == page.layout and page.layout != "section":   # 连着两张隔断页各管一节，不算版式重复
            issues.append(("相邻两页版式相同", where(page),
                           f"和上一页「{prev.title}」都是 {page.layout}，连着看像同一页",
                           "合并两页，或给其中一页换更贴合内容的版式"))
    seen_steps = []
    lead = mdparse.uses_lead(doc.meta)
    for page in pages:
        title = page.action_title.strip()
        if is_directory_subtitle(title, chinese(page.title)):
            issues.append(("目录式副标题", where(page), f"副标题「{title}」没有给出结论",
                           "改成一句结论，如「试点 8 周，检验单全部电子化」"))
        if not page.conclusion and page.layout not in NAV_LAYOUTS:
            issues.append(("缺结论行", where(page),
                           "没有普通文字结论行，" + ("标题下没有 lead" if lead else "页底没有红条"),
                           "在版式字段后加一行本页结论，最好带数字"))
        if lead:
            issues += check_lead(page) + check_process_wrap(page)
        if page.layout == "process":
            names = {name for name, _ in page.fields["steps"]}
            for other_title, other in seen_steps:
                if len(names & other) * 2 >= min(len(names), len(other)):
                    issues.append(("内容重复", where(page),
                                   f"流程步骤和页「{other_title}」大半相同",
                                   "删掉重复的一页，或改写成不同的内容"))
            seen_steps.append((page.title, names))
        if page.layout in ("stats", "stats-side"):
            for item in page.fields["items"]:
                if "[数据缺失]" in item["value"]:
                    continue
                if not is_number(item["value"]):
                    issues.append(("数字卡不是数字", where(page),
                                   f"卡片「{item['label']}」的数字栏是「{item['value']}」",
                                   "数字栏只放数字；文字内容改写进说明、结论要点或别的版式"))
        if page.layout in ("timeline", "timeline-phase"):
            keys = (page.fields["keys"] if page.layout == "timeline"
                    else [item[4] for item in page.fields["items"]])
            if len(keys) >= 3 and sum(keys) * 2 > len(keys):
                issues.append(("关键标记过多", where(page),
                               f"{sum(keys)}/{len(keys)} 个节点标了 !",
                               "只给 1–2 个真正关键的节点加 !"))
        issues += check_images(doc, page)
        text = page_text(page)
        if "[数据缺失]" in text:
            issues.append(("数据缺失", where(page), "有 [数据缺失] 占位",
                           "交付时告诉用户要补哪项数据；没有数据的指标不单独做数字卡"))
        words = [w for w in VOICE if w in text]
        if words:
            issues.append(("用词", where(page), "用了 " + "、".join(words),
                           "换成具体做了什么、结果是多少"))
    issues += check_agenda_sections(doc)
    if mdparse.is_project(doc.meta):
        issues += review_project(doc)
    return issues


# 项目汇报模式（report: project）的 8 节：SKILL.md §1 的固定英文名、顺序和 PDCA。
# 英文名不在表里的页是子节点，归到它前面最近的一节。
CHAPTERS = ("BACKGROUND", "TARGET", "SCOPE", "USE CASES", "MILESTONE", "OUTCOMES", "TEAM", "LESSONS")
CHAPTER_PDCA = {"BACKGROUND": "P", "TARGET": "P", "SCOPE": "P", "USE CASES": "P",
                "MILESTONE": "D", "OUTCOMES": "DC", "TEAM": "", "LESSONS": "A"}
# 06 节 [DO/CHECK]：讲做了什么的页标 D、讲结果验证的页标 C，一页只放一个（§2）


def chapter_of(page):
    """页属于 §1 的哪一节（0–7）；英文名不是 8 节之一返回 None。"""
    name = english(page.title)
    return CHAPTERS.index(name) if name in CHAPTERS else None


def review_project(doc):
    """项目汇报模式检查（§1 章节顺序、§2 PDCA）。返回 [(规则, 位置, 问题, 怎么改)]。"""
    issues = []
    order = " → ".join(CHAPTERS)
    latest = None                       # 已出现的最靠后的一节
    current = None                      # 当前页所属的节（子节点沿用前一节）
    for page in doc.pages:
        n = chapter_of(page)
        if n is not None:
            if latest is not None and n < latest:
                issues.append(("章节顺序", where(page),
                               f"{CHAPTERS[n]} 排在 {CHAPTERS[latest]} 后面",
                               f"按 SKILL.md §1 的顺序排：{order}"))
            latest = n if latest is None else max(latest, n)
            current = CHAPTERS[n]
        issues += _check_pdca(page, current)
        if page.layout in MISFIT.get(english(page.title), set()):
            issues.append(("章节和版式不配", where(page),
                           f"{english(page.title)} 章节用了 {page.layout}",
                           "目标、范围写目标值和边界（stats / stats-side），流程和改前改后放到对应章节"))
    for page in doc.pages:
        if page.layout != "agenda":
            continue
        listed = [e for e in (item["english"].upper() for item in page.fields["items"])
                  if e in CHAPTERS]
        if listed != sorted(listed, key=CHAPTERS.index):
            issues.append(("章节顺序", where(page),
                           f"目录页的顺序是 {' → '.join(listed)}",
                           f"按 SKILL.md §1 的顺序排：{order}"))
    missing = [c for c in CHAPTERS if not any(english(p.title) == c for p in doc.pages)]
    if missing:
        issues.append(("缺章节", "整份文稿", "少了 " + "、".join(missing),
                       f"项目汇报 8 节不可去节，按 {order} 补齐（主标题写「中文 / 英文固定名」）"))
    return issues


def _check_pdca(page, chapter):
    """一页的 PDCA 标注：团队页不带；导航页可不写；其余页要写且和所属节对应。"""
    letters = "/".join(page.pdca)
    if page.layout == "team" or chapter == "TEAM":
        if page.pdca:
            return [("团队页带 PDCA", where(page), f"项目团队页标了 PDCA: {letters}",
                     "删掉这一行 PDCA:；第 07 节项目团队不带 PDCA 标签")]
        return []
    if not page.pdca:
        if page.layout in NAV_LAYOUTS:
            return []
        return [("缺 PDCA", where(page), "项目汇报模式的这一页没有标 PDCA",
                 "加一行 PDCA: <字母>，按所属节填（§1 表的 PDCA 列）")]
    allowed = CHAPTER_PDCA.get(chapter)
    if allowed is None:                 # 还没进入 8 节中的任何一节，对不上表
        return []
    want = " 或 ".join(allowed)
    if len(page.pdca) > 1:
        return [("PDCA 和章节不符", where(page), f"一页标了 {letters} 多个阶段",
                 f"一页只放一个图标；{chapter} 节标 {want}，做了什么和结果验证分两页")]
    if page.pdca[0] not in allowed:
        return [("PDCA 和章节不符", where(page), f"{chapter} 节标了 {letters}",
                 f"按 SKILL.md §1 改成 PDCA: {want}")]
    return []


def review_pptx(doc, path):
    """真伪检查：pptx 是不是生成器按当前文稿做的——页数对得上，标题填在模板标题框
    （idx 0 / 10）里，banner 上没有另画的文本框。返回 [(规则, 位置, 问题, 怎么改)]。"""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(path)
    slides = list(prs.slides)
    width, height = prs.slide_width, prs.slide_height
    closing = 1 if doc.meta.get("closing", "").strip() or mdparse.is_monthly(doc.meta) else 0
    want = 1 + len(doc.pages) + closing
    issues = []
    if len(slides) != want:
        issues.append(("页数不符", "pptx", f"有 {len(slides)} 页，文稿应生成 {want} 页",
                       "用 generate_ppt.py 按当前文稿重新生成"))
    for n, (slide, page) in enumerate(zip(slides[1:], doc.pages), start=2):
        place = f"pptx 第 {n} 页（{where(page)}）"
        titles = {ph.placeholder_format.idx: ph.text_frame.text.strip()
                  for ph in slide.placeholders if ph.has_text_frame}
        if (titles.get(0) != page.title
                or (page.action_title and titles.get(10) != page.action_title)):
            issues.append(("不是生成脚本做的", place,
                           "主标题、副标题没有填进模板标题框（idx 0 / 10）", REDRAW))
        for sh in slide.shapes:
            if (sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX and sh.top is not None
                    and sh.top < BANNER_BOTTOM):
                issues.append(("不是生成脚本做的", place,
                               f"banner 上另画了文本框「{sh.name}」", REDRAW))
        issues += check_drawn(slide, page, place, mdparse.uses_lead(doc.meta))
        if mdparse.uses_lead(doc.meta):
            issues += check_red_area(slide, page, place, width, height)
    return issues


# 红色面积配额（SKILL.md §6.3 面积表）：通用 / 项目模式正文页 banner 以外红色 ≤ 5%；月报豁免。
RED_QUOTA = 0.05
STATUS_LAYOUTS = {"timeline", "timeline-phase"}   # §6.3 例外：时间线状态色（红描边）不计入
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def _red(el):
    return el is not None and el.get("val", "").upper() == BRAND_RED


def red_area(sh, page):
    """一个形状 banner 以外的品牌红面积（EMU²）：填充或描边是品牌红时按外框算，直线按
    长 × 线宽算；时间线类版式里只有描边是红的（状态色）不算。红色文字不计入。"""
    sp = sh._element.find(f"{_P}spPr")
    if sp is None or None in (sh.left, sh.top, sh.width, sh.height):
        return 0
    ln = sp.find(f"{_A}ln")
    fill_red = _red(sp.find(f"{_A}solidFill/{_A}srgbClr"))
    line_red = ln is not None and _red(ln.find(f"{_A}solidFill/{_A}srgbClr"))
    if sh._element.tag == f"{_P}cxnSp":
        if not line_red:
            return 0
        top, bottom = sh.top, sh.top + sh.height
        if bottom <= top:                               # 横线：整条在 banner 上就不算
            below = 1.0 if top >= BANNER_BOTTOM else 0.0
        else:                                           # 斜线 / 竖线：按 banner 以下的比例算
            below = max(0, bottom - max(top, BANNER_BOTTOM)) / (bottom - top)
        return math.hypot(sh.width, sh.height) * int(ln.get("w", 12700)) * below
    if not fill_red and not (line_red and page.layout not in STATUS_LAYOUTS):
        return 0
    return sh.width * max(0, sh.top + sh.height - max(sh.top, BANNER_BOTTOM))


def check_red_area(slide, page, place, width, height):
    """banner 以外的品牌红面积超过 §6.3 配额（方向 C：通用 / 项目模式）。"""
    share = sum(red_area(sh, page) for sh in slide.shapes) / (width * height)
    if share <= RED_QUOTA:
        return []
    return [("红色面积超配额", place,
             f"banner 以外红色约占页面 {share:.1%}，超过 SKILL.md §6.3 配额 {RED_QUOTA:.0%}",
             "红色只留 lead 竖线、1 个重点数字和 `!` 标记；去掉另加的红色块 / 红线，"
             "改 md 后用 generate_ppt.py 重新生成")]


def _texts(slide, name):
    """名为 name 的形状里的各段文字（去掉首尾空白、跳过空段）。"""
    return [p.text.strip() for sh in slide.shapes if sh.name == name and sh.has_text_frame
            for p in sh.text_frame.paragraphs if p.text.strip()]


def check_drawn(slide, page, place, lead):
    """结论和 Next 画进了 pptx：方向 C 下结论在 Lead、Next 在 Next Line（文字一致）；
    月报模式下两句都要出现在底部红条 Caption Bar 里。"""
    if page.layout in NAV_LAYOUTS:
        return []
    issues = []
    for rule, want, shape, label in (
            ("结论没画出来", conclusion_text(page), "Lead" if lead else "Caption Bar", "结论行"),
            ("Next 没画出来", next_text(page), "Next Line" if lead else "Caption Bar", "Next")):
        if not want:
            continue
        got = _texts(slide, shape)
        if (got != [want]) if lead else (want not in got):
            seen = "没有这个形状" if not got else f"里面是「{' / '.join(got)}」"
            issues.append((rule, place, f"文稿的{label}「{want}」没画进「{shape}」（{seen}）",
                           REDRAW))
    return issues


def _inches(emu):
    return f'{emu / EMU_PER_IN:.2f}"'


def check_bounds(slide, width, height, place):
    """出界 / 压 banner / 压页脚。banner 里的模板标题框（idx 0 / 10）不算。"""
    issues = []
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.idx in (0, 10):
            continue
        if None in (sh.left, sh.top, sh.width, sh.height):
            continue
        right, bottom = sh.left + sh.width, sh.top + sh.height
        name = f"形状「{sh.name}」"
        if sh.left < 0 or sh.top < 0 or right > width or bottom > height:
            issues.append(("出界", place, f"{name}超出页面",
                           "把它挪回页面内，或缩小尺寸"))
        elif sh.top < BANNER_BOTTOM:
            issues.append(("压 banner", place,
                           f"{name}上沿在 {_inches(sh.top)}，压住顶部 banner（0–1.0\"）",
                           "下移到 1.0\" 以下；标题只写进模板标题框"))
        elif bottom > FOOTER_TOP:
            issues.append(("压页脚", place,
                           f"{name}下沿在 {_inches(bottom)}，压住页脚（7.25\" 以下）",
                           "上移或缩小，下沿不超过 7.25\""))
    return issues


# 重叠容差：两个形状横、纵两个方向都压进对方超过它才算重叠。生成器坐标由浮点取整，
# 相邻形状（细线连到标记、卡片贴边）可能互相压进几个 EMU；0.02"（96 dpi 下约 2 像素）
# 远小于肉眼能看出的叠压，又能吸收取整误差。
OVERLAP_TOL = int(0.02 * EMU_PER_IN)
# 有意叠放的豁免（按形状名，和 references/layouts.md「审核的几何豁免」一致）：
OVERLAP_EXEMPT = {
    "Caption Bar",   # 底部红条叠在最上层，盖住内容区下沿属设计
    "Month Tag",     # 月份标识框 1.12–1.42"，下沿伸进内容区 0.12"；字只占框的右上角
}
OVERLAP_PAIRS = {("Current Ring", "Milestone"),  # 进行中节点的外圈套在节点上
                 }
# 两两豁免（名字完全相同才算）：方向 C 的 lead 和它左侧的红竖线、页底 Next 行和它上方的细线
OVERLAP_TWINS = {frozenset({"Lead", "Lead Accent"}), frozenset({"Next Rule", "Next Line"})}
_COMPONENT = re.compile(r"^(\S+) .*?(\d+)$")   # 「Process Card 2」→ 组件 (Process, 2)


def _component(name):
    m = _COMPONENT.match(name)
    return (m.group(1), m.group(2)) if m else None


def _overlap_exempt(a, b):
    if a.is_placeholder or b.is_placeholder or {a.name, b.name} & OVERLAP_EXEMPT:
        return True
    if frozenset({a.name, b.name}) in OVERLAP_TWINS:
        return True
    ca, cb = _component(a.name), _component(b.name)
    if ca is not None and ca == cb:
        return True
    return any((a.name == x and b.name.split()[0] == y) or (b.name == x and a.name.split()[0] == y)
               for x, y in OVERLAP_PAIRS)


def _placed(slide):
    return [sh for sh in slide.shapes if None not in (sh.left, sh.top, sh.width, sh.height)]


def check_overlap(slide, place):
    """形状矩形互压（豁免见 OVERLAP_EXEMPT / OVERLAP_PAIRS 和同一组件）。"""
    issues = []
    shapes = _placed(slide)
    for i, a in enumerate(shapes):
        for b in shapes[i + 1:]:
            dx = min(a.left + a.width, b.left + b.width) - max(a.left, b.left)
            dy = min(a.top + a.height, b.top + b.height) - max(a.top, b.top)
            if dx > OVERLAP_TOL and dy > OVERLAP_TOL and not _overlap_exempt(a, b):
                issues.append(("重叠", place,
                               f"形状「{a.name}」和「{b.name}」互相压住"
                               f"（横向 {_inches(dx)}、纵向 {_inches(dy)}）",
                               "挪开其中一个或缩小尺寸；文字放不下就精简文稿后重新生成"))
    return issues


# 错位阈值：同类形状（名字去掉编号后相同，如 Process Card 1–4）的上沿 / 中线 / 下沿
# （同一行）或左沿 / 中线 / 右沿（同一列）都没对齐、最近的差值小于它时报告。
# 5 像素（96 dpi，约 0.05"）：手工微调挪歪常见 1–4 像素，肉眼难发现、投影时明显；
# 生成器有意的错开（放大的关键节点、进行中的节点 / 阶段段）中线仍对齐，不会报。
PX_EMU = EMU_PER_IN // 96
MISALIGN_TOL = 5 * PX_EMU
# 差值不到半个像素视为对齐：生成器按浮点算坐标再取整，同类形状间实测有约 0.2 像素的误差，
# 屏幕上看不出来；挪动 1 像素及以上都会被报出。
ALIGN_SNAP = PX_EMU // 2
_NUMBERED = re.compile(r"^(.*\S) \d+$")


def _edges(sh, axis):
    if axis == "row":
        return sh.top, sh.top + sh.height // 2, sh.top + sh.height
    return sh.left, sh.left + sh.width // 2, sh.left + sh.width


def _near_miss(a, b, axis):
    """a、b 在 axis 方向三条边都没对齐、最近的差值小于 MISALIGN_TOL → 差值，否则 None。"""
    diffs = [abs(x - y) for x, y in zip(_edges(a, axis), _edges(b, axis))]
    if min(diffs) <= ALIGN_SNAP or min(diffs) >= MISALIGN_TOL:
        return None
    return min(diffs)


def _aligned(a, b, axis):
    return min(abs(x - y) for x, y in zip(_edges(a, axis), _edges(b, axis))) <= ALIGN_SNAP


def check_alignment(slide, place):
    """同类带编号形状差几像素没对齐。每个挪歪的形状报一次，和对齐得多的那个比。"""
    roles = {}
    for sh in _placed(slide):
        m = _NUMBERED.match(sh.name)
        if m and not sh.is_placeholder:
            roles.setdefault(m.group(1), []).append(sh)
    issues = []
    for group in roles.values():
        for axis, edge in (("row", "上沿"), ("col", "左沿")):
            score = {id(s): sum(_aligned(s, t, axis) for t in group if t is not s) for s in group}
            for s in group:
                for t in group:
                    diff = None if t is s else _near_miss(s, t, axis)
                    if diff is None or score[id(s)] > score[id(t)]:
                        continue
                    want = t.top if axis == "row" else t.left
                    issues.append(("错位", place,
                                   f"形状「{s.name}」和「{t.name}」差 {diff / PX_EMU:.1f} 像素没对齐",
                                   f"把「{s.name}」的{edge}对齐到 {_inches(want)}（和「{t.name}」一致）"))
                    break
    return issues


# 文字溢出：行高按字号 × 1.2 估算，和 process / stats / stats-side / timeline / timeline-phase /
# before-after 估算溢出警告的行高同一口径；section / agenda / team / table 排框高用更宽松的
# common.LINE_SPACING（1.25），它们的框按 1.2 估不会报溢出。字宽用生成器的 text_width_in。估算高度超出文本框可用高度（去掉内边距）OVERFLOW_TOL 以上才报：
# 正常稿里最紧的大数字框估算高度恰好等于框高，0.05" 吸收估算与取整误差。
LINE_HEIGHT = 1.2
OVERFLOW_TOL = int(0.05 * EMU_PER_IN)
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _list_style(sh):
    """文本框自带 lstStyle 第一级（红条的字号、行距、缩进写在这里）；没有返回 None。"""
    return sh.text_frame._txBody.find(f"{_A}lstStyle/{_A}lvl1pPr")


def _para_metrics(para, lvl1):
    """一段的 (各 run 的 (文字, 磅)、行距倍数、段前磅、左缩进 EMU)；字号读不到返回 None。"""
    default_pt = None
    spacing, indent = 1.0, 0
    if lvl1 is not None:
        rpr = lvl1.find(f"{_A}defRPr")
        if rpr is not None and rpr.get("sz"):
            default_pt = int(rpr.get("sz")) / 100
        pct = lvl1.find(f"{_A}lnSpc/{_A}spcPct")
        if pct is not None:
            spacing = int(pct.get("val")) / 100000
        indent = int(lvl1.get("marL", 0))
    runs = []
    for run in para.runs:
        pt = run.font.size.pt if run.font.size is not None else default_pt
        if pt is None:
            return None
        runs.append((run.text, pt))
    if isinstance(para.line_spacing, float):
        spacing = para.line_spacing
    before = para.space_before.pt if para.space_before is not None else 0
    return runs, spacing, before, indent


def _text_height_in(sh):
    """估算文字占的高度（英寸）；字号继承模板（占位框）读不到时返回 None，不估。"""
    tf = sh.text_frame
    width = sh.width - tf.margin_left - tf.margin_right
    lvl1 = _list_style(sh)
    total = 0.0
    for para in tf.paragraphs:
        m = _para_metrics(para, lvl1)
        if m is None:
            return None
        runs, spacing, before, indent = m
        if not runs:
            continue
        pt = max(p for _, p in runs)
        line_w = (width - indent) / EMU_PER_IN
        text_w = sum(text_width_in(t, p) for t, p in runs)
        lines = 1 if tf.word_wrap is False else max(1, -(-text_w // line_w))
        total += lines * pt * LINE_HEIGHT * spacing / 72 + before / 72
    return total


def check_overflow(slide, place):
    """文字估计溢出文本框（自己写了字号的文本框和形状；占位框按模板排，不估）。"""
    issues = []
    for sh in _placed(slide):
        if not sh.has_text_frame or not sh.text_frame.text.strip() or sh.is_placeholder:
            continue
        need = _text_height_in(sh)
        if need is None:
            continue
        tf = sh.text_frame
        room = (sh.height - tf.margin_top - tf.margin_bottom) / EMU_PER_IN
        if (need - room) * EMU_PER_IN > OVERFLOW_TOL:
            issues.append(("文字溢出", place,
                           f"形状「{sh.name}」的文字估计要 {need:.2f}\"高，框只有 {room:.2f}\"",
                           "精简文稿里对应的文字后重新生成；不要缩字号或拉大框挤进去"))
    return issues


def _cell_height_in(cell, width_emu):
    """单元格文字估算高度（英寸，含上下内边距）；字号读不到返回 None。"""
    tf = cell.text_frame
    width = (width_emu - cell.margin_left - cell.margin_right) / EMU_PER_IN
    total = (cell.margin_top + cell.margin_bottom) / EMU_PER_IN
    for para in tf.paragraphs:
        m = _para_metrics(para, None)
        if m is None:
            return None
        runs, spacing, before, _ = m
        if not runs:
            continue
        pt = max(p for _, p in runs)
        lines = max(1, -(-sum(text_width_in(t, p) for t, p in runs) // width))
        total += lines * pt * LINE_HEIGHT * spacing / 72 + before / 72
    return total


def check_table_overflow(slide, place):
    """原生表格（table 版式的 Data Table）：某行文字估计高过行高。PowerPoint 会把行撑高、
    表格往下长，压到红条或页脚；按和文本框同一口径估算，超出 OVERFLOW_TOL 才报。"""
    issues = []
    for sh in _placed(slide):
        if not getattr(sh, "has_table", False):
            continue
        tbl = sh.table
        for r, row in enumerate(tbl.rows):
            need = [_cell_height_in(tbl.cell(r, c), col.width)
                    for c, col in enumerate(tbl.columns)]
            need = max((n for n in need if n is not None), default=0)
            room = row.height / EMU_PER_IN
            if (need - room) * EMU_PER_IN > OVERFLOW_TOL:
                issues.append(("文字溢出", place,
                               f"形状「{sh.name}」第 {r + 1} 行的文字估计要 {need:.2f}\"高，"
                               f"行高只有 {room:.2f}\"",
                               "精简文稿里对应的单元格后重新生成；不要缩字号或拉高行挤进去"))
    return issues


def review_geometry(doc, path):
    """几何检查（只用 python-pptx 和计算）：逐页查正文页。返回 [(规则, 位置, 问题, 怎么改)]。"""
    from pptx import Presentation
    prs = Presentation(path)
    slides = list(prs.slides)
    issues = []
    for n, (slide, page) in enumerate(zip(slides[1:], doc.pages), start=2):
        place = f"pptx 第 {n} 页（{where(page)}）"
        issues += check_bounds(slide, prs.slide_width, prs.slide_height, place)
        issues += check_overlap(slide, place)
        issues += check_alignment(slide, place)
        issues += check_overflow(slide, place)
        issues += check_table_overflow(slide, place)
    return issues


def main(argv):
    if len(argv) not in (1, 2):
        print("用法: python3 review_deck.py deck.md [输出.pptx]", file=sys.stderr)
        return 1
    try:
        doc = mdparse.parse(argv[0])      # 解析时已按注册表把各页交给版式模块
    except DeckError as e:
        print(f"错误: {e}", file=sys.stderr)
        return 1
    issues = review_md(doc)
    if len(argv) == 2:
        try:
            issues += review_pptx(doc, argv[1])
        except Exception as e:   # 文件不存在、不是 pptx：python-pptx 抛的异常种类不一
            print(f"错误: 读不了 pptx {argv[1]}: {e}", file=sys.stderr)
            return 1
        issues += review_geometry(doc, argv[1])
    for rule, place, problem, fix in issues:
        print(f"[{rule}] {place}：{problem} → {fix}")
    print(f"REVIEW: {len(issues)} 条问题" if issues else "REVIEW OK")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))

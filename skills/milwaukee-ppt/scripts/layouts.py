"""版式渲染层：每种版式一个渲染函数，签名统一为 fn(prs, page, ctx) -> slide。"""
import hashlib
import os
import re
import sys
import urllib.request
from dataclasses import dataclass

from xml.sax.saxutils import escape

from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from mdparse import DeckError

LAYOUT_BODY = 0          # 模板 layout 0「Blank Slide」：idx 0 主标题、idx 10 action title

# 证据页正文区几何（英寸）：banner 占顶部 1.0"，底部留给红条
IMG_LEFT, IMG_TOP, IMG_W, IMG_H = 0.56, 1.3, 12.2, 4.3
IMG_GAP = 0.3
MAX_IMAGES = 3

# 底部红条：月报模板固定样式，逐项复刻，不要改。
# 品牌红 DB011C（与 banner 同色）+ 0.75pt 黑描边，白字 ❖（Wingdings）14pt，行距 150%，垂直居中。
# 底边固定；文字多时向上加高，但不高过图片区底缘。
CAPTION_BOTTOM = 5379048 + 1097952  # EMU，底边位置固定
CAPTION_MIN_CY = 1097952            # 原版高度 = 最小高度
CAPTION_MAX_TOP = 5121840           # 不高过正文图片区底缘
CAPTION_LINE_EMU = 266700           # 每行 14pt × 150% 行距 = 21pt
CAPTION_INSETS_EMU = 45720 * 2      # tIns + bIns
CAPTION_CHARS_PER_LINE = 105        # 14pt Arial 在 12.3in 可用宽度下的保守估计

CAPTION_SP_XML = """<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <p:nvSpPr>
    <p:cNvPr id="{id}" name="Caption Bar"/>
    <p:cNvSpPr txBox="1"><a:spLocks noChangeArrowheads="1"/></p:cNvSpPr>
    <p:nvPr/>
  </p:nvSpPr>
  <p:spPr bwMode="auto">
    <a:xfrm><a:off x="155575" y="{y}"/><a:ext cx="11874874" cy="{cy}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
    <a:solidFill><a:srgbClr val="DB011C"/></a:solidFill>
    <a:ln w="9525" cap="flat" cmpd="sng" algn="ctr">
      <a:solidFill><a:schemeClr val="tx1"/></a:solidFill>
      <a:prstDash val="solid"/><a:round/>
      <a:headEnd type="none" w="med" len="med"/><a:tailEnd type="none" w="med" len="med"/>
    </a:ln>
    <a:effectLst/>
  </p:spPr>
  <p:txBody>
    <a:bodyPr vert="horz" wrap="square" lIns="91440" tIns="45720" rIns="91440" bIns="45720" numCol="1" rtlCol="0" anchor="ctr" anchorCtr="0" compatLnSpc="1"><a:prstTxWarp prst="textNoShape"><a:avLst/></a:prstTxWarp></a:bodyPr>
    <a:lstStyle>
      <a:defPPr><a:defRPr lang="en-US"/></a:defPPr>
      <a:lvl1pPr marL="342900" indent="-342900" eaLnBrk="0" hangingPunct="0">
        <a:lnSpc><a:spcPct val="150000"/></a:lnSpc>
        <a:buFont typeface="Wingdings" panose="05000000000000000000" pitchFamily="2" charset="2"/>
        <a:buChar char="v"/>
        <a:defRPr sz="1400" b="0">
          <a:solidFill><a:schemeClr val="bg1"/></a:solidFill>
          <a:ea typeface="微软雅黑" panose="020B0503020204020204" pitchFamily="34" charset="-122"/>
          <a:cs typeface="Arial" panose="020B0604020202020204" pitchFamily="34" charset="0"/>
        </a:defRPr>
      </a:lvl1pPr>
    </a:lstStyle>
  </p:txBody>
</p:sp>"""

CAPTION_P_XML = ('<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                 '<a:r><a:rPr lang="en-US" altLang="zh-CN" dirty="0"/>'
                 '<a:t>{text}</a:t></a:r></a:p>')


@dataclass
class Context:
    source_dir: str      # 文稿所在目录，相对图片路径以它为基准
    year: str            # Tag 不含年份时补上的年份


def _fill_titles(slide, page):
    slide.placeholders[0].text = page.title
    if page.action_title:
        slide.placeholders[10].text = page.action_title


def warn(msg):
    print(f"警告: {msg}", file=sys.stderr)


def url_cache_dir():
    """URL 图片的下载缓存目录（PPT 必须内嵌图片字节；缓存后可离线重复生成）。"""
    return os.environ.get("MILWAUKEE_PPT_CACHE") or os.path.join(
        os.path.expanduser("~"), ".cache", "milwaukee-ppt", "url-images")


def load_image(ref, ctx, page_title):
    """图片引用 → 本地路径。支持相对路径（相对文稿）、绝对路径、http(s) URL。"""
    if ref.startswith(("http://", "https://")):
        cache_dir = url_cache_dir()
        os.makedirs(cache_dir, exist_ok=True)
        ext = os.path.splitext(ref.split("?")[0])[1] or ".png"
        cache = os.path.join(cache_dir, hashlib.md5(ref.encode()).hexdigest() + ext)
        if not os.path.exists(cache):
            try:
                req = urllib.request.Request(ref, headers={"User-Agent": "Mozilla/5.0"})
                with urllib.request.urlopen(req, timeout=30) as resp:
                    data = resp.read()
                with open(cache, "wb") as fh:
                    fh.write(data)
            except Exception as e:                      # noqa: BLE001
                raise DeckError(f"页「{page_title}」下载图片失败: {ref}\n  {e}")
        return cache
    p = ref if os.path.isabs(ref) else os.path.join(ctx.source_dir, ref)
    if not os.path.exists(p):
        raise DeckError(f"页「{page_title}」引用的图片不存在: {ref}")
    return p


def _cell_widths(files, usable):
    """多图宽度分配：宽高比差异大时按比例分（宽图多占），否则等分。

    需要 Pillow；没装就等分。下限防最窄图被挤没。"""
    n = len(files)
    widths = [usable / n] * n
    if n < 2:
        return widths
    try:
        from PIL import Image
    except ImportError:
        return widths
    aspects = []
    for p in files:
        with Image.open(p) as im:
            aspects.append(im.width / im.height)
    if max(aspects) / min(aspects) <= 1.5:
        return widths
    floor = 0.28 if n == 2 else 0.22
    shares = [max(a / sum(aspects), floor) for a in aspects]
    total = sum(shares)
    return [usable * s / total for s in shares]


def place_images(slide, refs, ctx, page_title):
    """1–3 张图：各自保持比例、在 cell 内居中；整行按最高图垂直居中。"""
    if not refs:
        return
    if len(refs) > MAX_IMAGES:
        warn(f"页「{page_title}」有 {len(refs)} 张图, 仅使用前 {MAX_IMAGES} 张")
        refs = refs[:MAX_IMAGES]
    files = [load_image(r, ctx, page_title) for r in refs]
    widths = _cell_widths(files, IMG_W - IMG_GAP * (len(files) - 1))

    placed = []
    cell_left = IMG_LEFT
    for p, cell_w in zip(files, widths):
        pic = slide.shapes.add_picture(p, Inches(cell_left), Inches(IMG_TOP))
        scale = min(Inches(cell_w) / pic.width, Inches(IMG_H) / pic.height)
        pic.width = int(pic.width * scale)
        pic.height = int(pic.height * scale)
        pic.left = Inches(cell_left) + int((Inches(cell_w) - pic.width) / 2)
        placed.append(pic)
        cell_left += cell_w + IMG_GAP
    row_h = max(pic.height for pic in placed)
    row_top = Inches(IMG_TOP) + int((Inches(IMG_H) - row_h) / 2)
    for pic in placed:
        pic.top = row_top


def _next_shape_id(slide):
    ids = [int(e.get("id")) for e in slide.shapes._spTree.iter(qn("p:cNvPr"))
           if e.get("id")]
    return (max(ids) + 1) if ids else 100


def _caption_geometry(lines, page_title):
    """底边固定，按估算行数向上加高（CJK 宽字符按 2 计，额外半行余量防贴框）。"""
    rows = 0
    for text in lines:
        w = sum(2 if ord(ch) > 0x2E7F else 1 for ch in text)
        rows += max(1, -(-w // CAPTION_CHARS_PER_LINE))
    cy = max(CAPTION_MIN_CY,
             rows * CAPTION_LINE_EMU + CAPTION_INSETS_EMU + CAPTION_LINE_EMU // 2)
    y = CAPTION_BOTTOM - cy
    if y < CAPTION_MAX_TOP:
        warn(f"页「{page_title}」红条文字过多 ({rows} 行), 已顶到图片区底缘, "
             f"建议精简结论/Next")
        y = CAPTION_MAX_TOP
        cy = CAPTION_BOTTOM - y
    return y, cy


def caption_lines(page):
    """页面的一句话结论 + 可选 Next，组成红条的各行。"""
    lines = []
    conclusion = " ".join(page.conclusion).strip()
    if conclusion:
        lines.append(conclusion)
    if page.next:
        sep = "：" if page.next_label == "下一步" else ": "
        lines.append(f"{page.next_label}{sep}{page.next}")
    return lines


def add_caption_bar(slide, page):
    """有结论或 Next 时画底部红条；样式全部由 CAPTION_SP_XML 的 lstStyle 继承。"""
    lines = caption_lines(page)
    if not lines:
        return None
    y, cy = _caption_geometry(lines, page.title)
    sp = parse_xml(CAPTION_SP_XML.format(id=_next_shape_id(slide), y=y, cy=cy))
    tx_body = sp.find(qn("p:txBody"))
    for text in lines:
        tx_body.append(parse_xml(CAPTION_P_XML.format(text=escape(text))))
    slide.shapes._spTree.append(sp)
    return sp


def add_month_tag(slide, tag, year):
    """正文区右上角极简月份标识（无框、灰字）；只在页面写了 Tag 时显示。

    Tag 不含 4 位年份时补上文稿年份，如 "Mar" → "Mar 2026"。"""
    if not tag:
        return None
    text = tag if re.search(r"\d{4}", tag) else f"{tag} {year}"
    w, h = Inches(1.7), Inches(0.3)
    tb = slide.shapes.add_textbox(Inches(IMG_LEFT + IMG_W) - w, Inches(1.12), w, h)
    tb.name = "Month Tag"
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Pt(0)
    tf.margin_top = tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = text
    r.font.size = Pt(11)
    r.font.bold = False
    r.font.name = "Arial"
    r.font.color.rgb = RGBColor(0x6E, 0x6E, 0x74)
    return tb


def render_evidence(prs, page, ctx):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_BODY])
    _fill_titles(slide, page)
    place_images(slide, page.fields.get("images", []), ctx, page.title)
    add_month_tag(slide, page.tag, ctx.year)
    add_caption_bar(slide, page)
    return slide


RENDERERS = {
    "evidence": render_evidence,
}


def render(prs, page, ctx):
    return RENDERERS[page.layout](prs, page, ctx)

#!/usr/bin/env python3
"""检查 pptx：按页列出每个形状的名字、类型、位置大小（英寸）和文字前约 30 字。

用法：python3 inspect_ppt.py deck.pptx
输出：每页一个标题行 `== 第 N 页 | layout: <名>`，其下每个形状一行
`名字 | 类型 | L= T= W= H= | 文字`。
"""
import sys
from pathlib import Path

EMU = 914400
TEXT_LEN = 30


def fail(msg):
    print(f"错误：{msg}", file=sys.stderr)
    sys.exit(1)


def inches(v):
    return "-" if v is None else f"{v / EMU:.2f}"


def shape_type(sh):
    try:
        t = sh.shape_type
    except Exception:
        t = None
    if t is None:
        return "PLACEHOLDER" if sh.is_placeholder else "UNKNOWN"
    return t.name if hasattr(t, "name") else str(t)


def shape_text(sh):
    if not getattr(sh, "has_text_frame", False):
        return ""
    return " ".join(sh.text_frame.text.split())[:TEXT_LEN]


def main(argv):
    if len(argv) != 2:
        fail("用法：inspect_ppt.py <deck.pptx>")
    path = Path(argv[1])
    if not path.is_file():
        fail(f"文件不存在：{path}")
    from pptx import Presentation
    try:
        prs = Presentation(str(path))
    except Exception as e:
        fail(f"不是有效的 pptx 文件：{path}（{type(e).__name__}）")
    for n, slide in enumerate(prs.slides, 1):
        print(f"== 第 {n} 页 | layout: {slide.slide_layout.name}")
        for sh in slide.shapes:
            print(f"{sh.name} | {shape_type(sh)} | L={inches(sh.left)} "
                  f"T={inches(sh.top)} W={inches(sh.width)} H={inches(sh.height)} "
                  f"| {shape_text(sh)}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))

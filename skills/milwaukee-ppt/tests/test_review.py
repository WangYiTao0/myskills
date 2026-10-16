"""审核脚本 review_deck.py：md + pptx 进、问题清单出。

运行：python3 skills/milwaukee-ppt/tests/test_review.py
做法：先用生成器按文稿生成 pptx，再跑审核脚本，只看它的输出与退出码。
负样本里的 pptx 用 python-pptx 把生成器的正常输出改坏。文稿内容全部虚构。
从 copilot-agent 分支 commit 30678fb 的 tests/test_review.py 移植。
"""
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx.util import Inches

from deckcheck import (GENERATOR, SKILL, broken, check, finish, generate, generated, lines_with,
                       placeholder, review)

DEMO = SKILL / "examples" / "demo" / "deck.md"

GOOD = """\
---
title: 来料检验记录电子化 项目汇报
---

# 项目背景 / BACKGROUND
## 查一批次的检验结论要翻纸质单
Layout: before-after

### Before: 纸质检验单
手写纸质检验单，下班前录入 Excel，每周五手工汇总周报

### After: 电子检验单
平板填写并拍照上传，提交即入库，每周一自动生成周报

检验单提交即入库，供应商质量周报改为自动生成

# 新流程 / PROCESS
## 检验记录三步完成，提交即入库
Layout: process
1. 平板填写 | 现场填写电子检验单
2. 拍照上传 | 不良品照片随单上传
3. 自动出周报 | 每周一生成质量周报
A 车间检验员 5 月起全部使用新流程

# 关键数据 / KEY RESULTS
## 试点 8 周，检验单全部电子化
Layout: stats
- 电子检验单 | 2,360 张 | 试点 8 周全部电子化
- !单张录入用时 | 约 1.5 分钟 | 改前约 4 分钟
口径：A 车间试点 8 周
单张录入用时从约 4 分钟降到约 1.5 分钟

# 项目节点 / MILESTONE
## 3 月立项，6 月完成 8 周试点复盘
Layout: timeline
- Mar | 立项 | 确定 A 车间试点 | 完成
- Apr | 试用上线 | 电子检验单开放试用 | 完成
- !May | 全员切换 | A 车间检验员全部切换 | 完成
- Jun | 试点复盘 | 完成 8 周试点复盘 | 完成
试点复盘 6 月完成
"""


def test_good_deck_passes():
    r = review(*generated(GOOD))
    check("合格文稿：退出码 0", r.returncode == 0, r.stderr)
    check("合格文稿：最后一行 REVIEW OK", r.stdout.strip().endswith("REVIEW OK"), r.stdout)


BAD = """\
---
title: 来料检验记录电子化 项目汇报
report: project
---

# 项目目标 / TARGET
## 项目目标
Layout: process
1. 平板填写 | 现场填写电子检验单
2. 拍照上传 | 不良品照片随单上传
3. 自动出周报 | 每周一生成质量周报

# 应用场景 / USE CASES
## 检验流程全面赋能，效率大幅提升
Layout: process
1. 平板填写 | 检验员填写
2. 拍照上传 | 照片上传
3. 提交入库 | 直接入库
检验流程已切换

# 关键数据 / KEY RESULTS
## 试点 8 周，检验单全部电子化
Layout: stats
- 电子检验单 | 2,360 张 | 试点期间全部电子化
- 周报生成方式 | 自动生成 | 替代每周五人工汇总
- 不良率 | [数据缺失] | 仍在核对
检验单全部电子化

# 项目节点 / MILESTONE
## 3 月立项，6 月完成 8 周试点复盘
Layout: timeline
- !Mar | 立项 | 确定试点 | 完成
- !Apr | 试用上线 | 开放试用 | 完成
- !May | 全员切换 | 全部切换 | 完成
试点复盘 6 月完成
"""


def test_bad_deck_flags():
    r = review(*generated(BAD))
    out = r.stdout
    check("问题文稿：退出码仍为 0（审核只报告）", r.returncode == 0, r.stderr)
    check("问题文稿：不报 REVIEW OK，最后一行给出问题条数",
          "REVIEW OK" not in out and out.strip().splitlines()[-1].startswith("REVIEW: "), out)
    expect = {
        "相邻两页版式相同": "页「应用场景 / USE CASES」",
        "目录式副标题": "页「项目目标 / TARGET」",
        "缺结论行": "页「项目目标 / TARGET」",
        "章节和版式不配": "页「项目目标 / TARGET」",
        "内容重复": "页「应用场景 / USE CASES」",
        "数字卡不是数字": "自动生成",
        "数据缺失": "[数据缺失]",
        "用词": "赋能",
        "关键标记过多": "页「项目节点 / MILESTONE」",
    }
    for rule, needle in expect.items():
        lines = lines_with(out, rule)
        check(f"报出「{rule}」并指明位置", any(needle in line for line in lines), lines or out)
    check("每条问题都是「[规则] 位置：问题 → 怎么改」",
          all("：" in line and " → " in line
              for line in out.splitlines() if line.startswith("[")), out)
    check("合格的页不误报：关键数据页没有目录式副标题",
          not any("KEY RESULTS" in line for line in lines_with(out, "目录式副标题")), out)


def test_too_many_pages():
    pages = "\n".join(
        f"# 第 {n} 页 / PAGE {n}\n## 第 {n} 个车间完成切换，用时 {n} 周\n"
        + ("![](img/a.png)\n" if n % 2 else
           "Layout: process\n1. 准备 | 培训检验员\n2. 切换 | 改用平板\n3. 复盘 | 汇总问题\n")
        + f"第 {n} 个车间切换完成\n"
        for n in range(1, 12))
    r = generate("---\ntitle: 多车间推广汇报\n---\n" + pages, images={"img/a.png": (1600, 900)})
    rr = review(r.workdir / "deck.md", r.path)
    check("11 页正文：报出页数过多", any("11 页" in line for line in lines_with(rr.stdout, "页数过多")),
          rr.stdout)


def test_number_prefixes_same_as_generator():
    """生成器认作「前缀 + 大数字」的写法（～ ＞ ＜ ± 等），审核不能报「数字卡不是数字」。"""
    values = ("～5 天", "＞95%", "＜3 台", "±0.5 mm", "约 1.5 分钟", "~2 周", "≥99%")
    pages = "".join(
        f"# 关键数据 {k} / KEY RESULTS {k}\n## 各项指标都达标\nLayout: stats\n"
        + "\n".join(f"- 指标{n} | {v} | 说明{n}" for n, v in enumerate(group, 1))
        + "\n各项指标都达标\n"
        for k, group in enumerate((values[:4], values[4:]), 1))
    md = "---\ntitle: 数字前缀\n---\n" + pages
    md_path, pptx = generated(md)
    out = review(md_path, pptx).stdout
    check("～ ＞ ＜ ± 等前缀的数字：不报「数字卡不是数字」", not lines_with(out, "数字卡不是数字"), out)


def hand_drawn_title(prs):
    """模仿“模型自己画页面”：第 3 页（流程页）清空标题占位框，在 banner 上另画文本框写标题。"""
    slide = prs.slides[2]
    placeholder(slide, 0).text_frame.text = ""
    box = slide.shapes.add_textbox(Inches(5.1), Inches(0.1), Inches(8), Inches(0.4))
    box.text_frame.text = "新流程 / PROCESS"


def textbox_on_banner(prs):
    """标题仍在占位框里，但第 2 页 banner 上另画了一个文本框。"""
    prs.slides[1].shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(3), Inches(0.4))


def drop_last_slide(prs):
    ids = prs.slides._sldIdLst
    ids.remove(ids[-1])


def test_pptx_authenticity():
    md, pptx = generated(GOOD)
    good = review(md, pptx)
    check("生成器做的 pptx：不报不是生成脚本做的", not lines_with(good.stdout, "不是生成脚本做的"),
          good.stdout)

    r = review(md, broken(pptx, hand_drawn_title))
    lines = lines_with(r.stdout, "不是生成脚本做的")
    check("标题画在普通文本框里：报出不是生成脚本做的，指明第 3 页",
          any("第 3 页" in line for line in lines), r.stdout)
    check("标题画在普通文本框里：退出码仍为 0", r.returncode == 0, r.stderr)

    r = review(md, broken(pptx, textbox_on_banner))
    check("banner 上另画文本框：报出不是生成脚本做的，指明第 2 页",
          any("第 2 页" in line for line in lines_with(r.stdout, "不是生成脚本做的")), r.stdout)

    r = review(md, broken(pptx, drop_last_slide))
    check("pptx 少一页：报出页数不符", any("文稿应生成 5 页" in line
                                       for line in lines_with(r.stdout, "页数不符")), r.stdout)


def test_demo_deck_passes():
    out = Path(tempfile.mkdtemp(prefix="mppt-review-")) / "demo.pptx"
    g = subprocess.run([sys.executable, str(GENERATOR), str(DEMO), "-o", str(out)],
                       capture_output=True, text=True)
    check("示例文稿：生成成功", g.returncode == 0, g.stderr)
    r = review(DEMO, out)
    check("示例文稿：审核只打印 REVIEW OK", r.stdout.strip() == "REVIEW OK", r.stdout)


def test_unreadable_deck():
    r = review(Path(tempfile.gettempdir()) / "no-such-deck-xyz.md")
    check("文稿找不到：退出码 1，错误写到 stderr", r.returncode == 1 and "错误" in r.stderr,
          (r.returncode, r.stderr))
    bad = generate("---\ntitle: 写错版式\n---\n# 某页\nLayout: no-such\n")
    r = review(bad.workdir / "deck.md")
    check("文稿写错：退出码 1，错误含页标题", r.returncode == 1 and "某页" in r.stderr,
          (r.returncode, r.stderr))
    md, _ = generated(GOOD)
    r = review(md, md.parent / "no-such.pptx")
    check("pptx 读不了：退出码 1，错误写到 stderr", r.returncode == 1 and "pptx" in r.stderr,
          (r.returncode, r.stderr))


def test_review_does_not_modify_files():
    md, pptx = generated(BAD)
    before = (md.read_bytes(), pptx.read_bytes())
    review(md, pptx)
    check("审核不改文稿和 pptx", (md.read_bytes(), pptx.read_bytes()) == before)


if __name__ == "__main__":
    test_good_deck_passes()
    test_bad_deck_flags()
    test_too_many_pages()
    test_number_prefixes_same_as_generator()
    test_pptx_authenticity()
    test_demo_deck_passes()
    test_unreadable_deck()
    test_review_does_not_modify_files()
    finish()

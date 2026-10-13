"""测试专用版式 `_probe`：把共用部件的结果画成文本框，供 test_common_parts.py 检查。

只在 deckcheck.generate(..., probe=True) 时注册（见 deckcheck._PROBE_RUNNER），
不进正式注册表，文稿里写 `Layout: _probe` 平时会报“版式名未知”。

正文行写法：
    num: [!]数值        → Probe Number N：common.add_number 写一段（! 为重点，红色）
    status: a, b, c    → Probe Status N：common.parse_statuses 的结果，
                         每个状态一个 run（颜色取 STATUS_COLORS）；全部不写时文字为 none
    status!: a, b, c   → 同上，required=True
    其他行              → Probe Line N：「缩进宽度|原文」；
                         common.sub_item 认出的子列表项另画 Probe Sub N
"""
import re

from pptx.util import Inches

from layouts.common import (BRAND_RED, STATUS_COLORS, add_number, add_run, fill_titles,
                            parse_statuses, sub_item)

NAME = "_probe"
LAYOUT_INDEX = 5
INK = "1A1A1A"

_STATUS = re.compile(r"^status(!?):(.*)$")


def parse(page):
    page.fields = {"lines": list(zip(page.lines, page.indents)),
                   "subs": [s for s in (sub_item(page, i) for i in range(len(page.lines)))
                            if s is not None],
                   "statuses": []}
    for line in page.lines:
        m = _STATUS.match(line)
        if m:
            raws = [x.strip() for x in m.group(2).split(",")]
            page.fields["statuses"].append(
                parse_statuses(page, raws, required=bool(m.group(1))))


def _box(slide, name, i):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(1.2 + 0.1 * i), Inches(6), Inches(0.1))
    tb.name = f"{name} {i}"
    return tb.text_frame.paragraphs[0]


def render(prs, page, ctx):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_INDEX])
    fill_titles(slide, page)
    n_num = n_line = 0
    for line, indent in page.fields["lines"]:
        if not line:
            continue
        if line.startswith("num:"):
            n_num += 1
            value = line[4:].strip()
            key = value.startswith("!")
            add_number(_box(slide, "Probe Number", n_num), value.lstrip("!"), 60,
                       BRAND_RED if key else INK)
        elif not _STATUS.match(line):
            n_line += 1
            add_run(_box(slide, "Probe Line", n_line), f"{indent}|{line}", 12, False, INK)
    for i, text in enumerate(page.fields["subs"], start=1):
        add_run(_box(slide, "Probe Sub", i), text, 12, False, INK)
    for i, sts in enumerate(page.fields["statuses"], start=1):
        p = _box(slide, "Probe Status", i)
        if sts is None:
            add_run(p, "none", 12, False, INK)
            continue
        for j, st in enumerate(sts):
            add_run(p, ("" if j == 0 else " ") + st, 12, False, STATUS_COLORS[st])
    return slide

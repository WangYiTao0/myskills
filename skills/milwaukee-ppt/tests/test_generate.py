"""milwaukee-ppt 生成器测试：只走“md 进、pptx 出”这一个切入点。

纯 Python 脚本，不用 pytest：python3 skills/milwaukee-ppt/tests/test_generate.py
测试图片用 Pillow 现生成纯色图；文稿内容全部虚构。
"""
import re
import subprocess
import sys
import tempfile
import textwrap
import zipfile
from pathlib import Path

from pptx import Presentation

SKILL = Path(__file__).resolve().parent.parent
GENERATOR = SKILL / "scripts" / "generate_ppt.py"
TEMPLATE = SKILL / "assets" / "MS_Template.pptx"
EMU = 914400
BANNER_BOTTOM = EMU          # 母版红 banner 占顶部 1.0"
FOOTER_TOP = int(7.25 * EMU)  # 母版页脚文本框上沿

FAILS = []


def check(name, ok, detail=""):
    print(("✓ " if ok else "✗ ") + name + ("" if ok else f"  —— {detail}"))
    if not ok:
        FAILS.append(name)


def make_image(path, size, color):
    from PIL import Image
    Image.new("RGB", size, color).save(path)


def run(md_text, workdir, out_name="out.pptx", use_o=True):
    """写文稿 → 调生成器 → 返回 (退出码, stdout, stderr, Presentation|None)。"""
    md = Path(workdir) / "deck.md"
    md.write_text(textwrap.dedent(md_text).lstrip(), encoding="utf-8")
    cmd = [sys.executable, str(GENERATOR), str(md)]
    out = Path(workdir) / out_name
    if use_o:
        cmd += ["-o", str(out)]
    r = subprocess.run(cmd, capture_output=True, text=True)
    prs = Presentation(str(out)) if use_o and out.exists() else None
    return r.returncode, r.stdout, r.stderr, prs


# 虚构的旧格式月报（与原月报 md 写法一致，不改一字）
MONTHLY_MD = """
---
month: 2026-03
author: Alex.Sample
date: 9 Apr
---

<!-- 注释不会进入 PPT -->

# Sample Dashboard Program
## Build a Reusable Dashboard Model
![](img/wide.png)
Standard dashboard workflow drafted and shared with two pilot teams.
Next: Roll the workflow out to the remaining pilot teams.

# Mock Tooling Review
## First Draft Review
![](img/wide.png)
![](img/tall.png)
Tag: Mar
First-draft review finished with sample stakeholders.
下一步：整理反馈并进入正式开发。

# Demo Knowledge Page
![](img/wide.png)
![](img/tall.png)
![](img/square.png)
Tag: Dec 2025
Demo pages live; usage tracking enabled.
"""


def texts(slide):
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]


def test_monthly_regression():
    with tempfile.TemporaryDirectory() as tmp:
        (Path(tmp) / "img").mkdir()
        make_image(Path(tmp) / "img" / "wide.png", (1600, 900), (120, 120, 120))
        make_image(Path(tmp) / "img" / "tall.png", (600, 1200), (180, 180, 180))
        make_image(Path(tmp) / "img" / "square.png", (800, 800), (60, 60, 60))
        rc, _, err, prs = run(MONTHLY_MD, tmp)
        check("月报：退出码 0", rc == 0, err)
        if prs is None:
            check("月报：生成了 pptx", False, err)
            return
        slides = list(prs.slides)
        check("月报：页数 = 封面 + 3 页 + Thank You", len(slides) == 5,
              f"实际 {len(slides)}")

        cover = slides[0]
        check("封面用 Title Slide", cover.slide_layout.name == "Title Slide",
              cover.slide_layout.name)
        check("封面主标题", cover.placeholders[0].text == "Mar. 2026 Working Summary",
              repr(cover.placeholders[0].text))
        check("封面副标题 = date author",
              cover.placeholders[1].text == "9 Apr   Alex.Sample",
              repr(cover.placeholders[1].text))

        thanks = slides[-1]
        check("Thank You 页", thanks.slide_layout.name == "Title Slide"
              and thanks.placeholders[0].text == "Thank You",
              repr(texts(thanks)))

        body = slides[1:-1]
        check("证据页都用 layout 0", all(
            s.slide_layout == prs.slide_layouts[0] for s in body),
            [s.slide_layout.name for s in body])
        check("主标题在 idx 0（顺序与文稿一致）",
              [s.placeholders[0].text for s in body] ==
              ["Sample Dashboard Program", "Mock Tooling Review",
               "Demo Knowledge Page"],
              [s.placeholders[0].text for s in body])
        ph10 = [s.placeholders[10].text if 10 in [p.placeholder_format.idx
                for p in s.placeholders] else None for s in body]
        check("action title 在 idx 10",
              ph10[:2] == ["Build a Reusable Dashboard Model", "First Draft Review"],
              ph10)
        check("没写 ## 的页 idx 10 不填字", not ph10[2], repr(ph10[2]))
        pics = [len([sh for sh in s.shapes if sh.shape_type == 13]) for s in body]
        check("每页图数 = 文稿里的图数 (1/2/3)", pics == [1, 2, 3], pics)

        bars = [caption_bar(s) for s in body]
        check("每页都有底部红条", all(bars), [b is not None for b in bars])
        if all(bars):
            check("红条文字 = 结论 + Next（保留中文「下一步」前缀）",
                  [paragraphs(b) for b in bars] == [
                      ["Standard dashboard workflow drafted and shared with two pilot teams.",
                       "Next: Roll the workflow out to the remaining pilot teams."],
                      ["First-draft review finished with sample stakeholders.",
                       "下一步：整理反馈并进入正式开发。"],
                      ["Demo pages live; usage tracking enabled."]],
                  [paragraphs(b) for b in bars])
            for i, b in enumerate(bars):
                problems = caption_style_problems(b)
                check(f"红条样式与月报一致（第 {i + 1} 页）", not problems, problems)

        tags = [tag_texts(s) for s in body]
        check("Tag：不写不显示；写月份补年份；含年份原样",
              tags == [[], ["Mar 2026"], ["Dec 2025"]], tags)
        badge = [sh for sh in body[1].shapes
                 if sh.has_text_frame and sh.text_frame.text == "Mar 2026"]
        if badge:
            run_ = badge[0].text_frame.paragraphs[0].runs[0]
            check("Tag 是正文区右上的灰字 11pt",
                  str(run_.font.color.rgb) == "6E6E74" and run_.font.size.pt == 11
                  and badge[0].top >= EMU and badge[0].left > 9 * EMU,
                  (run_.font.color.rgb, run_.font.size, badge[0].top, badge[0].left))

        bad = []
        for n, s in enumerate(body, start=2):
            for sh in s.shapes:
                if sh.is_placeholder and sh.placeholder_format.idx in (0, 10):
                    continue          # banner 里的标题占位框
                if not (sh.left >= 0 and sh.left + sh.width <= prs.slide_width
                        and sh.top >= BANNER_BOTTOM
                        and sh.top + sh.height <= FOOTER_TOP):
                    bad.append(f"第{n}页 {sh.name}")
        check("所有形状在页面内、不压 banner 和页脚", not bad, bad)
        sample_texts = {t for sl in Presentation(str(TEMPLATE)).slides
                        for t in texts(sl) if t.strip()}
        leftover = [t for sl in slides for t in texts(sl) if t in sample_texts]
        check("没有残留模板样例页（样例文字不出现）", not leftover, leftover)
        with zipfile.ZipFile(Path(tmp) / "out.pptx") as z:
            parts = [n for n in z.namelist()
                     if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
        check("包里的 slide 部件数 = 页数（样例页部件已删）",
              len(parts) == len(slides), parts)

def test_errors():
    with tempfile.TemporaryDirectory() as tmp:
        rc, _, err, prs = run("""
            ---
            month: 2026-03
            ---
            # Only Page
            Some conclusion.
            """, tmp)
        check("缺 author：非零退出", rc != 0, rc)
        check("缺 author：错误信息点名 author", "author" in err, err)
        check("缺 author：不生成文件", prs is None)

        rc, _, err, prs = run("""
            ---
            month: 2026-03
            author: Alex.Sample
            ---
            # Fine Page
            Fine.

            # Broken Image Page
            ![](img/missing.png)
            Conclusion.
            """, tmp)
        check("图片不存在：非零退出", rc != 0, rc)
        check("图片不存在：错误信息含页标题和路径",
              "Broken Image Page" in err and "img/missing.png" in err, err)
        check("图片不存在：不生成半坏的文件", prs is None)


def test_default_output_name():
    with tempfile.TemporaryDirectory() as tmp:
        rc, out, err, _ = run("""
            ---
            month: 2026-11
            author: Alex.Sample
            ---
            # Page
            Conclusion.
            """, tmp, use_o=False)
        expected = Path(tmp) / "Monthly Working Summary_Alex.Sample_Nov_2026.pptx"
        check("不写 -o：按 frontmatter 推导文件名（放在文稿旁）",
              rc == 0 and expected.exists(),
              (rc, err, sorted(p.name for p in Path(tmp).iterdir())))


def test_caption_grows_upward():
    long = " ".join(["This sample sentence is deliberately long."] * 8)
    with tempfile.TemporaryDirectory() as tmp:
        rc, _, err, prs = run(f"""
            ---
            month: 2026-03
            author: Alex.Sample
            ---
            # Short
            Short conclusion.

            # Long
            {long}
            Next: {long}
            """, tmp)
        if prs is None:
            check("长红条：生成成功", False, err)
            return
        short_bar, long_bar = (caption_bar(s) for s in list(prs.slides)[1:3])
        check("短红条 = 原版高度", short_bar.height == CAPTION_MIN_CY, short_bar.height)
        check("长红条向上加高、底边不动",
              long_bar.height > CAPTION_MIN_CY
              and long_bar.top + long_bar.height == CAPTION_BOTTOM,
              (long_bar.top, long_bar.height))


# 原月报红条：底边固定在 5379048 + 1097952 EMU，最小高度 1097952，左 155575，宽 11874874
CAPTION_BOTTOM = 6477000
CAPTION_MIN_CY = 1097952
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def caption_bar(slide):
    for sh in slide.shapes:
        if sh.name == "Caption Bar":
            return sh
    return None


def tag_texts(slide):
    """非占位框、非红条的文本框文字（证据页上只有月份标识会是这种）。"""
    return [sh.text_frame.text for sh in slide.shapes
            if sh.has_text_frame and not sh.is_placeholder
            and sh.name != "Caption Bar" and sh.text_frame.text]


def paragraphs(shape):
    return [p.text for p in shape.text_frame.paragraphs if p.text]


def caption_style_problems(sh):
    el = sh._element
    out = []
    if sh.left != 155575 or sh.width != 11874874:
        out.append(f"x/宽 {sh.left}/{sh.width}")
    if sh.top + sh.height != CAPTION_BOTTOM:
        out.append(f"底边 {sh.top + sh.height}")
    if sh.height < CAPTION_MIN_CY:
        out.append(f"高 {sh.height} 小于原版")
    fill = el.find(f".//{{http://schemas.openxmlformats.org/presentationml/2006/main}}spPr/{A}solidFill/{A}srgbClr")
    if fill is None or fill.get("val") != "DB011C":
        out.append("底色不是 DB011C")
    ln = el.find(f".//{A}ln")
    if ln is None or ln.get("w") != "9525" or ln.find(f"{A}solidFill/{A}schemeClr").get("val") != "tx1":
        out.append("描边不是 0.75pt 黑")
    lvl1 = el.find(f".//{A}lstStyle/{A}lvl1pPr")
    if lvl1 is None:
        return out + ["缺 lvl1pPr"]
    if lvl1.find(f"{A}lnSpc/{A}spcPct").get("val") != "150000":
        out.append("行距不是 150%")
    if lvl1.find(f"{A}buFont").get("typeface") != "Wingdings" or lvl1.find(f"{A}buChar").get("char") != "v":
        out.append("项目符号不是 Wingdings ❖")
    rpr = lvl1.find(f"{A}defRPr")
    if rpr.get("sz") != "1400" or rpr.find(f"{A}solidFill/{A}schemeClr").get("val") != "bg1":
        out.append("文字不是白色 14pt")
    if el.find(f".//{A}bodyPr").get("anchor") != "ctr":
        out.append("不是垂直居中")
    return out


def main():
    test_monthly_regression()
    test_errors()
    test_default_output_name()
    test_caption_grows_upward()
    print()
    if FAILS:
        print(f"{len(FAILS)} 项失败")
        sys.exit(1)
    print("全部通过")


if __name__ == "__main__":
    main()

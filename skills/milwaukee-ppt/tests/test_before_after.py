"""before-after 版式测试（票 #14）：只走“md 进、pptx 出”。

运行：python3 skills/milwaukee-ppt/tests/test_before_after.py
覆盖：槽位 11–16 的文字 / 图片、lead 与红条结论的切分、图在槽内等比居中、
缺小节等文稿错误、caption / lead / 短文的估算警告、红色纪律、红条不盖住 caption。
"""
from pptx.enum.shapes import MSO_SHAPE_TYPE

from deckcheck import (LAYOUT_INDEX, caption_bar, caption_style_problems, check,
                       finish, generate, out_of_bounds, paragraphs, placeholder,
                       red_elements,
                       tag_shape, warnings)

NAME = "before-after"
TITLE = "Mock Before / After Page"
MD = """
---
title: Mock Deck
date: 13 Oct 2026
---

# Mock Before / After Page
## Mock Action Title
Layout: before-after
Tag: Mar
### Before
Manual copy-paste between three mock spreadsheets.
Caption: Three manual hand-offs
### After
One mock form feeds every sheet.
Caption: One automated hand-off

Mock one-line conclusion.
Next: Mock next step.
"""
IMAGES = {"img/shot.png": (1600, 900)}


def body_slide(r):
    return r.slides[1] if len(r.slides) > 1 else None


def test_chrome():
    """所有版式通用：用模板对应 layout、标题进 idx 0/10、Tag 和底部红条、不越界。"""
    r = generate(MD, images=IMAGES)
    check("before-after：退出码 0", r.rc == 0, r.stderr)
    s = body_slide(r)
    if s is None:
        check("before-after：生成了正文页", False, r.stderr)
        return
    check("before-after：用模板 layout 3「Before / After」",
          s.slide_layout == r.prs.slide_layouts[LAYOUT_INDEX[NAME]],
          s.slide_layout.name)
    check("before-after：主标题在 idx 0、action title 在 idx 10",
          placeholder(s, 0).text == TITLE
          and placeholder(s, 10).text == "Mock Action Title",
          (placeholder(s, 0).text, placeholder(s, 10).text))
    bar = caption_bar(s)
    check("before-after：底部红条 = 结论 + Next",
          bar is not None and paragraphs(bar) ==
          ["Mock one-line conclusion.", "Next: Mock next step."],
          bar and paragraphs(bar))
    if bar is not None:
        problems = caption_style_problems(bar)
        check("before-after：红条样式与月报一致", not problems, problems)
    tag = tag_shape(s)
    check("before-after：Tag 补年份显示", tag is not None and tag.text_frame.text == "Mar 2026",
          tag and tag.text_frame.text)
    bad = out_of_bounds(r.prs, [s])
    check("before-after：形状在页面内、不压 banner 和页脚", not bad, bad)


def test_text_slots():
    """改前、改后都是短文：标题默认 BEFORE / AFTER，短文进 13/14，caption 进 15/16。"""
    r = generate(MD, images=IMAGES)
    s = body_slide(r)
    if s is None:
        check("before-after（文字）：生成了正文页", False, r.stderr)
        return
    got = {i: (placeholder(s, i).text if placeholder(s, i) is not None else None)
           for i in (11, 12, 13, 14, 15, 16)}
    check("before-after（文字）：槽位 11–16 内容正确", got == {
        11: "BEFORE", 12: "AFTER",
        13: "Manual copy-paste between three mock spreadsheets.",
        14: "One mock form feeds every sheet.",
        15: "Three manual hand-offs", 16: "One automated hand-off"}, got)
    check("before-after（文字）：没有警告", not warnings(r.stderr), r.stderr)


IMG_MD = """
---
title: Mock Deck
---

# Mock Image Page
Layout: before-after
### Before
![](img/wide.png)
Caption: Old mock screen
### After
![](img/tall.png)
Caption: New mock screen
"""
IMG_FILES = {"img/wide.png": (1600, 900), "img/tall.png": (600, 1200)}


def pictures(slide):
    return [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]


def test_image_slots():
    """改前、改后都是图：槽内文字清空，图在槽矩形内、比例不变、居中。"""
    r = generate(IMG_MD, images=IMG_FILES)
    check("before-after（图）：退出码 0", r.rc == 0, r.stderr)
    s = body_slide(r)
    if s is None:
        return
    pics = pictures(s)
    check("before-after（图）：两张原生图片（不是图片占位框）", len(pics) == 2,
          [sh.shape_type for sh in s.shapes])
    if len(pics) != 2:
        return
    for idx, pic, aspect in ((13, min(pics, key=lambda p: p.left), 16 / 9),
                             (14, max(pics, key=lambda p: p.left), 0.5)):
        slot = placeholder(s, idx)
        check(f"before-after（图）：槽 {idx} 文字已清空", slot.text == "", slot.text)
        inside = (pic.left >= slot.left and pic.top >= slot.top
                  and pic.left + pic.width <= slot.left + slot.width
                  and pic.top + pic.height <= slot.top + slot.height)
        check(f"before-after（图）：槽 {idx} 的图在槽矩形内", inside,
              (pic.left, pic.top, pic.width, pic.height,
               slot.left, slot.top, slot.width, slot.height))
        check(f"before-after（图）：槽 {idx} 的图比例不变",
              abs(pic.width / pic.height - aspect) < 0.01, pic.width / pic.height)
        cx = abs((pic.left + pic.width / 2) - (slot.left + slot.width / 2))
        cy = abs((pic.top + pic.height / 2) - (slot.top + slot.height / 2))
        check(f"before-after（图）：槽 {idx} 的图居中", cx < 2 and cy < 2, (cx, cy))
        fills = (pic.width >= slot.width - 2) or (pic.height >= slot.height - 2)
        check(f"before-after（图）：槽 {idx} 的图撑满一边", fills,
              (pic.width, pic.height, slot.width, slot.height))
    check("before-after（图）：caption 进 15/16",
          (placeholder(s, 15).text, placeholder(s, 16).text)
          == ("Old mock screen", "New mock screen"))


LEAD_MD = """
---
title: Mock Deck
---

# Mock Lead Page
Layout: before-after
Tag: Mar
Mock lead sentence about the change.
### Before
Mock old way.
Caption: Old
### After
![](img/wide.png)
Caption: New

Mock conclusion after the last section.
"""


def overlaps(a, b):
    return (a.left < b.left + b.width and b.left < a.left + a.width
            and a.top < b.top + b.height and b.top < a.top + a.height)


def test_lead_and_conclusion():
    """第一个小节之前是 lead（模板约定位置），最后一个小节之后是红条结论；文字 + 图混搭。"""
    r = generate(LEAD_MD, images=IMG_FILES)
    check("before-after（lead）：退出码 0", r.rc == 0, r.stderr)
    s = body_slide(r)
    if s is None:
        return
    leads = [sh for sh in s.shapes if sh.name == "Lead"]
    check("before-after（lead）：有一个 Lead 文本框，文字正确",
          len(leads) == 1 and leads[0].text_frame.text
          == "Mock lead sentence about the change.", [x.name for x in s.shapes])
    if leads:
        lead = leads[0]
        E = 914400
        check("before-after（lead）：位置在模板约定的 0.56 / 1.30，高 0.42",
              abs(lead.left - 0.56 * E) < 2000 and abs(lead.top - 1.30 * E) < 2000
              and abs(lead.height - 0.42 * E) < 2000,
              (lead.left / E, lead.top / E, lead.height / E))
        tag = tag_shape(s)
        check("before-after（lead）：Lead 不和 Tag 重叠",
              tag is not None and not overlaps(lead, tag))
    bar = caption_bar(s)
    check("before-after（lead）：最后小节之后的文字进红条",
          bar is not None and paragraphs(bar) == ["Mock conclusion after the last section."],
          bar and paragraphs(bar))
    check("before-after（lead）：改前是文字、改后是图",
          placeholder(s, 13).text == "Mock old way." and placeholder(s, 14).text == ""
          and len(pictures(s)) == 1)


def deck(body):
    return "---\ntitle: Mock Deck\n---\n\n# Mock Broken Page\nLayout: before-after\n" + body


def test_errors():
    """文稿错误：非零退出、不出文件，错误信息含页标题和原因。"""
    cases = [
        ("缺 After 小节", "### Before\nOld.\nCaption: a\n", "### After"),
        ("缺 Before 小节", "### After\nNew.\nCaption: b\n", "### Before"),
        ("一节里图和文字混放",
         "### Before\n![](img/wide.png)\nOld.\n### After\nNew.\n", "一张图或一段短文"),
        ("小节没有内容", "### Before\nCaption: a\n### After\nNew.\n", "没有内容"),
        ("未知小节名", "### Before\nOld.\n### During\nMid.\n### After\nNew.\n",
         "### During"),
        ("两个小节之间有普通文字",
         "### Before\nOld.\nCaption: a\nStray line.\n### After\nNew.\n", "Stray line."),
    ]
    for name, body, needle in cases:
        r = generate(deck(body), images=IMG_FILES)
        check(f"before-after（报错）：{name} → 非零退出、含页标题与原因",
              r.rc != 0 and not r.path.exists()
              and "Mock Broken Page" in r.stderr and needle in r.stderr,
              (r.rc, r.stderr))


def test_caption_too_long():
    """caption 超过一行（按字宽估算）只警告、照常生成；中文 20 字左右不警告。"""
    long_cn = "改前需要在三张模拟表格之间来回复制粘贴并逐行人工核对结果"   # 28 个汉字
    ok_cn = "改后一张模拟表单自动同步到全部表格"                          # 17 个汉字
    long_en = "Every mock sheet is now refreshed from one shared form within seconds"
    r = generate(deck(f"### Before\nOld.\nCaption: {long_cn}\n"
                      f"### After\nNew.\nCaption: {ok_cn}\n"), images=IMG_FILES)
    w = [x for x in warnings(r.stderr) if "Caption" in x or "caption" in x]
    check("before-after（caption）：中文过长 → 一条警告，含页标题和 BEFORE，仍生成",
          r.rc == 0 and len(w) == 1 and "Mock Broken Page" in w[0] and "BEFORE" in w[0],
          r.stderr)
    r = generate(deck(f"### Before\nOld.\nCaption: Old\n"
                      f"### After\nNew.\nCaption: {long_en}\n"), images=IMG_FILES)
    w = [x for x in warnings(r.stderr) if "Caption" in x or "caption" in x]
    check("before-after（caption）：英文过长 → 一条警告，含 AFTER",
          r.rc == 0 and len(w) == 1 and "AFTER" in w[0], r.stderr)


def test_red_only_after_title_and_bar():
    """红色只在 AFTER 标题（版式继承）和底部红条上。"""
    r = generate(LEAD_MD, images=IMG_FILES)
    s = body_slide(r)
    if s is None:
        return
    check("before-after（红色）：页面上自带红色的只有底部红条",
          red_elements(s) == ["Caption Bar"], red_elements(s))
    check("before-after（红色）：版式里红的只有 AFTER 标题，且它填的是 AFTER",
          red_elements(s.slide_layout) == ["After Title"]
          and placeholder(s, 12).text == "AFTER", red_elements(s.slide_layout))


def test_bar_does_not_cover_content():
    """有红条时，caption、内容区和图都在红条上方，不被盖住；没有红条时 caption 留在模板位置。"""
    for md, files in ((MD, IMAGES), (LEAD_MD, IMG_FILES)):
        r = generate(md, images=files)
        s = body_slide(r)
        bar = caption_bar(s)
        if bar is None:
            check("before-after（红条）：有红条", False, r.stderr)
            continue
        hidden = [sh.name for sh in s.shapes
                  if sh.name not in ("Caption Bar", "Month Tag")
                  and not (sh.is_placeholder and sh.placeholder_format.idx in (0, 10))
                  and sh.top + sh.height > bar.top]
        check(f"before-after（红条）：{placeholder(s, 0).text} 没有形状伸进红条", not hidden,
              hidden)
    r = generate(deck("### Before\nOld.\nCaption: a\n### After\nNew.\nCaption: b\n"))
    s = body_slide(r)
    E = 914400
    check("before-after（红条）：没有红条时 caption 在模板位置 6.12\"",
          caption_bar(s) is None and abs(placeholder(s, 15).top - 6.12 * E) < 2000
          and abs(placeholder(s, 16).top - 6.12 * E) < 2000,
          placeholder(s, 15).top / E)


def test_long_text_wraps_and_warns():
    """短文在槽宽内换行；估计超出槽高时警告（含页标题和 BEFORE），仍生成。"""
    para = "Mock sentence that describes one manual step of the old way. "
    long_text = "\n".join([para * 3] * 6)
    r = generate(deck(f"### Before\n{long_text}\n### After\nNew.\n"))
    s = body_slide(r)
    check("before-after（短文）：内容槽开启自动换行",
          s is not None and placeholder(s, 13).text_frame.word_wrap is True)
    w = [x for x in warnings(r.stderr) if "溢出" in x]
    check("before-after（短文）：过长 → 溢出警告，含页标题和 BEFORE",
          r.rc == 0 and len(w) == 1 and "Mock Broken Page" in w[0] and "BEFORE" in w[0],
          r.stderr)


def test_lead_too_long():
    """lead 估计超过一行时警告（含页标题和 lead）。"""
    lead = "这是一句很长很长的模拟正文导语，用来说明改前改后的差别以及证据来源和统计口径，" * 2
    r = generate(deck(f"{lead}\n### Before\nOld.\n### After\nNew.\n"))
    w = [x for x in warnings(r.stderr) if "lead" in x.lower()]
    check("before-after（lead）：过长 → 一条警告，含页标题，仍生成",
          r.rc == 0 and len(w) == 1 and "Mock Broken Page" in w[0], r.stderr)


if __name__ == "__main__":
    test_chrome()
    test_text_slots()
    test_image_slots()
    test_lead_and_conclusion()
    test_errors()
    test_caption_too_long()
    test_red_only_after_title_and_bar()
    test_bar_does_not_cover_content()
    test_long_text_wraps_and_warns()
    test_lead_too_long()
    finish()

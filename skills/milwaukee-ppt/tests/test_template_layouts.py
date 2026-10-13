"""验证模板 assets/MS_Template.pptx 的版式契约（票 #13）。

- 共 8 个 layout；0–3 与基线逐项一致；4 张样例页不动。
  基线写死在本文件里（取自提交 702cd8a 的模板，即 10-08 的 Before/After 几何），
  不依赖 git，单独拷出 skill 目录也能跑。
- 4–7 是 Process / Stats / Timeline / Image + Text：idx 0/10 与 layout 0 同位置，
  命名 BODY 槽位在 banner 以下、红条最高位以上，左右边距不小于 0.5 英寸。
- 新 layout 不预画形状（只有占位框），banner 以外没有新增红色。

运行：python3 skills/milwaukee-ppt/tests/test_template_layouts.py
"""
import hashlib
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

SKILL = Path(__file__).resolve().parent.parent
TEMPLATE = SKILL / "assets" / "MS_Template.pptx"

# 基线（EMU）：layout 0–3 的 (版式名, [(idx, 占位框名, (L, T, W, H))])
_BANNER_TITLES = [(0, "Title 1", (4644601, 87850, 7543800, 369350)),
                  (10, "Text Placeholder 13", (4644601, 457200, 7543800, 381000))]
BASELINE_LAYOUTS = [
    ("Blank  Slide", _BANNER_TITLES),
    ("Title Slide", [(0, "Rectangle 3", (914400, 4419600, 10363200, 457200)),
                     (1, "Rectangle 4", (1828800, 4953000, 8534400, 457200))]),
    ("Blank  Slide WITH BULLETS",
     [_BANNER_TITLES[0], (1, "Content Placeholder 2", (508000, 1295400, 11379200, 5257800)),
      _BANNER_TITLES[1]]),
    ("Before / After", _BANNER_TITLES + [
        (11, "Before Title", (1860804, 1792224, 2651760, 393192)),
        (12, "After Title", (7676388, 1792224, 2651760, 393192)),
        (13, "Before Content", (713232, 2542032, 4946904, 2660904)),
        (14, "After Content", (6528816, 2542032, 4946904, 2660904)),
        (15, "Before Caption", (512064, 5596128, 5349240, 320040)),
        (16, "After Caption", (6327648, 5596128, 5349240, 320040))]),
]
# layout 3 上的原生形状（对比框和中间箭头）
BASELINE_BA_SHAPES = [("Before Frame", (512064, 2331720, 5349240, 3081528)),
                      ("After Frame", (6327648, 2331720, 5349240, 3081528)),
                      ("Conversion", (5879592, 3584448, 429768, 457200))]
# 4 张样例页：(版式名, slide XML 的 sha256)
BASELINE_SLIDES = [
    ("Title Slide", "8808d2c1fc8060e1399a572aa9d347d0693c87ca15ecea5de52ef6d7a119e12f"),
    ("Blank  Slide", "1217054dea7395fd41f10d197e79249c43ab84021fc81fc0566055c503e1fdaa"),
    ("Blank  Slide WITH BULLETS",
     "0443c02ee03c6d8ea3394352a63dbd1b9a81f64dff9e2392cd1b6253e466e10c"),
    ("Before / After", "d41a8be7e5a3add9a1a02d6840345a15d6194988d986982b4c6334278bbe7fcc"),
]

EMU_IN = 914400
SLIDE_W_IN = 13.333
BANNER_BOTTOM_IN = 1.0
SLOT_BOTTOM_LIMIT_IN = 5.6  # 证据页红条最高位（=图片区底缘），槽位不能压到它下面
MIN_MARGIN_IN = 0.5

# index -> (名字, {idx: 槽位名})
NEW_LAYOUTS = {
    4: ("Process", {17: "Process Area"}),
    5: ("Stats", {18: "Stats Area"}),
    6: ("Timeline", {19: "Timeline Area"}),
    7: ("Image + Text", {20: "Image Frame", 21: "Text Area"}),
}

failures = []


def check(cond, msg):
    if not cond:
        failures.append(msg)


def geom(shape):
    return (shape.left, shape.top, shape.width, shape.height)


def describe_layout(layout):
    return (layout.name, sorted(
        (ph.placeholder_format.idx, ph.name, geom(ph)) for ph in layout.placeholders))


def is_grey(hex6):
    r, g, b = hex6[0:2], hex6[2:4], hex6[4:6]
    return r == g == b


prs = Presentation(str(TEMPLATE))
layouts = list(prs.slide_layouts)

# 1. 一共 8 个 layout
check(len(layouts) == 8, f"layout 数应为 8，实际 {len(layouts)}")

# 2. layout 0–3 与基线一致
for i, want in enumerate(BASELINE_LAYOUTS):
    if i >= len(layouts):
        break
    got = describe_layout(layouts[i])
    check(got == (want[0], sorted(want[1])), f"layout {i} 与基线不一致：{got} != {want}")
if len(layouts) > 3:
    got = [(s.name, geom(s)) for s in layouts[3].shapes if not s.is_placeholder]
    check(got == BASELINE_BA_SHAPES, f"layout 3 的对比框 / 箭头与基线不一致：{got}")

# 3. 4 张样例页不动（版式名与 XML 完全一致）
check(len(prs.slides) == len(BASELINE_SLIDES),
      f"样例页数应为 {len(BASELINE_SLIDES)}，实际 {len(prs.slides)}")
for i, (slide, (layout_name, digest)) in enumerate(zip(prs.slides, BASELINE_SLIDES)):
    check(slide.slide_layout.name == layout_name, f"样例页 {i} 版式变了")
    check(hashlib.sha256(slide.part.blob).hexdigest() == digest, f"样例页 {i} 内容变了")

# 4. 新 layout
ref = {ph.placeholder_format.idx: geom(ph) for ph in layouts[0].placeholders}
for index, (name, slots) in NEW_LAYOUTS.items():
    if index >= len(layouts):
        failures.append(f"缺少 layout {index}（{name}）")
        continue
    lay = layouts[index]
    tag = f"layout {index}（{name}）"
    check(lay.name == name, f"{tag} 名字是 {lay.name!r}")

    phs = {ph.placeholder_format.idx: ph for ph in lay.placeholders}
    for idx in (0, 10):
        check(idx in phs and geom(phs[idx]) == ref[idx], f"{tag} idx {idx} 与 layout 0 位置不同")
    check(set(phs) == {0, 10, *slots}, f"{tag} 占位框 idx 应为 {{0, 10, {sorted(slots)}}}，实际 {sorted(phs)}")

    # 不预画形状：layout 上只有占位框
    extra = [s.name for s in lay.shapes if not s.is_placeholder]
    check(not extra, f"{tag} 预画了非占位框形状：{extra}")

    for idx, slot_name in slots.items():
        ph = phs.get(idx)
        if ph is None:
            continue
        check(ph.name == slot_name, f"{tag} idx {idx} 名字是 {ph.name!r}，应为 {slot_name!r}")
        check(ph.placeholder_format.type == PP_PLACEHOLDER.BODY, f"{tag} idx {idx} 不是 BODY")
        l, t, w, h = (v / EMU_IN for v in geom(ph))
        check(t >= BANNER_BOTTOM_IN, f"{tag} {slot_name} 顶边 {t:.3f} 压 banner")
        check(t + h <= SLOT_BOTTOM_LIMIT_IN + 1e-6, f"{tag} {slot_name} 底边 {t + h:.3f} 超过 {SLOT_BOTTOM_LIMIT_IN}")
        check(l >= MIN_MARGIN_IN, f"{tag} {slot_name} 左边距 {l:.3f} < 0.5")
        check(SLIDE_W_IN - (l + w) >= MIN_MARGIN_IN, f"{tag} {slot_name} 右边距 {SLIDE_W_IN - (l + w):.3f} < 0.5")

        # 槽位文字必须自带非红色：母版正文默认是红色 Wingdings 项目符号
        lvl1 = ph._element.find(".//{*}lstStyle/{*}lvl1pPr")
        check(lvl1 is not None, f"{tag} {slot_name} 没有 lvl1pPr 样式")
        if lvl1 is not None:
            fill = lvl1.find("./{*}defRPr/{*}solidFill/{*}srgbClr")
            check(fill is not None and is_grey(fill.get("val")), f"{tag} {slot_name} 文字没有显式灰阶颜色")
            sz = lvl1.find("./{*}defRPr")
            check(sz is not None and int(sz.get("sz", "0")) >= 1200, f"{tag} {slot_name} 字号 < 12pt")
            bu_none = lvl1.find("./{*}buNone") is not None
            bu_clr = lvl1.find("./{*}buClr/{*}srgbClr")
            check(bu_none or (bu_clr is not None and is_grey(bu_clr.get("val"))),
                  f"{tag} {slot_name} 项目符号会继承母版红色")

    # banner 以外没有新增红色：整份 layout XML 只允许灰阶 srgb 与 bg1/tx1（白/黑）
    xml = lay.part.blob.decode("utf-8")
    colors = set(re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', xml))
    check(all(is_grey(c) for c in colors), f"{tag} 出现非灰阶颜色：{sorted(colors)}")
    schemes = set(re.findall(r'schemeClr val="(\w+)"', xml))
    check(schemes <= {"bg1", "tx1"}, f"{tag} 引用了主题色：{sorted(schemes)}")

if failures:
    print("FAIL")
    for f in failures:
        print(" -", f)
    sys.exit(1)
print("OK：8 个 layout，0–3 与样例页未变，4 个新 layout 契约满足")

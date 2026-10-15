"""审核脚本 review_deck.py 的「红色面积超配额」（spec #47 / #52）：banner 以外的品牌红面积
超过 SKILL.md §6.3 的配额就报，月报豁免。

运行：python3 skills/milwaukee-ppt/tests/test_review_quota.py
做法：生成器按文稿出 pptx（正样本，不报），再用 python-pptx 加红色块 / 红线（负样本，报）；
只看审核脚本的输出。配额值从 SKILL.md §6.3 面积表读，守住文档和脚本一致。文稿内容全部虚构。
"""
import re

from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from deckcheck import (BRAND_RED, SKILL, broken, check, finish, generated, lines_with, review)
import test_review_project as full

RULE = "红色面积超配额"
SKILL_MD = (SKILL / "SKILL.md").read_text(encoding="utf-8")

PAGES = """
# 现场证据 / EVIDENCE
## 充电位高峰排队超过 20 分钟
![](img/a.png)
早班交接时 6 台叉车同时排队充电
Next: 9 月起按电量排班
"""
GENERAL = "---\ntitle: 叉车充电排班 项目汇报\n---\n" + PAGES
MONTHLY = "---\nmonth: 2026-08\nauthor: Mock.Author\n---\n" + PAGES
IMAGES = {"img/a.png": (1600, 900)}


def quota():
    """SKILL.md §6.3 面积表「标准正文页」一行的配额（百分数）。"""
    m = re.search(r"^\| 标准正文页\s*\| ≤ (\d+)%", SKILL_MD, re.M)
    return int(m.group(1)) if m else None


def red_block(share, page=2, top_in=1.5):
    """在第 page 页 banner 下加一块占页面面积 share 的品牌红实心矩形（宽占页宽一半）。"""
    def change(prs):
        w = prs.slide_width // 2
        h = int(prs.slide_width * prs.slide_height * share / w)
        sp = prs.slides[page - 1].shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(0.6 * 914400)),
                                                   Emu(int(top_in * 914400)), w, h)
        sp.name = "My Red Block"
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(BRAND_RED)
        sp.line.fill.background()
    return change


def theme_red_block(share, page=2):
    """同 red_block，但填充写成主题色 schemeClr accent1（模板主题里 accent1 就是品牌红）。"""
    def change(prs):
        red_block(share, page)(prs)
        sp = [sh for sh in prs.slides[page - 1].shapes if sh.name == "My Red Block"][0]
        fill = sp._element.spPr.find(qn("a:solidFill"))
        for child in list(fill):
            fill.remove(child)
        fill.append(fill.makeelement(qn("a:schemeClr"), {"val": "accent1"}))
    return change


def grouped_red_block(share, page=2):
    """在组合形状里放一块超配额的品牌红矩形（组本身没有填充）。"""
    def change(prs):
        w = prs.slide_width // 2
        h = int(prs.slide_width * prs.slide_height * share / w)
        group = prs.slides[page - 1].shapes.add_group_shape()
        group.name = "My Group"
        sp = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(0.6 * 914400)),
                                    Emu(int(1.5 * 914400)), w, h)
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(BRAND_RED)
        sp.line.fill.background()
    return change


def red_lines(count, pt, page=2):
    """加 count 条横贯正文区的品牌红直线（线宽 pt）。"""
    def change(prs):
        slide = prs.slides[page - 1]
        for k in range(count):
            y = int((1.6 + 0.2 * k) * 914400)
            c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(0.6 * 914400), y,
                                           int(12.6 * 914400), y)
            c.name = f"My Red Line {k + 1}"
            c.line.color.rgb = RGBColor.from_string(BRAND_RED)
            c.line.width = Pt(pt)
    return change


def page2_hits(out):
    return [ln for ln in lines_with(out, RULE) if "第 2 页" in ln and " → " in ln]


def test_quota_in_skill_md():
    check("SKILL.md §6.3：标准正文页的配额是 5%", quota() == 5, quota())
    section = SKILL_MD.split("### 6.3", 1)[1].split("\n### 6.4", 1)[0]
    check("SKILL.md §6.3：写明审核脚本按模式查配额、月报豁免",
          RULE in section and "月报" in section and "豁免" in section, section[:200])
    check("SKILL.md §6.3：写明红色文字不计入面积（局限）", "红色文字不计入" in section)
    check("SKILL.md §6.3：面积表不再说「底部红条不计入配额」适用于所有模式",
          "月报的底部红条与其他版式相同,不计入配额" not in section)


def test_generator_output_passes():
    for label, md in (("项目模式全版式", full.PROJECT_FULL),
                      ("通用模式全版式", full.plain_full())):
        r = review(*generated(md, full.IMAGES))
        check(f"{label}：生成器原样输出不报「{RULE}」", not lines_with(r.stdout, RULE), r.stdout)


def test_big_red_block():
    share = (quota() or 5) / 100
    md, pptx = generated(GENERAL, IMAGES)
    r = review(md, broken(pptx, red_block(share + 0.01)))
    hits = page2_hits(r.stdout)
    check("通用模式加一块超配额的大红块：报「红色面积超配额」，带第 2 页和占比",
          len(hits) == 1 and "%" in hits[0] and "REVIEW: " in r.stdout, r.stdout)
    r = review(md, broken(pptx, red_block(share - 0.01)))
    check("通用模式加一块配额以内的红块：不报", not lines_with(r.stdout, RULE), r.stdout)
    project = GENERAL.replace("title:", "report: project\ntitle:", 1).replace(
        "## 充电位", "PDCA: P\n## 充电位")
    md, pptx = generated(project, IMAGES)
    r = review(md, broken(pptx, red_block(share + 0.01)))
    check("项目模式加大红块：报", bool(page2_hits(r.stdout)), r.stdout)


def test_theme_and_group_red_counted():
    share = (quota() or 5) / 100
    md, pptx = generated(GENERAL, IMAGES)
    r = review(md, broken(pptx, theme_red_block(share + 0.01)))
    check("大红块填充写成主题色 accent1（按母版主题解析成品牌红）：报",
          len(page2_hits(r.stdout)) == 1, r.stdout)
    r = review(md, broken(pptx, theme_red_block(share - 0.01)))
    check("主题色红块在配额以内：不报", not lines_with(r.stdout, RULE), r.stdout)
    r = review(md, broken(pptx, grouped_red_block(share + 0.01)))
    check("大红块放进组合形状（grpSp，递归算）：报", len(page2_hits(r.stdout)) == 1, r.stdout)
    r = review(md, broken(pptx, grouped_red_block(share - 0.01)))
    check("组合形状里的红块在配额以内：不报", not lines_with(r.stdout, RULE), r.stdout)


def test_red_lines_count_by_length():
    md, pptx = generated(GENERAL, IMAGES)
    # 12" 长、6pt 宽的线一条约 0.0833 平方英寸 × 12 = 1 平方英寸 ≈ 1% 页面；10 条 ≈ 10%
    r = review(md, broken(pptx, red_lines(10, 6)))
    check("加 10 条 6pt 红线（按长 × 线宽约 10%）：报", bool(page2_hits(r.stdout)), r.stdout)
    r = review(md, broken(pptx, red_lines(2, 1)))
    check("加 2 条 1pt 细红线（不按外框算）：不报", not lines_with(r.stdout, RULE), r.stdout)


def test_banner_not_counted():
    md, pptx = generated(GENERAL, IMAGES)
    r = review(md, broken(pptx, red_block(0.06, top_in=0.0)))
    hits = [ln for ln in lines_with(r.stdout, RULE)]
    check("红块大半压在 banner 上：只按 banner 以下部分算，不报", not hits, r.stdout)


PHASES = """---
title: 叉车充电排班 项目汇报
---

# 推广计划 / ROADMAP
## 六个阶段都还没开始，9 月立项
Layout: timeline-phase
- Sep | 立项 | 确定范围 | 未完成
- Oct | 开发 | 排班功能 | 未完成
- Nov | 试运行 | B 仓 | 未完成
- Dec | 推广 | A 仓 | 未完成
- Jan | 复盘 | 汇总数据 | 未完成
- Feb | 验收 | 正式验收 | 未完成
9 月立项后按月推进
"""


def test_timeline_status_not_counted():
    r = review(*generated(PHASES))
    check("阶段条全部未完成（红描边状态色，§6.3 例外）：不报", not lines_with(r.stdout, RULE),
          r.stdout)


def test_monthly_exempt():
    md, pptx = generated(MONTHLY, IMAGES)
    r = review(md, broken(pptx, red_block(0.08)))
    check("月报模式加大红块：不报「红色面积超配额」", not lines_with(r.stdout, RULE), r.stdout)


if __name__ == "__main__":
    test_quota_in_skill_md()
    test_generator_output_passes()
    test_big_red_block()
    test_theme_and_group_red_counted()
    test_red_lines_count_by_length()
    test_banner_not_counted()
    test_timeline_status_not_counted()
    test_monthly_exempt()
    finish()
